"""Convert the shared, sanitized Markdown tree to native Word elements."""
import re
import unicodedata
from copy import deepcopy
from lxml import html
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from safe_markup import md, ordered_list_start
from .pagination import keep_table_together
from .numbering import WordList
from .images import word_image, fit_width, figure_layout, paragraph_width, align_inline_picture
from report_layout import table_layout


def _link_script(character):
    """Separate common CJK/complex font groups without relabelling text.

    This local grouping covers the reproduced Western/CJK, Arabic, Indic and
    Hangul transitions. It does not emulate every Writer/Unicode script rule.
    """
    value = ord(character)
    if any(first <= value <= last for first, last in (
        (0x1100, 0x11FF),
        (0x2E80, 0x2FFF), (0x3000, 0x303F), (0x3040, 0x30FF),
        (0x3100, 0x312F), (0x3130, 0x318F), (0x31A0, 0x31FF),
        (0x3400, 0x4DBF), (0x4E00, 0x9FFF), (0xA960, 0xA97F),
        (0xAC00, 0xD7FF), (0xF900, 0xFAFF), (0xFE30, 0xFE4F),
        (0xFF66, 0xFFDC), (0x1B000, 0x1B2FF), (0x20000, 0x323AF),
    )):
        return 'asian'
    if (unicodedata.bidirectional(character) in ('R', 'AL')
            or any(first <= value <= last for first, last in (
                (0x0900, 0x109F), (0x1700, 0x18AF), (0x1900, 0x1CFF),
                (0xA800, 0xA95F), (0xA980, 0xABFF), (0x11000, 0x11FFF),
            ))):
        return 'complex'
    return 'western'


def _link_cluster_starts(text):
    """Safe local boundaries for marks, selectors, modifiers and ZWJ chains.

    This is deliberately scoped to inline link segmentation, not a general
    Unicode grapheme implementation. The original characters are never changed.
    """
    for index, character in enumerate(text):
        if index and (unicodedata.category(character).startswith('M')
                      or character == '\u200d' or text[index - 1] == '\u200d'
                      or 0x1F3FB <= ord(character) <= 0x1F3FF):
            continue
        yield index


def _separate_hyperlink_runs(link):
    """Keep external fragment links clickable after Writer's DOCX save.

    Writer exports these links as HYPERLINK fields and may close a field at
    the first font/script run. Adjacent native hyperlinks retain the same
    destination and character formatting without adding text or break hints.
    Nontext children remain whole. No new text/run boundary is introduced inside
    a cluster; existing style-run boundaries receive their own link wrapper.
    """
    text_tag, properties_tag, run_tag = qn('w:t'), qn('w:rPr'), qn('w:r')
    def text_run(child):
        return (child.tag == run_tag
                and all(node.tag in (text_tag, properties_tag) for node in child)
                and any(node.tag == text_tag and node.text for node in child))
    groups, pending = [], []
    for child in link:
        if text_run(child):
            pending.append(child)
        else:
            if pending:
                groups.append(pending); pending = []
            groups.append([child])
    if pending:
        groups.append(pending)
    segments = []
    for group in groups:
        if not all(text_run(child) for child in group):
            segments.append(group)
            continue
        values = [''.join(node.text or '' for node in run if node.tag == text_tag) for run in group]
        text = ''.join(values)
        starts = list(_link_cluster_starts(text))
        boundaries, offset = {0, len(text)}, 0
        for value in values:
            offset += len(value)
            # A style boundary already exists here. Keeping its runs inside
            # one hyperlink would still let Writer unlink a styled combining
            # mark, even though the preceding base retains its destination.
            boundaries.add(offset)
        for first, second in zip(starts, starts[1:]):
            if _link_script(text[first]) != _link_script(text[second]):
                boundaries.add(second)
        boundaries = sorted(boundaries)
        if len(boundaries) <= 2:
            segments.append(group)
            continue
        run_index, run_offset = 0, 0
        for first, last in zip(boundaries, boundaries[1:]):
            part = []
            while run_index < len(group) and run_offset < last:
                run, value = group[run_index], values[run_index]
                end = run_offset + len(value)
                if first < end and last > run_offset:
                    if first <= run_offset and end <= last:
                        part.append(run)
                    else:
                        clone = deepcopy(run)
                        node_offset = run_offset
                        for node in list(clone):
                            if node.tag != text_tag:
                                continue
                            content = node.text or ''
                            low, high = max(0, first - node_offset), min(len(content), last - node_offset)
                            if low < high:
                                node.text = content[low:high]
                                if node.text[:1].isspace() or node.text[-1:].isspace():
                                    node.set(qn('xml:space'), 'preserve')
                            else:
                                clone.remove(node)
                            node_offset += len(content)
                        part.append(clone)
                if end <= last:
                    run_offset = end
                    run_index += 1
                else:
                    break
            segments.append(part)
    if len(segments) <= 1:
        return
    parent, position = link.getparent(), link.getparent().index(link)
    for segment in segments:
        clone = OxmlElement('w:hyperlink')
        clone.attrib.update(link.attrib)
        for child in segment:
            clone.append(child)
        parent.insert(position, clone)
        position += 1
    parent.remove(link)


def write_markdown(document, text, root, code_block, keep_tables=True, retest=False, column_weights=None):
    tree = html.fragment_fromstring(str(md(text, img_base="")), create_parent="div")
    available = document.sections[-1].page_width - document.sections[-1].left_margin - document.sections[-1].right_margin
    flowing_picture_paragraphs = set()

    def picture(node, paragraph=None, parent=None, caption=None, href=None):
        if not node.get("src"):
            return
        source, size = word_image(node.get("src"), root)
        paragraph = paragraph or document.add_paragraph()
        containing_width = paragraph_width(document, paragraph)
        width = min(containing_width, Cm(15))
        requested = re.search(r"width:\s*([\d.]+)(%|cm|mm|in|px|pt)", node.get("style", ""))
        if requested:
            value, unit = float(requested[1]), requested[2]
            factor = containing_width / 100 if unit == "%" else Cm(1) * {"cm": 1, "mm": .1, "in": 2.54, "px": 2.54/96, "pt": 2.54/72}[unit]
            # Clamp before integer conversion: valid decimal syntax can exceed
            # floating-point range, while the page width is always finite.
            width = int(min(containing_width, value * factor))
        caption_space, keep_caption = figure_layout(document,
            caption if caption is not None else node.get('alt', ''), width=containing_width)
        width = fit_width(document, size, width, caption_space)
        run = paragraph.add_run()
        shape = align_inline_picture(run.add_picture(source, width=max(Cm(.1), width)))
        if parent is not None: parent.append(run._r)
        shape._inline.docPr.set("descr", node.get("alt", ""))
        if href:
            if href.startswith('#'):
                from urllib.parse import unquote
                from .navigation import reference_name
                href = '#' + reference_name(unquote(href[1:]))
            link = OxmlElement('a:hlinkClick')
            link.set(qn('r:id'), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            shape._inline.docPr.append(link)
        # In a table cell the image and caption share this paragraph. Let an
        # oversized caption continue while the inline drawing remains whole.
        paragraph.paragraph_format.keep_together = keep_caption
        paragraph.paragraph_format.keep_with_next = True
        if not keep_caption:
            flowing_picture_paragraphs.add(paragraph._p)

    def inline(node, paragraph, bold=False, italic=False, code=False, parent=None):
        def add(value):
            if not value: return
            run = paragraph.add_run(value)
            run.bold = bold; run.italic = italic
            if code:
                run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(9)
            if parent is not None:
                if parent.tag == qn('w:hyperlink'):
                    style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink')
                    run._r.get_or_add_rPr().append(style)
                parent.append(run._r)
        def linebreak():
            run = paragraph.add_run(); run.add_break()
            if parent is not None: parent.append(run._r)
        add(node.text)
        for child in node:
            tag = child.tag
            if tag == "br": linebreak()
            elif tag == "img": picture(child, paragraph, parent)
            elif tag == "figure":
                # Cells retain block figures inside one Word paragraph. Their
                # captions must stay separated from any following prose or
                # figure even when the authored Markdown has no spaces.
                inline(child, paragraph, bold, italic, code, parent=parent)
                linebreak()
            elif tag == "a" and child.get("href"):
                href = child.get('href')
                def linked_contents(contents, link_bold=bold, link_italic=italic, link_code=code):
                    link = OxmlElement("w:hyperlink")
                    if href.startswith('#'):
                        from urllib.parse import unquote
                        from .navigation import reference_name
                        link.set(qn('w:anchor'), reference_name(unquote(href[1:])))
                    else:
                        link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
                    paragraph._p.append(link)
                    inline(contents, paragraph, link_bold, link_italic, link_code, parent=link)
                    if not href.startswith('#') and '#' in href:
                        _separate_hyperlink_runs(link)
                if child.xpath('.//figure'):
                    # A table cell keeps the drawing and caption in one
                    # paragraph. Link their authored text fragments separately:
                    # layout breaks must not become part of the link label, and
                    # the drawing carries its own clickable destination.
                    parts = list(paragraph_parts(child, preserve_whitespace=True))
                    for index, part in enumerate(parts):
                        if part.tag != 'figure':
                            if part.text_content().strip():
                                linked_contents(part)
                            continue
                        caption = part.find('figcaption')
                        for img in part.findall('img'):
                            picture(img, paragraph, caption=caption.text_content() if caption is not None else None,
                                    href=href)
                        if caption is not None:
                            linebreak()
                            caption = deepcopy(caption)
                            # A literal separator between adjacent figures is
                            # part of the preceding link label. Keep it even
                            # though ordinary paragraph splitting discards an
                            # otherwise empty whitespace-only fragment.
                            if index + 1 < len(parts):
                                following = parts[index + 1]
                                if following.tag != 'figure' and not following.text_content().strip():
                                    separator = following.text_content()
                                    if len(caption):
                                        caption[-1].tail = (caption[-1].tail or '') + separator
                                    else:
                                        caption.text = (caption.text or '') + separator
                            wrappers = figure_wrappers.get(part, ())
                            linked_contents(caption,
                                bold or any(n.tag in {'b', 'strong'} for n in wrappers),
                                italic or any(n.tag in {'i', 'em'} for n in wrappers),
                                code or any(n.tag == 'code' for n in wrappers))
                            if any(p.tag == 'figure' or p.text_content().strip() for p in parts[index + 1:]):
                                linebreak()
                else:
                    linked_contents(child)
            elif tag == "figcaption":
                linebreak()
                inline(child, paragraph, bold, italic, code, parent=parent)
            else:
                inline(child, paragraph, bold or tag in {"b", "strong"}, italic or tag in {"i", "em"}, code or tag == "code", parent=parent)
            add(child.tail)

    figure_wrappers = {}

    def paragraph_parts(node, wrappers=(), preserve_whitespace=False):
        # The shared Markdown tree promotes images to block figures even
        # inside links/emphasis in sentences and list paragraphs. Split at any
        # inline depth, retaining formatting on either side and on the caption.
        fragment = html.Element(node.tag, dict(node.attrib))
        fragment.text = node.text
        for child in node:
            if child.tag == 'figure':
                if (fragment.text_content() if preserve_whitespace else fragment.text_content().strip()) or len(fragment):
                    yield fragment
                figure_wrappers[child] = wrappers
                yield child
                fragment = html.Element(node.tag, dict(node.attrib))
                fragment.text = child.tail
            elif child.xpath('.//figure'):
                for part in paragraph_parts(child, wrappers + (child,), preserve_whitespace):
                    if part.tag == 'figure':
                        if (fragment.text_content() if preserve_whitespace else fragment.text_content().strip()) or len(fragment):
                            yield fragment
                        yield part
                        fragment = html.Element(node.tag, dict(node.attrib))
                    else:
                        fragment.append(part)
                if len(fragment):
                    fragment[-1].tail = (fragment[-1].tail or '') + (child.tail or '')
                else:
                    fragment.text = (fragment.text or '') + (child.tail or '')
            else:
                fragment.append(deepcopy(child))
        if (fragment.text_content() if preserve_whitespace else fragment.text_content().strip()) or len(fragment):
            yield fragment

    def block(node, depth=0, list_paragraph=None, list_offset=0):
        tag = node.tag
        if tag == 'p' and node.xpath('.//figure'):
            for part in paragraph_parts(node):
                block(part, depth, list_offset=list_offset)
            return
        if tag in {"ul", "ol"}:
            previous = node.getprevious()
            keep_short_labelled_list = (previous is not None and previous.tag == 'p'
                and not (previous.text or '').strip() and len(previous) == 1
                and previous[0].tag in {'strong', 'b'} and not (previous[0].tail or '').strip()
                and len(node) <= 4 and len(node.text_content()) <= 240
                and not node.xpath('.//ul|.//ol|.//table|.//pre|.//figure|.//img')
                and len(node.xpath('.//br')) <= 4)
            first_paragraph = len(document.paragraphs)
            numbering = WordList(document, tag == "ol",
                ordered_list_start(node) if tag == 'ol' else 1, len(node), depth, list_offset)
            nested_offset = list_offset + numbering.extra_indent
            for index, item in enumerate(node):
                p = document.add_paragraph()
                numbering.add(p, index)
                started = bool((item.text or '').strip())
                if started: p.add_run(item.text)
                current = p
                def continuation():
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.left_indent = numbering.indent
                    return paragraph
                for child in item:
                    if child.tag in {"ul", "ol"}:
                        block(child, depth+1, list_offset=nested_offset); current = None; started = True
                    elif child.tag == "p":
                        for part in paragraph_parts(child):
                            if part.tag == 'figure':
                                block(part, depth+1, p if not started else None, nested_offset)
                            else:
                                target = p if not started else continuation()
                                inline(part, target)
                            started = True; current = None
                    elif child.tag in {'pre', 'table', 'blockquote', 'figure', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'}:
                        if child.tag != 'figure' and not started:
                            # A block-led item has an ordinal paragraph before
                            # its first content block. Keep only this boundary,
                            # never an entire long table or transcript.
                            p.paragraph_format.keep_with_next = True
                        block(child, depth+1, p if child.tag == 'figure' and not started else None, nested_offset)
                        current = None; started = True
                    elif child.xpath('.//figure'):
                        wrapper = html.Element('p')
                        wrapper.append(deepcopy(child)); wrapper[0].tail = None
                        for part in paragraph_parts(wrapper):
                            if part.tag == 'figure':
                                block(part, depth+1, p if not started else None, nested_offset); current = None
                            else:
                                if current is None: current = continuation()
                                inline(part, current)
                            started = True
                    else:
                        if current is None: current = continuation()
                        wrapper = html.Element("span"); wrapper.append(html.fromstring(html.tostring(child, with_tail=False)))
                        inline(wrapper, current); started = True
                    if child.tail and (child.tail.strip() or current is not None):
                        if current is None: current = continuation()
                        current.add_run(child.tail); started = True
            if keep_short_labelled_list:
                # A short affected-assets list should not leave its last IP
                # alone on a page before the next major section. Bound the
                # chain; arbitrary long/nested lists must still paginate.
                paragraphs = document.paragraphs[first_paragraph:]
                for index, paragraph in enumerate(paragraphs):
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.keep_with_next = index < len(paragraphs)-1
        elif tag == "pre":
            text = node.text_content().removesuffix("\n")
            p = code_block(text)
            if depth:
                from .pagination import code_fits_page
                p.paragraph_format.left_indent = Cm(.65 * depth) + list_offset
                p.paragraph_format.keep_together = code_fits_page(document, text, width=paragraph_width(document, p))
        elif tag == "table":
            rows = node.xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr")
            cols = max([len(row) for row in rows]+[1])
            table = document.add_table(rows=0, cols=cols); table.style = "Table Grid"; table.autofit = False
            shares, _ = table_layout([[cell.text_content() for cell in row] for row in rows], column_weights)
            indent = Cm(.65 * depth) + list_offset
            table_width = max(Cm(1), available - indent)
            if depth:
                indentation = OxmlElement('w:tblInd')
                indentation.set(qn('w:type'), 'dxa')
                indentation.set(qn('w:w'), str(round(indent / 635)))
                table._tbl.tblPr.append(indentation)
            widths = [int(table_width * share / 100) for share in shares]
            for col, width in zip(table.columns, widths): col.width = width
            width_el = table._tbl.tblPr.find(qn("w:tblW")); width_el.set(qn("w:type"), "dxa"); width_el.set(qn("w:w"), str(int(table_width/635)))
            for row in rows:
                cells = table.add_row().cells
                for i, cell in enumerate(row):
                    cells[i].width = widths[i]
                    paragraph = cells[i].paragraphs[0]
                    # The shared Markdown parser carries each column's colon
                    # alignment on its cells. Preserve it in editable Word
                    # tables, including headers and numeric evidence columns.
                    alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                                 "center": WD_ALIGN_PARAGRAPH.CENTER,
                                 "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(cell.get("align"))
                    if alignment is not None:
                        paragraph.alignment = alignment
                    inline(cell, paragraph, bold=cell.tag == "th")
                if any(cell.tag == "th" for cell in row):
                    repeat = OxmlElement("w:tblHeader"); table.rows[-1]._tr.get_or_add_trPr().append(repeat)
            # As in the summary/inventory exporter, long tables must flow into
            # the remaining page space. Chaining 80+ rows with keepNext forces
            # a mostly empty preceding page before Word/Writer finally splits.
            keep_table_together(table, whole=keep_tables and len(table.rows) <= 24)
            # The generic table policy protects row paragraphs. An evidence
            # cell with a multipage caption is explicitly allowed to flow;
            # otherwise that policy restores the impossible keep-lines chain.
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if paragraph._p in flowing_picture_paragraphs:
                            paragraph.paragraph_format.keep_together = False
                            paragraph.paragraph_format.keep_with_next = False
                            row._tr.get_or_add_trPr().find(qn('w:cantSplit')).set(qn('w:val'), '0')
            document.add_paragraph()
        elif tag == "figure":
            if retest and document.paragraphs:
                document.paragraphs[-1].paragraph_format.keep_with_next = True
            caption = ' '.join(node.xpath('./figcaption//text()'))
            wrappers = figure_wrappers.get(node, ())
            href = next((part.get('href') for part in reversed(wrappers)
                         if part.tag == 'a' and part.get('href')), None)
            for child in node:
                if child.tag == "img":
                    picture_paragraph = list_paragraph or document.add_paragraph()
                    if depth and list_paragraph is None:
                        picture_paragraph.paragraph_format.left_indent = Cm(.65 * depth) + list_offset
                    picture(child, picture_paragraph, caption=caption, href=href)
                elif child.tag == "figcaption":
                    contents = deepcopy(child)
                    for wrapper in reversed(wrappers):
                        outer = html.Element(wrapper.tag, dict(wrapper.attrib))
                        outer.text = contents.text
                        for nested in list(contents): outer.append(nested)
                        contents = html.Element('span'); contents.append(outer)
                    p = document.add_paragraph(style="Caption"); inline(contents, p)
                    if depth: p.paragraph_format.left_indent = Cm(.65 * depth) + list_offset
                    _, keep_caption = figure_layout(document, p.text, width=paragraph_width(document, p))
                    p.paragraph_format.keep_together = keep_caption
        elif tag == "img": picture(node)
        elif tag == "blockquote":
            for child in node: block(child, depth, list_offset=list_offset)
        elif tag == "hr": document.add_paragraph(" ")
        elif tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6"}:
            p = document.add_paragraph(style=f"Heading {min(int(tag[1]), 4)}" if tag.startswith("h") else None)
            if depth: p.paragraph_format.left_indent = Cm(.65 * depth) + list_offset
            # A bold list label (e.g. generated affected assets) belongs with
            # the first item; otherwise it can be orphaned at the page bottom.
            if (tag == 'p' and not (node.text or '').strip() and len(node) == 1
                    and node[0].tag in {'strong', 'b'} and not (node[0].tail or '').strip()
                    and node.getnext() is not None and node.getnext().tag in {'ul', 'ol'}):
                p.paragraph_format.keep_with_next = True
            if retest and node.getprevious() is not None and node.getprevious().tag == 'h3':
                p.paragraph_format.keep_with_next = True
            # Generated retest starts with counts followed by its register.
            # Keep this short boundary with the table header/first row, not
            # the entire register or any later authored narrative paragraph.
            if (retest and tag == 'p' and node.getparent() is tree
                    and node.getprevious() is None and node.getnext() is not None
                    and node.getnext().tag == 'table'):
                p.paragraph_format.keep_with_next = True
            if (tag == 'p' and node.getnext() is not None and node.getnext().tag == 'pre'
                    and len(node.text_content()) <= 160 and node.xpath('.//strong|.//b')):
                p.paragraph_format.keep_with_next = True
            inline(node, p)
        else:
            p = document.add_paragraph(); inline(node, p)
    if tree.text and tree.text.strip(): document.add_paragraph(tree.text)
    for node in tree: block(node)
