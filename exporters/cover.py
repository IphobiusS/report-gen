"""Native Word cover with independent header, title and lower metadata zones."""
from functools import lru_cache
from pathlib import Path

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .images import word_image


@lru_cache(maxsize=2)
def _identity_font(bold=False):
    from PIL import ImageFont
    # Carlito is the bundled metric-compatible substitute for Word's Calibri.
    name = 'Carlito-Bold.ttf' if bold else 'Carlito-Regular.ttf'
    return ImageFont.truetype(str(Path(__file__).resolve().parents[1] / 'themes' / 'fonts' / name), 100)


def _identity_needs_columns(wordmark, classification, width):
    """Keep short covers identical; isolate text that cannot share one line."""
    if '\n' in wordmark + classification or '\r' in wordmark + classification:
        return True
    points = (_identity_font(True).getlength(wordmark) * 15 / 100
              + _identity_font().getlength(classification) * 8 / 100)
    return points + Cm(.5) / Pt(1) > width / Pt(1)


def add_cover(doc, meta, labels, engagement_dir, hx):
    design = meta.get('design')
    if design:
        from design_library import validate_design
        design=validate_design(design)
    theme = ('profile-'+design['layout']) if design else meta.get('theme', 'serio')
    branding = meta.get('branding') or {}
    illustration = bool(branding.get('cover_image'))
    bg = hx('cover-bg', 'ffffff')
    body_color = hx('cover-fg', '182b3b')
    title_color = hx('cover-title', '182b3b')
    muted = hx('cover-muted', '566777')
    accent = hx('cover-rule', '1c5965')
    line = hx('cover-meta-border', 'ced6dd')
    grid = bool(design) or theme in ('ejecutivo', 'corporativo', 'laboratorio')
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    for name in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, name, Cm(0))

    def shade(cell, color):
        for previous in cell._tc.get_or_add_tcPr().findall(qn('w:shd')):
            previous.getparent().remove(previous)
        e = OxmlElement('w:shd'); e.set(qn('w:fill'), color)
        cell._tc.get_or_add_tcPr().append(e)

    def border(cell, **edges):
        node = cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
        if node is None:
            node = OxmlElement('w:tcBorders')
            cell._tc.get_or_add_tcPr().append(node)
        for name, value in edges.items():
            color, size = value
            for previous in node.findall(qn('w:'+name)): node.remove(previous)
            e = OxmlElement('w:'+name); e.set(qn('w:val'), 'single')
            e.set(qn('w:color'), color); e.set(qn('w:sz'), str(size)); node.append(e)

    def margins(cell, top=0, bottom=0, start=0, end=0):
        for previous in cell._tc.get_or_add_tcPr().findall(qn('w:tcMar')):
            previous.getparent().remove(previous)
        node = OxmlElement('w:tcMar')
        for side, value in dict(top=top, bottom=bottom, start=start, end=end).items():
            e = OxmlElement('w:'+side); e.set(qn('w:w'), str(round(value*56.6929)))
            e.set(qn('w:type'), 'dxa'); node.append(e)
        cell._tc.get_or_add_tcPr().append(node)

    def paragraph(cell, first=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        return p

    def run(p, text, size=10, color=body_color, bold=False):
        r = p.add_run('' if text is None else str(text)); r.font.name = 'Calibri'; r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color); r.bold = bold
        return r

    rail = (12 if design['layout']=='rail' else (3 if design['layout']=='timeline' else 0)) if design else (10 if theme=='laboratorio' else (4 if theme=='serio' else 0))
    cover = doc.add_table(rows=3, cols=2 if rail else 1)
    cover.autofit = False; cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [rail/10, 21-rail/10] if rail else [21]
    for col, width in zip(cover.columns, widths): col.width = Cm(width)
    heights = [70,120,106] if design else [70,130,96]
    if design and illustration:
        # An illustrated title can exceed the old 120 mm minimum. Reserving
        # almost a full page across all three minimum-height rows then pushed
        # the complete metadata row onto another page. Give the central zone
        # room to grow, with a smaller lower blank area and a 17 mm allowance
        # for the renderer's font and paragraph metrics. No text is shrunk.
        heights = [60,130,82]
    if theme=='corporativo': heights = [67,133,96] if illustration else [83,105,108]
    elif theme=='htb': heights = [77,116,103]
    # Word's minimum row content height excludes vertical cell padding.
    heights[0] -= 20
    heights[2] -= 20 if theme=='offsec' else 18
    # Leave room for the cover's outer rules as well as the section boundary.
    heights[0] -= {'ejecutivo':5, 'htb':6, 'offsec':3}.get(theme,0)
    if design and design['layout'] in ('band','split','frame'): heights[0]-=5
    cells = []
    for row, height in zip(cover.rows, heights):
        row.height = Cm(height/10); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width); shade(cell, bg); margins(cell)
            p = paragraph(cell, first=True); run(p, '', size=1)
        if rail: shade(row.cells[0], accent)
        cell = row.cells[-1]; cells.append(cell)
        margins(cell, start=30-rail if theme=='laboratorio' else 24-rail, end=24)
    header, middle, bottom = cells
    header.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    middle.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    bottom.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    margins(header, top=20, bottom=4 if design and illustration else 0,
            start=30-rail if theme=='laboratorio' else 24-rail, end=24)
    margins(bottom, top=4 if design and illustration else 0, bottom=18,
            start=30-rail if theme=='laboratorio' else 24-rail, end=24)
    width = 15.6 if theme=='laboratorio' else 16.2
    if design:
        layout=design['layout']
        if layout in ('band','split'):
            shade(middle,hx('profile-middle','182B3B'))
            if layout=='band': border(middle,top=(accent,24))
            else:
                shade(header,hx('profile-dark','182B3B'));border(middle,bottom=(accent,32))
        if layout=='frame':
            for c in cells: border(c,start=(accent,24),end=(accent,24))
            border(header,top=(accent,24));border(bottom,bottom=(accent,24))
    if theme=='offsec':
        border(header, top=(accent,12)); border(bottom, bottom=(accent,12))
        margins(bottom, bottom=20, start=35, end=35); width = 14
    elif theme=='ejecutivo': border(header, top=('182B3B',96))
    elif theme=='htb':
        border(header, top=('9FEF00',96)); border(middle, top=('304328',4), bottom=('304328',4))
        margins(bottom, bottom=18, start=47, end=24); width=13.9
    if theme=='corporativo':
        shade(middle, '14213D'); border(middle, start=('1F5FA8',96)); title_color='FFFFFF'

    hp = paragraph(header, first=True)
    identity_width = Cm(15.6 if theme=='laboratorio' else 16.2)
    wordmark = branding.get('wordmark', 'iphobiuss')
    classification = branding.get('confidential_text',labels['default_confidential'])
    wordmark = '' if wordmark is None else str(wordmark)
    classification = '' if classification is None else str(classification)
    if _identity_needs_columns(wordmark, classification, identity_width):
        # A right tab only aligns the first line. Separate cells also keep
        # continuation lines of the classification in their own right column.
        identity = header.add_table(rows=1, cols=2)
        identity.autofit = False
        identity.alignment = WD_TABLE_ALIGNMENT.LEFT
        for column, cell, fraction in zip(identity.columns, identity.rows[0].cells, (.55, .45)):
            column.width = cell.width = int(identity_width * fraction)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        left, right = identity.rows[0].cells
        margins(left, end=5)
        run(paragraph(left, first=True), wordmark, 15, hx('cover-wordmark','182b3b'), True)
        run(paragraph(right, first=True, align=WD_ALIGN_PARAGRAPH.RIGHT), classification, 8, hx('cover-tag-fg','566777'))
        header._tc.remove(hp._p)
        tail = header.paragraphs[-1]
        tail.paragraph_format.line_spacing = Pt(1)
        tail.paragraph_format.space_after = Pt(0)
        run(tail, '', 1)
    else:
        hp.paragraph_format.tab_stops.clear_all()
        hp.paragraph_format.tab_stops.add_tab_stop(identity_width, WD_TAB_ALIGNMENT.RIGHT)
        run(hp, wordmark, 15, hx('cover-wordmark','182b3b'), True)
        hp.add_run('\t')
        run(hp, classification, 8, hx('cover-tag-fg','566777'))
    if design:
        p=paragraph(header);p.paragraph_format.space_before=Pt(16)
        from provider_branding import provider_mark
        mark = provider_mark(design)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT)
        run(p,mark['provider']+' / '+design['certificate'],10,hx('cover-mark-ink','182B3B'))
        if mark['png'] is not None:
            p.add_run('\t')
            shape=p.add_run().add_picture(str(mark['png']),width=Cm(4.2))
            shape._inline.docPr.set('descr',mark['name'])
    if branding.get('logo') or branding.get('client_logo'):
        p = paragraph(header); p.paragraph_format.space_before=Pt(10 if design else 16)
        for key in ('logo','client_logo'):
            if branding.get(key):
                src,size=word_image(branding[key], engagement_dir)
                w=min(Cm(6.2), int(Cm(1.6 if design else 1.8)*size[0]/size[1]))
                shape=p.add_run().add_picture(src,width=w)
                shape._inline.docPr.set('descr',key.replace('_',' ')); p.add_run('    ')

    align=WD_ALIGN_PARAGRAPH.LEFT if theme in ('serio','laboratorio','profile-editorial') else WD_ALIGN_PARAGRAPH.CENTER
    title_style=doc.styles['Title']
    for previous in title_style.element.xpath('./w:pPr/w:pBdr'):
        previous.getparent().remove(previous)
    title_style.font.name='Calibri'
    tp=paragraph(middle,first=True,align=align);tp.style=title_style
    tp.paragraph_format.space_after=Pt(0);tp.paragraph_format.line_spacing=1.14
    tp.paragraph_format.keep_with_next=False
    if theme == 'profile-editorial':
        rules = OxmlElement('w:pBdr')
        rule = OxmlElement('w:top')
        for key, value in {'val':'single', 'sz':'24', 'space':'22', 'color':accent}.items():
            rule.set(qn('w:' + key), value)
        rules.append(rule)
        tp._p.get_or_add_pPr().insert_element_before(
            rules, 'w:shd', 'w:tabs', 'w:spacing', 'w:ind', 'w:jc', 'w:rPr')
    size=27 if (design and illustration) else 32 if design else 28 if illustration else (34 if theme=='ejecutivo' else (33 if theme=='serio' else 31))
    run(tp,meta.get('report_title',''),size,title_color,True)
    if meta.get('report_subtitle'):
        p=paragraph(middle,align=align);p.paragraph_format.space_before=Pt(17)
        run(p,meta['report_subtitle'],12,'C6D7EC' if theme=='corporativo' else hx('cover-subtitle','1c5965'))
    if illustration:
        src,size=word_image(branding['cover_image'],engagement_dir)
        p=paragraph(middle,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.space_before=Pt(17)
        height=Cm(4 if theme=='corporativo' else 5.2)
        shape=p.add_run().add_picture(src,width=min(Cm(14.5),int(height*size[0]/size[1])))
        shape._inline.docPr.set('descr','Cover image')

    fields=[(labels['client'],meta.get('client')),(labels['cover_osid'],meta.get('osid')),
            (labels['cover_email'],meta.get('email')),
            (labels['cover_assessor'],str(meta.get('assessor') or '')+(' · '+str(meta['assessor_title']) if meta.get('assessor_title') else '')),
            (labels['cover_date'],meta.get('date')),(labels['cover_version'],meta.get('version'))]
    fields=[(k,str(v)) for k,v in fields if v not in (None,'')]
    lead=paragraph(bottom,first=True);lead.paragraph_format.line_spacing=Pt(1);run(lead,'',1)
    mt=bottom.add_table(rows=0,cols=2);mt.autofit=False
    col_widths=[width/2,width/2] if grid else [3.4,width-3.4]
    for col,w in zip(mt.columns,col_widths):col.width=Cm(w)
    if grid:
        for start in range(0,len(fields),2):
            row=mt.add_row()
            for offset,cell in enumerate(row.cells):
                cell.width=Cm(width/2);margins(cell,top=3,bottom=4,end=5)
                if theme=='laboratorio':
                    shade(cell,'FFFFFF');margins(cell,top=4,bottom=4,start=4,end=4)
                    border(cell,bottom=('DBE5DF',4))
                elif theme=='corporativo':border(cell,top=('CDD8EA',4))
                elif start==0:border(cell,top=('BAC8D2',4))
                if start+offset>=len(fields):continue
                key,value=fields[start+offset]
                p=paragraph(cell,first=True);run(p,key,9,muted)
                p=paragraph(cell);p.paragraph_format.space_before=Pt(3);run(p,value,10.5)
    else:
        for key,value in fields:
            a,b=mt.add_row().cells
            for cell,w in zip((a,b),col_widths):cell.width=Cm(w);margins(cell,top=2.8,bottom=2.8)
            border(b,bottom=(line,4))
            if theme=='offsec':border(a,bottom=(line,4))
            run(paragraph(a,first=True),key,9,muted)
            run(paragraph(b,first=True,align=WD_ALIGN_PARAGRAPH.RIGHT if theme=='htb' else WD_ALIGN_PARAGRAPH.LEFT),value,10)
    p=bottom.paragraphs[-1]
    p.paragraph_format.space_before=Pt(14);p.paragraph_format.space_after=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if theme=='htb' else (WD_ALIGN_PARAGRAPH.LEFT if theme in ('serio','laboratorio','profile-editorial') else WD_ALIGN_PARAGRAPH.CENTER)
    run(p,branding.get('confidential_text',labels['default_confidential']),7.5,muted)
