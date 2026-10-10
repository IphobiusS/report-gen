"""Native Word TOC with readable cached links and real page-reference fields.

The cache contains titles, never guessed PDF page numbers. Word/LibreOffice
can update the native fields using their own pagination.
"""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER


def reference_name(fragment):
    """Word-safe stable names shared by body bookmarks and Markdown links."""
    from hashlib import sha256
    return 'rg_ref_' + sha256(fragment.encode('utf-8')).hexdigest()[:24]


def add_reference(paragraph, fragment):
    # The empty HTML fragment is a real link to the document's beginning.
    # None means the caller has no reference to add.
    if fragment is None:
        return
    name = reference_name(fragment)
    document = paragraph.part.document
    starts = document.element.xpath('//w:bookmarkStart')
    if any(node.get(qn('w:name')) == name for node in starts):
        return
    number = str(max([int(node.get(qn('w:id'))) for node in starts] + [0]) + 1)
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), number); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), number)
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _field_part(paragraph, kind, instruction=None, hidden=False):
    run = paragraph.add_run()
    run.font.hidden = hidden
    node = OxmlElement('w:instrText' if instruction is not None else 'w:fldChar')
    if instruction is not None:
        node.set(qn('xml:space'), 'preserve')
        node.text = instruction
    else:
        node.set(qn('w:fldCharType'), kind)
        if kind == 'begin':
            node.set(qn('w:dirty'), 'true')
    run._r.append(node)


def add_toc(document, entries, title, color):
    """Insert a visible index before the first recorded body heading.

    Entries are (heading paragraph, displayed title, nesting level). TC fields
    enumerate exactly these entries, so Markdown body headings cannot silently
    change the index when a reader refreshes it.
    """
    if not entries:
        return
    from docx.text.paragraph import Paragraph

    ids = [int(n.get(qn('w:id'))) for n in document.element.xpath('//w:bookmarkStart')]
    next_id = max(ids, default=0) + 1
    links = []
    for index, (paragraph, text, level) in enumerate(entries, next_id):
        name = f'rg_toc_{index}'
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(index)); start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(index))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)
        escaped = ' '.join(text.split()).replace('\\', '\\\\').replace('"', '\\"')
        _field_part(paragraph, 'begin', hidden=True)
        _field_part(paragraph, '', f' TC "{escaped}" \\f R \\l {level} ', hidden=True)
        _field_part(paragraph, 'end', hidden=True)
        links.append((name, text, level))

    before = entries[0][0]._p
    heading = document.add_paragraph(style='TOC Heading')
    heading.add_run(title)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_after = Pt(14)
    for run in heading.runs:
        run.font.color.rgb = color; run.font.size = Pt(17); run.bold = True
    before.addprevious(heading._p)

    sdt = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr'); obj = OxmlElement('w:docPartObj')
    gallery = OxmlElement('w:docPartGallery'); gallery.set(qn('w:val'), 'Table of Contents')
    obj.append(gallery); props.append(obj); sdt.append(props)
    content = OxmlElement('w:sdtContent'); sdt.append(content)
    available = document.sections[-1].page_width - document.sections[-1].left_margin - document.sections[-1].right_margin
    for i, (name, text, level) in enumerate(links):
        style_name = f'TOC {level}'
        if style_name not in document.styles:
            style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = document.styles['Normal']
        element = OxmlElement('w:p'); content.append(element)
        p = Paragraph(element, document._body); p.style = style_name
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.keep_together = True
        p.paragraph_format.left_indent = Cm(.6 * (level - 1))
        # No dangling dotted line when the viewer has not calculated PAGEREF.
        p.paragraph_format.tab_stops.add_tab_stop(available, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
        if i == 0:
            _field_part(p, 'begin')
            _field_part(p, '', ' TOC \\f R \\h \\z ')
            _field_part(p, 'separate')
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), name)
        link.set(qn('w:history'), 'true')
        run = p.add_run(text); run.font.color.rgb = color
        link.append(run._r); p._p.append(link)
        p.add_run('\t')
        _field_part(p, 'begin')
        _field_part(p, '', f' PAGEREF {name} \\h ')
        _field_part(p, 'separate')
        # Empty until a layout engine knows the actual Word page. No fake 1s.
        p.add_run('')
        _field_part(p, 'end')
        if i == len(links) - 1:
            _field_part(p, 'end')
    before.addprevious(sdt)
    separator = document.add_page_break()
    before.addprevious(separator._p)
    settings = document.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'), 'true')
