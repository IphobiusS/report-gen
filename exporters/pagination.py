"""Word pagination: keep a table together when it fits on one page."""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from math import ceil
from functools import lru_cache
from pathlib import Path


@lru_cache(maxsize=1)
def _code_font():
    from PIL import ImageFont
    return ImageFont.truetype(str(Path(__file__).resolve().parents[1] / 'themes' / 'fonts' / 'DejaVuSansMono.ttf'), 100)


def code_fits_page(document, text, size=9, width=None):
    """Do not ask Word to keep a code paragraph larger than a page together.

    Count hard breaks and wrapped monospace lines, reserving the code inset,
    paragraph spacing and a preceding heading. Word remains the final layout
    engine; this conservative bound only controls the keep-lines instruction.
    """
    section = document.sections[-1]
    width = (width if width is not None else section.page_width - section.left_margin - section.right_margin) / Pt(1) - 14
    height = (section.page_height - section.top_margin - section.bottom_margin) / Pt(1)
    if width <= 0:
        return False
    font = _code_font()
    lines = sum(max(1, ceil(font.getlength(line.expandtabs(8)) * size / 100 / width))
                for line in text.split('\n'))
    return lines * size * 1.4 + 6 + 4 * size <= height


def _tail_fits(table, start):
    """Conservatively bound rows before creating a keep-next chain.

    Word performs final layout. One em per character and an extra line per
    paragraph deliberately overestimate ordinary prose, reserving additional
    room for the repeated header and differences between font renderers.
    """
    document = table.part.document
    section = document.sections[-1]
    normal = document.styles['Normal']
    normal_size = normal.font.size or Pt(11)
    available = (section.page_height - section.top_margin - section.bottom_margin) / Pt(1)
    indexes = list(range(start, len(table.rows)))
    if start and table.rows[0]._tr.xpath('./w:trPr/w:tblHeader'):
        indexes.insert(0, 0)
    total = 0
    for index in indexes:
        row = table.rows[index]
        height = (row.height or 0) / Pt(1)
        for cell in row.cells:
            # Embedded drawings/nested tables need native layout; do not
            # create a speculative keep-next chain around them.
            if cell._tc.xpath('.//w:drawing|.//w:pict|.//w:tbl'):
                return False
            width = (cell.width or (section.page_width - section.left_margin - section.right_margin) / len(row.cells)) / Pt(1) - 12
            if width <= 0:
                return False
            cell_height = 8
            for paragraph in cell.paragraphs:
                size = max([normal_size / Pt(1), (paragraph.style.font.size or normal_size) / Pt(1)]
                           + [(run.font.size or normal_size) / Pt(1) for run in paragraph.runs])
                line_width = width - (paragraph.paragraph_format.left_indent or 0) / Pt(1) - (paragraph.paragraph_format.right_indent or 0) / Pt(1)
                if line_width <= 0:
                    return False
                lines = sum(max(1, ceil(len(line) * size / line_width)) for line in paragraph.text.split('\n')) + 1
                spacing = paragraph.paragraph_format.line_spacing
                line_height = max(size * 1.4, spacing.pt if hasattr(spacing, 'pt') else size * (spacing or 1))
                cell_height += lines * line_height
                cell_height += (paragraph.paragraph_format.space_before or 0) / Pt(1)
                cell_height += (paragraph.paragraph_format.space_after or normal.paragraph_format.space_after or Pt(8)) / Pt(1)
            height = max(height, cell_height)
        total += height
    return total <= available * .85


def keep_table_together(table, whole=True):
    # cantSplit alone only protects individual rows. keepNext chains every
    # row to the next one, ending at the table boundary (not the next finding).
    rows = table.rows
    # Writer gives explicit Normal paragraph spacing precedence over the
    # table style. Compact body spacing must not make Table Grid rows taller
    # than the same table in comfortable density. Materialize only the table's
    # own spacing and preserve any paragraph-specific formatting.
    if table.style.style_id == 'TableGrid':
        spacing = table.style.paragraph_format
        for row in rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for field in ('space_after', 'line_spacing'):
                        if getattr(paragraph.paragraph_format, field) is None:
                            setattr(paragraph.paragraph_format, field, getattr(spacing, field))
    # A small row count does not imply a short table: a dozen rows of long
    # evidence can span several pages. Do not push that oversized chain to a
    # new page before Word finally allows it to split.
    whole = whole and _tail_fits(table, 0)
    tail_start = max(1, len(rows)-3)
    keep_tail = not whole and _tail_fits(table, tail_start)
    for row_index, row in enumerate(rows):
        props = row._tr.get_or_add_trPr()
        if props.find(qn("w:cantSplit")) is None:
            props.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for i, paragraph in enumerate(cell.paragraphs):
                paragraph.paragraph_format.keep_together = True
                header = row_index == 0 and props.find(qn("w:tblHeader")) is not None
                paragraph.paragraph_format.keep_with_next = ((whole and row_index < len(rows)-1)
                    or (keep_tail and tail_start <= row_index < len(rows)-1)
                    or (header and len(rows) > 1) or i < len(cell.paragraphs)-1)
