"""A control review is a real outline and composition, not another color preset."""
from copy import deepcopy
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
import pytest

import design_library as lib
from design_samples import sample_data
import export
import project_history
from wizard import dump_yaml


PROFILE = 'enterprise-control-review'


def test_every_catalog_profile_has_bilingual_purpose_metadata():
    for profile in lib.catalog()['profiles']:
        assert set(profile['description']) == {'es', 'en'}
        assert all(20 <= len(text) <= 180 for text in profile['description'].values())
        categories = profile['categories']
        assert len(categories) == len(set(categories))
        assert set(categories) <= {'corporate', 'red-team', 'blue-team', 'certification'}
        assert ('corporate' in categories) == (profile['provider'] == 'Enterprise')
        assert ('certification' in categories) == (profile['provider'] != 'Enterprise')


@pytest.mark.parametrize('lang', ['es', 'en'])
def test_control_review_starter_has_distinct_evidence_bound_outline(lang):
    profile = lib.profile(PROFILE)
    data = lib.new_report(profile, 'Control testing', lang, 'quiet')
    assert data['findings'] == []
    sections = data['report']['sections']
    keys = [section['key'] for section in sections]
    assert keys == ['control-summary', 'control-criteria', 'control-sampling',
                    'control-test-record', 'findings', 'control-exceptions',
                    'control-retest', 'control-evidence']
    assert all(section.get('body', 'TODO: ').startswith('TODO: ') for section in sections)
    note = lib.guidance(data)['note']
    assert ('No es una plantilla oficial' if lang == 'es' else 'Not an official template') in note


@pytest.mark.parametrize('variant', lib.VARIANTS)
def test_editorial_composition_exports_left_alignment_without_changing_content(tmp_path, variant):
    data = sample_data(lib.profile(PROFILE), variant, 'en')
    before = deepcopy(data)
    source = tmp_path / 'engagement.yaml'
    dump_yaml(data, source)
    docx = export.export(source, 'docx', tmp_path / 'report.docx')
    with ZipFile(docx) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    title = root.xpath('//w:p[w:pPr/w:pStyle[@w:val="Title"]]', namespaces=ns)[0]
    assert title.find('w:pPr/w:jc', ns).get(qn('w:val')) == 'left'
    rule = title.find('w:pPr/w:pBdr/w:top', ns)
    assert rule is not None and rule.get(qn('w:val')) == 'single'
    children = [etree.QName(node).localname for node in title.find('w:pPr', ns)]
    assert children.index('pBdr') < children.index('spacing') < children.index('jc')
    html = export.export(source, 'html', tmp_path / 'report.html').read_text()
    assert 'profile-cover layout-editorial' in html
    assert '.cover.layout-editorial,.layout-editorial .cover-foot{text-align:left}' in html
    assert 'Control Review / Fictional sample' in html
    assert data == before


def test_personal_editorial_design_survives_library_deletion_and_project_restore(tmp_path):
    profile = lib.profile(PROFILE)
    library = tmp_path / 'designs.json'
    entry = lib.put_personal(library, {
        'name': 'My controls', 'design': lib.appearance(profile, 'quiet'),
        'sections': profile['sections'],
    })
    project = tmp_path / 'project'; project.mkdir()
    (project / 'img').mkdir()
    data = lib.new_report(entry, 'Control sampling', 'en')
    project_history.save(project, data)
    bundle = project_history.export_bundle(project)
    library.unlink()
    restored = tmp_path / 'restored'
    project_history.import_bundle(bundle, restored)
    import engine
    actual = engine.load_engagement(restored / 'engagement.yaml')
    assert actual['meta']['design'] == data['meta']['design']
    assert actual['meta']['design']['layout'] == 'editorial'
    assert actual['meta']['design']['guidance_id'] == PROFILE
    assert actual['report'] == data['report']


def test_unknown_composition_stays_an_input_error():
    design = lib.appearance(lib.profile(PROFILE))
    design['layout'] = 'editorial-unknown'
    with pytest.raises(ValueError):
        lib.validate_design(design)


@pytest.mark.parametrize('ident', ['offsec-oscp', 'enterprise-assurance',
    'enterprise-purple-team', 'offsec-oswe', 'enterprise-incident', PROFILE])
def test_illustrated_profile_cover_keeps_growth_space_without_shrinking_content(tmp_path, ident):
    from PIL import Image
    import engine
    Image.new('RGB', (720, 400), '#E9EEF3').save(tmp_path / 'illustration.png')
    data = sample_data(lib.profile(ident), 'quiet', 'en')
    title = ('Independent review of identity controls, operational resilience and '
             'evidence custody across shared banking services')
    data['meta']['report_title'] = title
    data['meta']['branding'].update(logo='illustration.png', client_logo='illustration.png',
                                    cover_image='illustration.png')
    output = tmp_path / 'illustrated.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, output)
    document = Document(output)
    cover = document.tables[0]
    # Minimum rows + cell padding must not consume the whole A4 page: long
    # titles and renderer-specific font metrics still need space to grow.
    required = sum(row.height for row in cover.rows)
    for row in cover.rows:
        margins = row.cells[-1]._tc.get_or_add_tcPr().find(qn('w:tcMar'))
        for side in ('top', 'bottom'):
            node = margins.find(qn('w:' + side))
            required += int(node.get(qn('w:w'))) * 635
    assert required <= document.sections[0].page_height - Cm(.7) + Pt(.1)
    paragraph = next(p for row in cover.rows for cell in row.cells
                     for p in cell.paragraphs if p.style.name == 'Title')
    assert paragraph.text == title
    assert all(run.font.size == Pt(27) for run in paragraph.runs if run.text)
    assert len(document.inline_shapes) >= 3
    assert abs(document.inline_shapes[-1].height - Cm(5.2)) < Pt(.1)
