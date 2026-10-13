"""Opt-in literal locale metadata preserves all visible Word source content."""
from copy import deepcopy
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor
from lxml import etree
import pytest

from exporters.technical_text import TECHNICAL_LOCALE, mark_technical_text, technical_spans

TOKEN = 'CASE_0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'
URL = 'https://bank.example.test/' + '0123456789abcdef' * 6 + '?locale=es%2Den&next=%252F#control'


def texts(root):
    return ''.join(root.xpath('.//w:t/text()'))


def marked(root):
    return ''.join(run.xpath('./w:t/text()')[0] if len(run.xpath('./w:t/text()')) == 1
                   else ''.join(run.xpath('./w:t/text()')) for run in root.xpath('.//w:r')
                   if run.xpath('./w:rPr/w:lang[@w:val="'+TECHNICAL_LOCALE+'"]'))


def xml(document):
    return etree.tostring(document.element)


def properties_without_locale(run):
    prop = deepcopy(run.find(qn('w:rPr')))
    if prop is None:
        return None
    lang = prop.find(qn('w:lang'))
    if lang is not None:
        lang.attrib.pop(qn('w:val'), None)
        if not lang.attrib and not len(lang):
            prop.remove(lang)
    return etree.tostring(prop)


@pytest.mark.parametrize('length', [23,24,31,32,64])
@pytest.mark.parametrize('kind', ['technical','alphabetic','url'])
def test_semantic_classifier_thresholds(length, kind):
    token = ('CASE0123456789'*8)[:length] if kind=='technical' else 'A'*length
    if kind=='url':
        token = 'https://a.test/' + 'x'*(length-len('https://a.test/'))
    spans = technical_spans('before '+token+' after')
    expected = length >= (32 if kind=='url' else 24) and kind!='alphabetic'
    assert spans == ([(7,7+length)] if expected else [])


def test_prose_and_unicode_clusters_keep_their_language_and_bytes():
    doc=Document()
    value=('pneumonoultramicroscopicsilicovolcanoconiosis '+ 'A'*80 + ' '
           + 'é'+TOKEN + ' '+TOKEN+'\u0301 '
           + 'https://bank.example.test/cafe\u0301/東京/Δ/'+'0'*64)
    doc.add_paragraph(value)
    before=xml(doc)
    assert mark_technical_text(doc)==0
    assert xml(doc)==before


@pytest.mark.parametrize('prefix', ['Label: ', '🧪 café 東京 Δ: ', '👩\u200d💻 e\u0301: '])
def test_partial_run_offsets_preserve_unicode_styles_and_xml_spaces(prefix):
    doc=Document();p=doc.add_paragraph()
    run=p.add_run(prefix+TOKEN+' tail')
    run.bold=True;run.italic=True;run.font.name='DejaVu Sans Mono';run.font.size=Pt(9)
    run.font.color.rgb=RGBColor(1,2,3)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'es-ES');lang.set(qn('w:eastAsia'),'ja-JP');lang.set(qn('w:bidi'),'ar-SA');run._r.get_or_add_rPr().append(lang)
    before=properties_without_locale(run._r)
    assert mark_technical_text(doc)==1
    assert p.text==prefix+TOKEN+' tail'
    assert marked(p._p)==TOKEN
    assert len(p.runs)==3
    for part in p.runs:
        assert properties_without_locale(part._r)==before
    assert p.runs[0]._r.find(qn('w:t')).get(qn('xml:space'))=='preserve'
    assert p.runs[-1]._r.find(qn('w:t')).get(qn('xml:space'))=='preserve'
    assert p.runs[0]._r.rPr.find(qn('w:lang')).get(qn('w:val'))=='es-ES'
    assert p.runs[-1]._r.rPr.find(qn('w:lang')).get(qn('w:val'))=='es-ES'
    once=xml(doc)
    assert mark_technical_text(doc)==0
    assert xml(doc)==once


def test_token_crosses_bold_runs_and_hyperlink_without_changing_target():
    doc=Document();p=doc.add_paragraph('See ')
    p.add_run(URL[:17]).bold=True
    link=OxmlElement('w:hyperlink')
    rid=doc.part.relate_to(URL, RT.HYPERLINK,is_external=True)
    link.set(qn('r:id'),rid);link.set(qn('w:tooltip'),'Original target')
    r=OxmlElement('w:r');prop=OxmlElement('w:rPr');prop.append(OxmlElement('w:i'));r.append(prop)
    t=OxmlElement('w:t');t.text=URL[17:-8];r.append(t);link.append(r);p._p.append(link)
    p.add_run(URL[-8:]+' after.').font.size=Pt(12)
    before=texts(p._p);relationships=doc.part.rels.xml
    assert mark_technical_text(doc)==3
    assert texts(p._p)==before
    assert marked(p._p)==URL
    assert doc.part.rels.xml==relationships
    assert link.get(qn('r:id'))==rid and link.get(qn('w:tooltip'))=='Original target'
    assert link.xpath('./w:r/w:rPr/w:i')
    assert p.runs[1].bold is True


@pytest.mark.parametrize('tag', ['tab','br','cr','drawing','sym','instrText'])
def test_nontext_run_children_are_real_token_boundaries(tag):
    doc=Document();p=doc.add_paragraph();r=p.add_run()._r
    left=OxmlElement('w:t');left.text='A123456789012';r.append(left)
    barrier=OxmlElement('w:'+tag)
    if tag=='sym':barrier.set(qn('w:char'),'F041')
    if tag=='instrText':barrier.text='DO NOT TOUCH'
    r.append(barrier)
    right=OxmlElement('w:t');right.text='B123456789012';r.append(right)
    before=xml(doc)
    assert mark_technical_text(doc)==0
    assert xml(doc)==before


def test_mixed_text_and_nontext_children_are_preserved_when_run_is_split():
    doc=Document();p=doc.add_paragraph();r=p.add_run()._r
    for tag,value in [('t','Label '+TOKEN),('tab',None),('t','ordinary after'),('br',None),('t',URL)]:
        node=OxmlElement('w:'+tag);node.text=value;r.append(node)
    before=texts(p._p)
    assert mark_technical_text(doc)==1
    assert texts(p._p)==before
    assert marked(p._p)==TOKEN+URL
    assert len(p._p.xpath('.//w:tab'))==len(p._p.xpath('.//w:br'))==1
    assert [node.tag for node in p._p.xpath('.//w:r/*[not(self::w:rPr)]')]==[
        qn('w:t'),qn('w:t'),qn('w:tab'),qn('w:t'),qn('w:br'),qn('w:t')]


def test_fields_are_untouched_and_complex_fields_can_cross_paragraphs():
    doc=Document();p=doc.add_paragraph(TOKEN+' ')
    simple=OxmlElement('w:fldSimple');simple.set(qn('w:instr'),' HYPERLINK "'+URL+'" ')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=URL;r.append(t);simple.append(r);p._p.append(simple)
    original_simple=etree.tostring(simple)
    begin=p.add_run()._r;field=OxmlElement('w:fldChar');field.set(qn('w:fldCharType'),'begin');begin.append(field)
    instr=p.add_run()._r;it=OxmlElement('w:instrText');it.text=' QUOTE "'+TOKEN+'" ';instr.append(it)
    nextp=doc.add_paragraph(TOKEN)
    end=nextp.add_run()._r;field=OxmlElement('w:fldChar');field.set(qn('w:fldCharType'),'end');end.append(field)
    nextp.add_run(' '+TOKEN)
    before_fields=etree.tostring(nextp.runs[0]._r)
    assert mark_technical_text(doc)==2
    assert etree.tostring(simple)==original_simple
    assert etree.tostring(nextp.runs[0]._r)==before_fields
    assert ''.join(doc.element.xpath('.//w:instrText/text()'))==' QUOTE "'+TOKEN+'" '
    assert marked(doc.element)==TOKEN*2


def test_story_parts_tables_existing_headers_and_footers_and_roundtrip():
    doc=Document();doc.add_paragraph(TOKEN)
    doc.add_table(rows=1,cols=1).cell(0,0).text='Table '+TOKEN
    doc.sections[0].header.paragraphs[0].text='Header '+URL
    doc.sections[0].footer.paragraphs[0].text='Footer '+TOKEN
    parts_before={p.partname for p in doc.part.package.parts}
    assert mark_technical_text(doc)==4
    assert {p.partname for p in doc.part.package.parts}==parts_before
    assert marked(doc.sections[0].header._element)==URL
    assert marked(doc.sections[0].footer._element)==TOKEN
    out=BytesIO();doc.save(out);out.seek(0);reopened=Document(out)
    assert mark_technical_text(reopened)==0
    assert marked(reopened.element)==TOKEN*2
    plain=Document();plain.add_paragraph('ordinary prose')
    before={p.partname for p in plain.part.package.parts}
    assert mark_technical_text(plain)==0
    assert {p.partname for p in plain.part.package.parts}==before


def test_content_controls_and_bookmarks_preserve_boundaries_and_metadata():
    doc=Document();p=doc.add_paragraph(TOKEN[:12])
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),'9');start.set(qn('w:name'),'literal');p._p.append(start)
    sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');alias=OxmlElement('w:alias');alias.set(qn('w:val'),'Original literal');pr.append(alias);sdt.append(pr)
    content=OxmlElement('w:sdtContent');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=TOKEN[12:];r.append(t);content.append(r);sdt.append(content);p._p.append(sdt)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'9');p._p.append(end)
    prop_before=etree.tostring(pr)
    assert mark_technical_text(doc)==2
    assert marked(p._p)==TOKEN and texts(p._p)==TOKEN
    assert etree.tostring(pr)==prop_before
    assert p._p.xpath('./w:bookmarkStart/@w:name')==['literal']
    assert p._p.xpath('./w:bookmarkEnd/@w:id')==['9']


def test_many_fragmented_tokens_keep_every_style_and_boundary():
    doc=Document();p=doc.add_paragraph()
    for index in range(120):
        p.add_run('Label '+str(index)+': ')
        for offset in range(0,len(TOKEN),5):
            p.add_run(TOKEN[offset:offset+5]).italic=bool(index%2)
        p.add_run(' / prose. ')
    before=texts(p._p)
    assert mark_technical_text(doc)==120*((len(TOKEN)+4)//5)
    assert texts(p._p)==before and marked(p._p)==TOKEN*120
    once=xml(doc)
    assert mark_technical_text(doc)==0 and xml(doc)==once
