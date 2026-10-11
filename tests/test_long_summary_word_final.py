"""Large summaries flow after their chart and preserve row/header structure."""
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn

import delivery
import export
import workflows
from wizard import dump_yaml


def test_long_summary_can_flow_by_row_without_linking_the_entire_inventory(tmp_path):
    data={'meta':{'report_title':'Large synthetic summary','lang':'en','theme':'corporativo'},
          'report':{'sections':[{'key':'findings_summary'},{'key':'findings'}]},
          'findings':[{'id':f'F{i:03}','mode':'vuln','title':f'Synthetic finding {i:03}',
                       'severity':'info','cvss':0,'description_md':f'BODY-CANARY-{i:03}'}
                      for i in range(1,101)]}
    workflows.normalize(data);before=deepcopy(data)
    dump_yaml(data,tmp_path/'engagement.yaml')
    captured=delivery.capture(tmp_path,tmp_path/'captured')
    output=export.export(tmp_path/'engagement.yaml','docx',tmp_path/'report.docx')
    document=Document(output)
    table=next(t for t in document.tables if t.cell(0,2).text=='Finding')
    assert len(table.rows)==101
    assert table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
    assert all(row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None for row in table.rows)
    assert all(p.paragraph_format.keep_with_next for cell in table.rows[0].cells for p in cell.paragraphs)
    assert not any(p.paragraph_format.keep_with_next for row in table.rows[1:-3]
                   for cell in row.cells for p in cell.paragraphs)
    assert all(p.paragraph_format.keep_with_next for row in table.rows[-3:-1]
               for cell in row.cells for p in cell.paragraphs)
    assert not any(p.paragraph_format.keep_with_next for cell in table.rows[-1].cells for p in cell.paragraphs)
    assert table.rows[-1].cells[2].text=='Synthetic finding 100'
    delivery.inspect_export(output,'docx',captured)
    assert data==before
