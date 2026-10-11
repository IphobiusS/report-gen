"""Synthetic regressions for missing findings and export layout defects."""
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn
from lxml import html
import pdfplumber
import pytest

import engine
import export


def pending_report(lang='es', section_based=True):
    data = {
        'meta': {'lang': lang, 'theme': 'ejecutivo', 'report_title': 'Export regression', 'client': 'Synthetic'},
        'findings': [
            {'id': 'F01', 'mode': 'vuln', 'title': 'Confirmed pending classification',
             'verification': 'confirmed', 'description_md': 'Own account verification.',
             'impact_md': 'Recorded impact.', 'remediation_md': 'Server side correction.'},
            {'id': 'F02', 'mode': 'vuln', 'title': 'Second pending finding', 'severity': '',
             'description_md': 'Recorded configuration.', 'impact_md': 'Conditional impact.',
             'remediation_md': 'Configuration correction.'},
        ],
    }
    if section_based:
        data['report'] = {'sections': [{'key': 'executive_summary', 'summary': 'Synthetic regression.'},
                                       {'key': 'findings_summary'}, {'key': 'findings'}]}
    return data


def markup(data, folder, page_map=None):
    return engine.render_html(engine.build_env(), data, folder, page_map or {}, False,
                              engine.load_lang(data['meta']['lang']))


def test_summary_retains_pending_and_does_not_assign_severity():
    data = pending_report()
    data['findings'] += [{'id': 'F03', 'severity': 'high', 'title': 'Rated'},
                         {'id': 'M01', 'mode': 'machine', 'title': 'Machine'}]
    before = deepcopy(data)
    assert [f['id'] for f in engine.summary_findings(data['findings'])] == ['F03', 'F01', 'F02', 'M01']
    assert engine.unclassified_count(data['findings']) == 2
    assert engine.severity_counts(data['findings']) == {'critical': 0, 'high': 1, 'medium': 0, 'low': 0, 'info': 0}
    text = engine.summary_text(data['findings'], engine.load_lang('es'))
    assert '3 hallazgos' in text and '2 pendientes de valoración' in text and '1 máquina evaluada' in text
    assert data == before


@pytest.mark.parametrize('lang', ['es', 'en'])
@pytest.mark.parametrize('mixed', [False, True])
def test_summary_does_not_claim_every_pending_finding_has_a_rating(lang, mixed):
    data = pending_report(lang)
    if mixed:
        data['findings'].append({'id': 'F03', 'title': 'Rated finding', 'severity': 'high'})
    before = deepcopy(data)
    text = engine.summary_text(data['findings'], engine.load_lang(lang))
    expected = 'including:' if lang == 'en' else 'con el siguiente desglose:'
    assert expected in text
    assert ('2 pending assessment' if lang == 'en' else '2 pendientes de valoración') in text
    total = 3 if mixed else 2
    noun = 'findings' if lang == 'en' else 'hallazgos'
    assert f'{total} {noun}' in text
    assert data == before


@pytest.mark.parametrize('section_based', [False, True])
@pytest.mark.parametrize('lang', ['es', 'en'])
def test_pending_findings_are_visible_in_every_summary(tmp_path, section_based, lang):
    data = pending_report(lang, section_based)
    before = deepcopy(data)
    L = engine.load_lang(lang)
    tree = html.document_fromstring(markup(data, tmp_path))
    rows = tree.xpath('//table[@class="sumtable"]/tbody/tr')
    assert len(rows) == 2
    assert all(L['labels']['severity_pending'] in row.text_content() for row in rows)
    assert not tree.xpath('//svg[contains(@class,"sevchart")]')
    text = export.to_markdown(data, data['meta'], L, tmp_path)
    assert engine.summary_text(data['findings'], L) in text
    assert text.count(L['labels']['severity_pending']) == 2
    out = tmp_path / 'pending.docx'
    export.to_docx(data, data['meta'], L, tmp_path, out)
    doc = Document(out)
    assert engine.summary_text(data['findings'], L) in '\n'.join(p.text for p in doc.paragraphs)
    summary = next(t for t in doc.tables if t.cell(0, 0).text == '#')
    assert len(summary.rows) == 3
    assert all(L['labels']['severity_pending'] in r.cells[1].text for r in summary.rows[1:])
    assert data == before


@pytest.mark.parametrize('findings,total,machines', [([], 0, 0),
    ([{'mode': 'machine', 'id': 'M01', 'title': 'One'}], 0, 1),
    ([{'id': 'F01', 'title': 'Zero', 'severity': 'info', 'cvss': 0}], 1, 0)])
def test_summary_zero_and_machine_counts(findings, total, machines):
    L = engine.load_lang('es')
    text = engine.summary_text(findings, L)
    assert f'identificó {total} ' in text
    assert ('máquina' in text) is bool(machines)
    from exporters.common import summary_sev
    if total:
        assert summary_sev(L, findings[0]) == 'Informativa (0)'


@pytest.mark.parametrize('section_based', [False, True])
def test_word_has_native_toc_and_resolved_bookmarks(tmp_path, section_based):
    data = pending_report(section_based=section_based)
    out = tmp_path / 'navigation.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, out)
    doc = Document(out)
    assert any(p.text == 'Tabla de contenidos' for p in doc.paragraphs)
    toc = doc.element.xpath('//w:sdt[w:sdtPr/w:docPartObj/w:docPartGallery[@w:val="Table of Contents"]]')
    assert len(toc) == 1
    anchors = [a.get(qn('w:anchor')) for a in toc[0].iter(qn('w:hyperlink'))]
    names = {a.get(qn('w:name')) for a in doc.element.xpath('//w:bookmarkStart')}
    assert anchors and len(anchors) == len(set(anchors)) and set(anchors) <= names
    instructions = [n.text or '' for n in doc.element.xpath('//w:instrText')]
    assert sum(i.strip().startswith('TC ') for i in instructions) == len(anchors)
    assert sum(i.strip().startswith('PAGEREF ') for i in instructions) == len(anchors)
    assert any(i.strip().startswith('TOC ') for i in instructions)
    assert doc.settings.element.find(qn('w:updateFields')).get(qn('w:val')) == 'true'
    # A field cache must not pretend the PDF page map is Word pagination.
    for instruction in toc[0].iter(qn('w:instrText')):
        if not (instruction.text or '').strip().startswith('PAGEREF '):
            continue
        for run in instruction.getparent().itersiblings():
            assert not any((t.text or '').strip() for t in run.iter(qn('w:t')))
            if any(f.get(qn('w:fldCharType')) == 'end' for f in run.iter(qn('w:fldChar'))):
                break


def test_word_toc_escapes_titles_in_field_instructions(tmp_path):
    data = pending_report()
    data['report']['sections'][0]['title'] = 'Title "quoted" \\f Z'
    out = tmp_path / 'quoted.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, out)
    doc = Document(out)
    instructions = [n.text or '' for n in doc.element.xpath('//w:instrText')]
    assert all(not n.startswith(' HYPERLINK') for n in instructions)
    # All field begin/separate/end elements remain balanced.
    depth = 0
    for node in doc.element.xpath('//w:fldChar'):
        kind = node.get(qn('w:fldCharType'))
        if kind == 'begin': depth += 1
        elif kind == 'end': depth -= 1
        assert depth >= 0
    assert depth == 0


TABLE = '''| ID | Condición | Resultado |
| --- | --- | --- |
| F01 | Configuración verificada | Una descripción extensa de las verificaciones y del contenido observado en la cuenta propia. |
| F02 | Pendiente de valoración | Otro resultado detallado que necesita más espacio que el identificador de la primera columna. |'''


def test_markdown_table_widths_shared_with_word(tmp_path):
    data = pending_report()
    data['report']['sections'][0]['summary'] = TABLE
    tree = html.document_fromstring(markup(data, tmp_path))
    table = tree.xpath('//table[contains(@class,"md-table")]')[0]
    shares = [float(c.get('style').split(':')[1].removesuffix('%')) for c in table.xpath('./colgroup/col')]
    assert len(shares) == 3 and shares[0] < shares[1] < shares[2]
    assert table.xpath('./tbody/tr[1]/td[1]')[0].get('class') == 'rg-compact-cell'
    out = tmp_path / 'tables.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, out)
    doc = Document(out)
    table = next(t for t in doc.tables if t.cell(0, 0).text == 'ID')
    widths = [c.width for c in table.columns]
    assert widths[0] < widths[1] < widths[2]
    assert all(abs(100*w/sum(widths)-share) < .1 for w, share in zip(widths, shares))
    assert table.cell(1, 0).text == 'F01' and table.cell(2, 0).text == 'F02'


def test_table_layout_handles_empty_ragged_and_long_values():
    from report_layout import table_layout
    assert table_layout([]) == ([], [])
    for rows in [[['A','B'],['short']], [['ID','URL'],['F01','x'*4000]], [['A']]]:
        shares, compact = table_layout(rows)
        assert len(shares) == len(compact) == max(map(len, rows))
        assert abs(sum(shares)-100) < .001 and all(s > 0 for s in shares)


def test_table_labels_have_room_for_their_longest_word():
    from report_layout import table_layout
    shares, _ = table_layout([['Pista', 'Respaldo', 'Límite'],
                              ['Preferencias', 'Explanation '*50, 'Limits '*50],
                              ['Consulta por RUT', 'Description '*40, 'Evidence '*30]])
    assert shares[0] >= (len('Preferencias') * 5.2 + 14) / 4.65


def test_code_label_stays_with_its_block_in_both_formats(tmp_path):
    data = pending_report()
    data['report']['sections'][0]['summary'] = '**example.headers** (946 bytes)\n\n```\n'+'a'*64+'\n```'
    tree = html.document_fromstring(markup(data, tmp_path))
    assert tree.xpath('//pre/preceding-sibling::p[1][@class="rg-code-label"]')
    out = tmp_path/'label.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, out)
    label = next(p for p in Document(out).paragraphs if p.text == 'example.headers (946 bytes)')
    assert label.paragraph_format.keep_with_next is True


def test_short_fenced_and_command_blocks_kept_in_word(tmp_path):
    data = pending_report()
    content = '\n'.join(f'CODELINE{i:02d}' for i in range(12))
    data['findings'][0]['walkthrough'] = [{'text_md': '```bash\n'+content+'\n```'}, {'command': content}]
    out = tmp_path / 'code.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, out)
    blocks = [p for p in Document(out).paragraphs if 'CODELINE00' in p.text]
    assert len(blocks) == 2
    assert all(p.paragraph_format.keep_together for p in blocks)
    assert all(p.text == content for p in blocks)


def _weasy(data, folder):
    from weasyprint import HTML
    from resources import restricted_fetcher
    return HTML(string=markup(data, folder), base_url=str(folder),
                url_fetcher=restricted_fetcher(folder, engine.THEMES)).render()


def test_fenced_code_has_block_style_and_stays_whole_in_pdf(tmp_path):
    data = pending_report()
    data['report']['sections'] = [{'key':'executive_summary', 'summary':
        '\n\n'.join('Paragraph '+str(i)+'. '+('Context sentence. '*8) for i in range(12))+
        '\n\n```bash\n'+ '\n'.join(f'CODELINE{i:02d}' for i in range(14))+'\n```'}]
    document = _weasy(data, tmp_path)
    boxes = [b for page in document.pages for b in page._page_box.descendants()
             if getattr(b, 'element_tag', None) == 'pre' and type(b).__name__ == 'BlockBox']
    assert boxes and all(b.style['background_color'].alpha > 0 for b in boxes)
    assert all(b.style['break_inside'] == 'avoid' for b in boxes)
    out = tmp_path / 'code.pdf'; document.write_pdf(out)
    with pdfplumber.open(out) as pdf:
        pages = [p.extract_text() or '' for p in pdf.pages]
    assert next(i for i,p in enumerate(pages) if 'CODELINE00' in p) == next(i for i,p in enumerate(pages) if 'CODELINE13' in p)


def test_long_fenced_code_can_flow_without_losing_lines(tmp_path):
    data = pending_report()
    lines = [f'LONGBLOCK{i:03d}' for i in range(130)]
    data['report']['sections'] = [{'key':'executive_summary', 'summary':'```\n'+'\n'.join(lines)+'\n```'}]
    document = _weasy(data, tmp_path)
    out = tmp_path/'long.pdf'; document.write_pdf(out)
    with pdfplumber.open(out) as pdf:
        text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
        assert all(text.count(line) == 1 for line in lines)
        for page in pdf.pages:
            assert all(c['bottom'] < page.height and c['top'] >= 0 for c in page.chars)


def test_toc_leaders_align_inside_their_rows(tmp_path):
    data = pending_report()
    document = _weasy(data, tmp_path)
    checked = 0
    for page in document.pages:
        for box in page._page_box.descendants():
            element = getattr(box, 'element', None)
            if element is None or element.tag != 'a' or type(box).__name__ != 'FlexBox':
                continue
            leaders = [c for c in box.children if getattr(c, 'element', None) is not None
                       and c.element.get('class') == 'leader']
            for leader in leaders:
                checked += 1
                midpoint = box.content_box_y() + box.height/2
                assert abs(leader.border_box_y() - midpoint) < 2
    assert checked >= 3
