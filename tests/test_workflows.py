"""Lifecycle regressions: restored bytes, archive attacks, redaction and report parity."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import stat
import zipfile
import pytest
import yaml
from PIL import Image, PngImagePlugin
import app as webapp
import engine
import workflows as wf
import project_history as hist
import finding_library as lib
import validate
from _support import auth_client


def finding(title='Acceso indebido'):
    return {'uid': wf.new_uid(), 'id': 'F1', 'mode': 'vuln', 'title': title, 'severity': 'info',
            'description_md': 'La respuesta expone información.', 'impact_md': 'Un tercero puede leer datos.',
            'remediation_md': 'Verificar la autorización.', 'walkthrough': [{'text_md': 'Solicitar el recurso y comparar la respuesta.'}]}


@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.setattr(webapp, 'PROJECTS', tmp_path)
    monkeypatch.setitem(webapp.app.config, 'LIBRARY_PATH', str(tmp_path / 'global-library.json'))
    c = auth_client(webapp.app)
    assert c.post('/api/projects', json={'slug': 'workflow', 'title': 'Informe de prueba'}).status_code == 200
    return c, tmp_path / 'workflow'


def get(c):
    r = c.get('/api/projects/workflow'); return r.get_json(), r.headers['ETag']


def post(c, path, body):
    _, tag = get(c)
    return c.post('/api/projects/workflow' + path, json=body, headers={'If-Match': tag})


def save(c, data):
    r = c.put('/api/projects/workflow', json=data); assert r.status_code == 200, r.get_json()


def png():
    out = BytesIO(); meta = PngImagePlugin.PngInfo(); meta.add_text('internal', 'SENSITIVE ORIGINAL')
    Image.new('RGB', (120, 80), 'red').save(out, 'PNG', pnginfo=meta); return out.getvalue()


def upload(c):
    r = c.post('/api/projects/workflow/image', data={'file': (BytesIO(png()), 'proof.png')}); assert r.status_code == 200
    source = r.get_json()['src']; response = post(c, '/evidence', {'src': source, 'title': 'Prueba 1'})
    assert response.status_code == 200, response.get_json()
    return response.get_json()['data']['evidence'][0]


def test_legacy_migration_preserves_content_and_identity(project):
    c, d = project
    legacy = {'meta': {'report_title': 'Antiguo', 'lang': 'es'}, 'findings': [dict(finding(), uid=None)]}
    del legacy['findings'][0]['uid']
    (d / 'engagement.yaml').write_text(yaml.safe_dump(legacy))
    data, tag = get(c); ident = data['findings'][0]['uid']; assert wf.UID.fullmatch(ident)
    assert data['findings'][0]['description_md'] == legacy['findings'][0]['description_md']
    assert get(c)[0]['findings'][0]['uid'] == ident
    assert get(c)[1] == tag
    data['findings'].append(dict(finding('Nuevo'), id='F2'))
    data['findings'].reverse()
    for n, f in enumerate(data['findings'], 1): f['id'] = f'F{n}'
    save(c, data)
    assert get(c)[0]['findings'][1]['uid'] == ident


def test_snapshot_restore_includes_images_and_recoverable_current_state(project):
    c, d = project
    e = upload(c); data, _ = get(c); data['findings'] = [dict(finding(), evidence_uids=[e['uid']])]; save(c, data)
    r = post(c, '/history', {'label': 'Entrega v1', 'stage': 'delivered'}); assert r.status_code == 200
    sid = r.get_json()['id']; before = dict(hist.project_files(d))
    data['findings'][0]['title'] = 'Cambió'; save(c, data)
    changed_image = BytesIO(); Image.new('RGB', (20, 20), 'blue').save(changed_image, 'PNG')
    (d / e['src']).write_bytes(changed_image.getvalue())
    response = post(c, '/history/' + sid + '/restore', {}); assert response.status_code == 200, response.get_json()
    assert dict(hist.project_files(d)) == before
    assert any(r['label'].startswith('Restaurada') for r in hist.history(d))
    assert any(hist.snapshot_data(d, r['id'])['findings'][0]['title'] == 'Cambió' for r in hist.history(d) if hist.snapshot_data(d, r['id']).get('findings'))


def test_history_deduplicates_unchanged_files(project):
    c, d = project
    upload(c); before = len(list((d / 'history' / 'objects').glob('*')))
    for n in range(4): post(c, '/history', {'label': f'Versión {n}', 'stage': 'draft'})
    assert len(list((d / 'history' / 'objects').glob('*'))) == before


def test_diff_matches_permanent_ids_and_two_snapshots(project):
    c, _ = project
    data, _ = get(c); data['findings'] = [finding(), dict(finding('Segundo'), id='F2')]; save(c, data)
    first = post(c, '/history', {'label': 'Primera', 'stage': 'draft'}).get_json()['id']
    data['findings'].reverse()
    for i, f in enumerate(data['findings'], 1): f['id'] = f'F{i}'
    save(c, data)
    second = post(c, '/history', {'label': 'Segunda', 'stage': 'retest'}).get_json()['id']
    changes = c.get(f'/api/projects/workflow/history/{first}/compare?to={second}').get_json()
    assert len(changes) == 2
    assert all(r['path'].endswith('.id') and r['finding_uid'] for r in changes)


@pytest.mark.parametrize('path,body', [('/history', {'label': 'A', 'stage': 'draft'}), ('/retest', {}), ('/evidence', {})])
def test_workflow_mutations_reject_stale_revision(project, path, body):
    c, _ = project
    r = c.post('/api/projects/workflow' + path, json=body, headers={'If-Match': 'old'})
    assert r.status_code == 409
    r = c.post('/api/projects/workflow' + path, json=body)
    assert r.status_code == 428


def test_portable_bundle_roundtrip_including_originals_and_versions(project):
    c, d = project
    upload(c); post(c, '/history', {'label': 'Entrega', 'stage': 'delivered'})
    r = c.get('/api/projects/workflow/bundle'); assert r.status_code == 200
    with zipfile.ZipFile(BytesIO(r.data)) as z:
        assert 'manifest.json' in z.namelist()
        assert any(name.startswith('originals/') for name in z.namelist())
        assert any(name.startswith('history/snapshots/') for name in z.namelist())
    result = c.post('/api/projects/import', data={'file': (BytesIO(r.data), 'copy.zip'), 'slug': 'copy'})
    assert result.status_code == 200, result.get_json()
    assert dict(hist.project_files(d)) == dict(hist.project_files(d.parent / 'copy'))
    assert [r['id'] for r in hist.history(d)] == [r['id'] for r in hist.history(d.parent / 'copy')]
    duplicate = c.post('/api/projects/import', data={'file': (BytesIO(r.data), 'copy.zip'), 'slug': 'copy'})
    assert duplicate.status_code == 409


def archive(entries):
    output = BytesIO()
    with zipfile.ZipFile(output, 'w') as z:
        for name, data in entries: z.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize('name', ['../escape', '/absolute', 'img/../../escape', 'img\\escape.png', 'C:/escape', 'webapp/app.py', 'img//a.png'])
def test_zip_paths_rejected_without_partial_import(tmp_path, name):
    with pytest.raises(ValueError): hist.import_bundle(archive([(name, 'bad')]), tmp_path / 'imported')
    assert not (tmp_path / 'imported').exists()
    assert not list(tmp_path.glob('.import-*'))


def test_zip_duplicate_case_and_symlinks_rejected(tmp_path):
    for entries in [[('img/A.png', b'a'), ('img/a.png', b'b')]]:
        with pytest.raises(ValueError): hist.import_bundle(archive(entries), tmp_path / 'bad')
    link = zipfile.ZipInfo('img/link.png'); link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(ValueError): hist.import_bundle(archive([(link, '../../outside')]), tmp_path / 'bad')


def test_bundle_rejects_changed_hash_and_resource_limits(project, monkeypatch):
    c, d = project
    blob = hist.export_bundle(d)
    with zipfile.ZipFile(BytesIO(blob)) as z: entries = [(n, z.read(n)) for n in z.namelist()]
    tampered = [(n, b'bad' if n == 'engagement.yaml' else b) for n, b in entries]
    with pytest.raises(ValueError, match='Integridad'): hist.import_bundle(archive(tampered), d.parent / 'bad')
    monkeypatch.setattr(hist, 'MAX_EXPANDED', 1)
    with pytest.raises(ValueError, match='grande'): hist.import_bundle(blob, d.parent / 'bad')


def test_redaction_changes_pixels_removes_metadata_and_preserves_original(project):
    c, d = project
    e = upload(c); data, _ = get(c)
    data['findings'] = [dict(finding(), evidence_uids=[e['uid']], walkthrough=[{'figure': {'src': e['src'], 'caption': 'Sensitive'}}]),
                        dict(finding('Otro'), id='F2', evidence_uids=[e['uid']], description_md=f'![Proof]({e["src"]})')]
    save(c, data)
    result = post(c, '/evidence/' + e['uid'] + '/edit', {'operations': [dict(type='redact', x=.1, y=.1, w=.5, h=.5)]})
    assert result.status_code == 200, result.get_json()
    changed = result.get_json()['data']; processed = changed['evidence'][0]['src']
    assert processed != e['src']
    with Image.open(d / processed) as im:
        assert im.getpixel((30, 25)) == (0, 0, 0)
        assert im.getpixel((100, 70)) == (255, 0, 0)
        assert im.info == {}
    assert (d / e['original']).read_bytes() == png()
    assert not (d / e['src']).exists()
    assert changed['findings'][0]['walkthrough'][0]['figure']['src'] == processed
    assert e['src'] not in changed['findings'][1]['description_md']
    # Report bundles include only the processed pixels, not originals/history.
    md = c.post('/api/projects/workflow/export/md')
    assert md.status_code == 200, md.get_json()
    with zipfile.ZipFile(BytesIO(md.data)) as z:
        assert set(z.namelist()) == {'report.md', processed}
        assert b'SENSITIVE ORIGINAL' not in z.read(processed)


@pytest.mark.parametrize('operation', [dict(type='redact',x=-1,y=0,w=1,h=1), dict(type='box',x=0,y=0,w=2,h=1), dict(type='blur',x=0,y=0,w=1,h=1), dict(type='redact',x='0',y=0,w=1,h=1), dict(type='redact',x=0,y=0,w=float('nan'),h=1)])
def test_bad_annotation_does_not_change_project(project, operation):
    c, d = project; e = upload(c); before = (d / 'engagement.yaml').read_bytes()
    r = post(c, '/evidence/' + e['uid'] + '/edit', {'operations': [operation]})
    assert r.status_code == 422
    assert (d / 'engagement.yaml').read_bytes() == before


def test_retest_updates_status_and_keeps_separate_records(project):
    c, _ = project; e = upload(c); data, _ = get(c); data['findings'] = [finding()]; save(c, data)
    fid = data['findings'][0]['uid']
    for result, status in [('partial', 'in_progress'), ('fixed', 'verified_fixed'), ('not_verified', 'not_verified')]:
        r = post(c, '/retest', {'finding_uid': fid, 'tested_on': '2026-09-08', 'assessor': 'QA', 'result': result, 'notes_md': 'Verificación controlada', 'evidence_uids': [e['uid']]})
        assert r.status_code == 200, r.get_json(); assert r.get_json()['data']['findings'][0]['status'] == status
    records = get(c)[0]['findings'][0]['retests']; assert len(records) == 3
    assert len({r['uid'] for r in records}) == 3


def test_delivery_review_finds_missing_evidence_steps_resources_and_todo(project):
    c, _ = project; data, _ = get(c); f = finding('TODO completar'); f['walkthrough'] = []; f['remediation_md'] = ''
    data['findings'] = [f, dict(finding('Ausente'), id='F2', walkthrough=[{'figure': {'src': 'img/missing.png'}}])]
    save(c, data)
    issues = c.get('/api/projects/workflow/review').get_json()['issues']; text = ' '.join(x['message'] for x in issues)
    assert all(word in text for word in ['TODO', 'reproducción', 'remediation_md', 'Imagen', 'evidencia'])
    assert all(x['finding_uid'] for x in issues if x['path'].startswith('findings['))


def test_library_sanitizes_links_and_replaces_variables(project):
    c, _ = project
    original = dict(finding('Falla de Client A'), affected='https://client.test', evidence_uids=[wf.new_uid()], owner='Person', retests=[], host={'ip': 'private'},
                    description_md='Client A usa https://client.test ![Secret](img/private.png)', walkthrough=[{'figure': {'src': 'img/x.png'}}])
    cleaned = lib.sanitize(original, {'cliente': 'Client A', 'objetivo': 'https://client.test'})
    assert 'Client A' not in json.dumps(cleaned)
    assert not any(k in cleaned for k in ['uid', 'affected', 'evidence_uids', 'owner', 'retests', 'host'])
    assert 'img/' not in json.dumps(cleaned)
    current = c.get('/api/library'); tag = current.headers['ETag']
    r = c.post('/api/library', json={'name': 'Auth', 'tags': ['web', 'auth'], 'content': cleaned}, headers={'If-Match': tag})
    assert r.status_code == 200, r.get_json(); entry = r.get_json()
    assert entry['variables'] == ['cliente', 'objetivo']
    assert len(entry['versions']) == 1
    bad = c.post('/api/library/' + entry['uid'] + '/instantiate', json={'values': {}})
    assert bad.status_code == 422
    a = c.post('/api/library/' + entry['uid'] + '/instantiate', json={'values': {'cliente': 'Client B', 'objetivo': 'https://new.test'}})
    b = c.post('/api/library/' + entry['uid'] + '/instantiate', json={'values': {'cliente': 'Client B', 'objetivo': 'https://new.test'}})
    assert a.status_code == 200; assert a.get_json()['uid'] != b.get_json()['uid']
    assert a.get_json()['title'] == 'Falla de Client B'
    assert a.get_json()['affected'] == 'https://new.test'
    assert c.post('/api/library', json={'name': 'Stale', 'content': cleaned}, headers={'If-Match': tag}).status_code == 409


def test_library_import_is_atomic_and_preserves_template_versions(project):
    c, _ = project; path = Path(webapp.app.config['LIBRARY_PATH'])
    row = lib.put(path, {'name': 'A', 'content': finding(), 'tags': ['web']})
    lib.put(path, {'name': 'A2', 'content': finding('Edited'), 'tags': ['web']}, row['uid'])
    exported = c.get('/api/library'); body = exported.get_json(); before = path.read_bytes()
    broken = deepcopy(body); broken['templates'].append({'bad': 'shape'})
    response = c.post('/api/library/import', json=broken, headers={'If-Match': exported.headers['ETag']})
    assert response.status_code == 422; assert path.read_bytes() == before
    response = c.post('/api/library/import', json=body, headers={'If-Match': exported.headers['ETag']})
    assert response.status_code == 200
    assert [len(r['versions']) for r in lib.load(path)['templates']] == [2, 2]


@pytest.mark.parametrize('data', [dict(assets='bad'), dict(evidence=[{'src': '../secret.png'}]), dict(findings=[{'uid': 'F1'}]), dict(findings=[dict(finding(), retests='bad')]), dict(report={'workflow': {'include_assets': 'yes'}})])
def test_new_fields_have_structure_validation(data):
    assert validate.structure_issues(data)


def rich_project():
    a = {'uid': wf.new_uid(), 'name': 'Portal', 'target': 'https://portal.test/api', 'kind': 'URL'}
    b = {'uid': wf.new_uid(), 'name': 'Servicio', 'target': '10.10.1.2', 'kind': 'Host'}
    e = {'uid': wf.new_uid(), 'title': 'Evidencia procesada', 'caption': 'Respuesta posterior a la corrección', 'src': 'img/proof.png'}
    f = dict(finding(), asset_uids=[a['uid'], b['uid']], evidence_uids=[e['uid']], status='verified_fixed', owner='Equipo App', due_date='2026-10-01',
             retests=[{'uid': wf.new_uid(), 'result': 'fixed', 'tested_on': '2026-09-08', 'assessor': 'QA', 'notes_md': 'El recurso responde 403 después de corregir.', 'evidence_uids': [e['uid']]}])
    return {'meta': {'lang': 'es', 'theme': 'corporativo', 'report_title': 'Prueba integral', 'client': 'Cliente ficticio'},
            'report': {'sections': [{'key': 'findings'}]}, 'findings': [f, dict(finding('Segundo hallazgo'), id='F2')], 'assets': [a, b], 'evidence': [e],
            'coverage': [{'uid': wf.new_uid(), 'test': 'Control de acceso', 'asset_uid': a['uid'], 'status': 'done', 'notes': 'Verificado con dos cuentas'}],
            'limitations_md': 'No se evaluó el entorno de producción.'}


@pytest.mark.parametrize('section_based', [False, True])
def test_new_content_in_html_docx_and_markdown_without_source_mutation(tmp_path, section_based):
    data = rich_project()
    # Canonical section key is discovered from the existing catalog.
    import sections
    key = next(s['key'] for s in sections.load_catalog()['sections'] if s.get('special') == 'findings')
    data['report']['sections'] = [{'key': key}] if section_based else []
    (tmp_path / 'img').mkdir(); (tmp_path / 'img/proof.png').write_bytes(png())
    original = deepcopy(data); lang = engine.load_lang('es')
    html = engine.render_html(engine.build_env(), data, tmp_path, {}, False, lang)
    from exporters.markdown import to_markdown
    md = to_markdown(data, data['meta'], lang, tmp_path)
    from exporters.docx import to_docx
    from docx import Document
    path = to_docx(data, data['meta'], lang, tmp_path, tmp_path / 'report.docx')
    doc = Document(path)
    doc_text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for text in [html, md.replace('\\', ''), doc_text]:
        assert 'Activos evaluados' in text
        assert '10.10.1.2' in text
        assert 'Control de acceso' in text
        assert 'Equipo App' in text
        assert 'responde 403' in text
        assert 'No se evaluó el entorno' in text
    assert len(doc.inline_shapes) >= 2
    assert data == original
    from lxml import html as html_parser
    tree = html_parser.fromstring(html)
    assert sum('Respuesta posterior a la corrección' in c.text_content() for c in tree.xpath('//figcaption')) == 2


def test_report_configuration_can_omit_workflow_sections():
    data = rich_project(); data['report']['workflow'] = dict(include_assets=False, include_coverage=False, include_retest=False)
    assert wf.report_sections(data) == []


@pytest.mark.parametrize('value', [[], {}, True, 42, None])
def test_enum_shapes_never_crash(value):
    data = {'findings': [dict(finding(), status=value, retests=[{'result': value}])], 'coverage': [{'status': value}]}
    assert validate.structure_issues(data)


def test_imported_backup_rejects_active_content_disguised_as_image(project):
    _, d = project
    files = dict(hist.project_files(d)); files['img/evil.html'] = b'<script src="evil.js"></script>'
    manifest = {'format': 'report-gen-project', 'version': 1, 'files': {k: hist.digest(v) for k, v in files.items()}}
    blob = archive([('manifest.json', json.dumps(manifest))] + list(files.items()))
    with pytest.raises(ValueError, match='Formato de imagen'): hist.import_bundle(blob, d.parent / 'bad')
    files['img/evil.png'] = files.pop('img/evil.html')
    manifest['files'] = {k: hist.digest(v) for k, v in files.items()}
    with pytest.raises(ValueError, match='Imagen inválida'): hist.import_bundle(archive([('manifest.json', json.dumps(manifest))] + list(files.items())), d.parent / 'bad')


def test_restore_rolls_back_if_directory_swap_fails(project, monkeypatch):
    c, d = project
    sid = post(c, '/history', {'label': 'Original', 'stage': 'draft'}).get_json()['id']
    data, _ = get(c); data['meta']['report_title'] = 'Estado actual'; save(c, data)
    before = (d / 'engagement.yaml').read_bytes()
    rename = Path.rename
    def fail_once(path, target):
        if path.name.startswith('.restore-'): raise OSError('Simulated swap failure')
        return rename(path, target)
    monkeypatch.setattr(Path, 'rename', fail_once)
    with pytest.raises(OSError): hist.restore(d, sid)
    assert (d / 'engagement.yaml').read_bytes() == before


def test_original_resources_are_never_allowed_in_reports(tmp_path):
    from resources import local_image
    (tmp_path / 'originals').mkdir(); (tmp_path / 'originals/proof.png').write_bytes(png())
    with pytest.raises(ValueError): local_image('originals/proof.png', tmp_path)


def test_workflow_toc_page_numbers_are_discovered_in_rendered_pdf(tmp_path):
    data = rich_project(); data['report']['sections'] = [{'key': 'findings'}]
    (tmp_path / 'img').mkdir(); (tmp_path / 'img/proof.png').write_bytes(png())
    html = engine.render_html(engine.build_env(), data, tmp_path, {}, True, engine.load_lang('es'))
    target = tmp_path / 'toc.pdf'; engine.render_pdf_weasyprint(html, target, tmp_path)
    keys = engine.toc_keys(data)
    assert {'rg_assets', 'rg_coverage', 'rg_retest'} <= set(keys)
    mapping = engine.read_page_map(target, keys)
    assert all(mapping[k] > 0 for k in ('rg_assets', 'rg_coverage', 'rg_retest'))


def test_annotation_layers_are_flattened_with_redaction_on_top(project):
    c, d = project; e = upload(c)
    operations = [dict(type='redact', x=.1, y=.1, w=.4, h=.4), dict(type='text', x=.1, y=.1, w=.4, h=.4, text='Secret'),
                  dict(type='box', x=.6, y=.1, w=.2, h=.2), dict(type='arrow', x=.6, y=.6, w=.2, h=.2)]
    r = post(c, '/evidence/' + e['uid'] + '/edit', {'operations': operations})
    assert r.status_code == 200, r.get_json()
    with Image.open(d / r.get_json()['data']['evidence'][0]['src']) as image:
        assert image.getpixel((25, 20)) == (0, 0, 0)
        assert not image.info


def test_default_restore_recovers_previous_evidence_bytes(project):
    c, d = project; e = upload(c); old = (d / e['src']).read_bytes()
    result = post(c, '/evidence/' + e['uid'] + '/edit', {'operations': [dict(type='redact',x=0,y=0,w=1,h=1)]})
    assert result.status_code == 200
    restored = post(c, '/restore', {})
    assert restored.status_code == 200
    data, _ = get(c); assert data['evidence'][0]['src'] == e['src']
    assert (d / e['src']).read_bytes() == old


def test_exports_reject_dangling_links_instead_of_silently_omitting_evidence(project):
    c, _ = project; data, _ = get(c)
    data['findings'] = [dict(finding(), evidence_uids=[wf.new_uid()])]; save(c, data)
    for fmt in ['md', 'docx', 'pdf']:
        r = c.post('/api/projects/workflow/export/' + fmt)
        assert r.status_code == 422
