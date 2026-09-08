"""Rendered pagination regressions: flowing findings and indivisible tables."""
import copy
from pathlib import Path

import pdfplumber
import pytest
from docx import Document
from docx.oxml.ns import qn

import engine
import export

ROOT = Path(__file__).resolve().parents[1]


def example(theme="corporativo", sections=True):
    data = {
        "meta": {"lang": "es", "theme": theme, "report_title": "Pagination", "client": "Test"},
        "findings": [
            {"id": f"F{i}", "mode": "vuln", "title": f"PAGINATION{i}", "severity": "info", "cvss": 0,
             "cwe": "CWE-200", "affected": f"test{i}.example.test", "description_md": f"START{i}: Example description.",
             "impact_md": "Example impact.", "remediation_md": f"END{i}: Example remediation."}
            for i in range(1, 4)
        ],
    }
    if sections:
        data["report"] = {"sections": [{"key": "findings"}]}
    return data


def render(data, folder, markers=False):
    source = engine.render_html(engine.build_env(), data, folder, {}, markers, engine.load_lang("es"))
    out = folder / "pagination.pdf"
    engine.render_pdf_weasyprint(source, out, folder)
    with pdfplumber.open(out) as document:
        pages = [page.extract_text() or "" for page in document.pages]
    return out, pages


def page_of(pages, text):
    matches = [i for i, page in enumerate(pages) if text in page]
    assert len(matches) == 1, (text, matches)
    return matches[0]


@pytest.mark.parametrize("theme", ["serio", "corporativo", "offsec", "htb"])
@pytest.mark.parametrize("sections", [True, False])
def test_short_findings_share_page_without_splitting_tables(tmp_path, theme, sections):
    _, pages = render(example(theme, sections), tmp_path)
    # The legacy summary also contains titles; use unique detail-cell text.
    assert page_of(pages, "START1") == page_of(pages, "START2")
    for i in range(1, 4):
        assert page_of(pages, f"START{i}") == page_of(pages, f"END{i}")


def test_main_table_and_its_heading_move_together_with_page_marker(tmp_path):
    data = example()
    data["findings"][0]["walkthrough"] = [{"text_md": "\n\n".join(
        f"Paragraph {i}. " + "Procedure explanation. " * 7 for i in range(14)
    )}]
    out, pages = render(data, tmp_path, markers=True)
    page = page_of(pages, "START2")
    assert page == page_of(pages, "END2") == page_of(pages, "PAGINATION2")
    assert page > page_of(pages, "START1")
    assert engine.read_page_map(out, ["F1", "F2", "F3"])["F2"] == page + 1


def test_markdown_table_moves_whole_to_next_page(tmp_path):
    data = example()
    data["findings"] = data["findings"][:1]
    paragraphs = "\n\n".join(f"Evidence paragraph {i}." for i in range(12))
    table = "| Evidence | Result |\n| --- | --- |\n" + "\n".join(
        f"| MDROW{i:02d} | Result for this evidence |" for i in range(20)
    )
    data["findings"][0]["walkthrough"] = [{"text_md": paragraphs + "\n\nPRETABLEEND\n\n" + table}]
    _, pages = render(data, tmp_path)
    page = page_of(pages, "MDROW00")
    assert page > page_of(pages, "PRETABLEEND")
    assert page == page_of(pages, "MDROW19")


def test_oversized_table_stops_pdf_without_replacing_previous_output(tmp_path):
    data = example()
    data["findings"] = data["findings"][:1]
    data["findings"][0]["description_md"] = "\n\n".join("Long description. " * 15 for _ in range(40))
    out = tmp_path / "pagination.pdf"
    out.write_bytes(b"previous report")
    with pytest.raises(ValueError, match="Hallazgo F1: una tabla no cabe"):
        render(data, tmp_path)
    assert out.read_bytes() == b"previous report"


@pytest.mark.parametrize("sections", [True, False])
def test_word_flows_findings_and_keeps_all_table_rows_together(tmp_path, sections):
    data = example(sections=sections)
    data["findings"][0]["walkthrough"] = [{"text_md": "| A | B |\n| --- | --- |\n| One | Two |\n| Three | Four |"}]
    out = tmp_path / "pagination.docx"
    export.to_docx(copy.deepcopy(data), data["meta"], engine.load_lang("es"), tmp_path, out)
    doc = Document(out)
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.startswith("F1 PAGINATION"))
    end = next(i for i, p in enumerate(paragraphs) if p.text.startswith("F3 PAGINATION"))
    assert not any(p._p.xpath('.//w:br[@w:type="page"]') for p in paragraphs[start:end])
    finding_tables = [t for t in doc.tables if t.cell(0, 0).text in {"CWE", "A"}]
    assert len(finding_tables) == 4
    for table in finding_tables:
        for i, row in enumerate(table.rows):
            assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
            for cell in row.cells:
                assert all(p.paragraph_format.keep_together for p in cell.paragraphs)
                assert cell.paragraphs[-1].paragraph_format.keep_with_next is (i < len(table.rows)-1)


def test_pdf_boundary_preflight_accepts_whole_tables_and_rejects_split_tables(tmp_path):
    from weasyprint import HTML
    from pdf_pagination import mark_finding_tables, verify_table_markers
    # Exercise the backend-neutral PDF verifier independently of the WeasyPrint
    # layout guard. Chromium itself is not required by this test.
    source = '<html><body><div class="finding" id="F9"><table>' + ''.join(
        f'<tr><td>Row {i}</td><td>Evidence</td></tr>' for i in range(4)
    ) + '</table></div></body></html>'
    marked, checks = mark_finding_tables(source)
    assert len(checks) == 1
    verify_table_markers(HTML(string=marked).write_pdf(), checks)
    split = marked.replace('</tr>', '</tr>' + '<tr><td style="height:300mm">Too tall</td><td>Too tall</td></tr>', 1)
    with pytest.raises(ValueError, match='Hallazgo F9'):
        verify_table_markers(HTML(string=split).write_pdf(), checks)


def test_api_returns_actionable_error_for_oversized_table_without_saving(tmp_path, monkeypatch):
    import app as webapp
    from _support import auth_client
    monkeypatch.setattr(webapp, "PROJECTS", tmp_path)
    client = auth_client(webapp.app)
    assert client.post('/api/projects', json={'slug': 'pagination', 'title': 'Original'}).status_code == 200
    original = client.get('/api/projects/pagination').get_json()
    data = example()
    data['findings'][0]['description_md'] = '\n\n'.join('Oversized table. ' * 15 for _ in range(40))
    response = client.post('/api/projects/pagination/export/pdf', json=data)
    assert response.status_code == 422
    assert 'Hallazgo F1' in response.get_json()['error']
    assert 'procedimiento detallado' in response.get_json()['error']
    assert client.get('/api/projects/pagination').get_json() == original
