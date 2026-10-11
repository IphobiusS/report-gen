"""Word tracking columns fit common owners, localized states and ISO dates."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from PIL import ImageFont
import pytest

import engine
import export
import workflows
from report_layout import WORKFLOW_COLUMNS


@pytest.mark.parametrize('lang', ['en', 'es'])
def test_retest_columns_fit_owner_status_and_dates_without_changing_other_exports(tmp_path, lang):
    owner = 'Authentication and Identity' if lang == 'en' else 'Autenticación e Identidad'
    findings = [dict(id=f'F{i:02}', title='Synthetic authorization control', severity='medium',
                     status=status, owner=owner, due_date='2026-11-15',
                     retests=[dict(result='open', tested_on='2026-10-10')])
                for i, status in enumerate(workflows.STATUSES, 1)]
    data = dict(meta=dict(lang=lang, theme='corporativo', report_density='compact'),
                report=dict(sections=[], workflow=dict(include_retest=True)), findings=findings)
    original = deepcopy(data)
    output = tmp_path / f'retest-{lang}.docx'
    export.to_docx(data, data['meta'], engine.load_lang(lang), tmp_path, output)
    document = Document(output)
    table = next(t for t in document.tables if t.cell(0, 0).text == ('Finding' if lang == 'en' else 'Hallazgo')
                 and len(t.columns) == 5)

    # Measure the actual written widths using the same bundled face used by
    # the independent renderer when Calibri is unavailable. A fixed 15%
    # owner column used to split Authentication inside the word.
    font = ImageFont.truetype(str(Path(export.__file__).parent / 'exporters/fonts/DejaVuSans.ttf'), 1050)
    for row in table.rows[1:]:
        for index in (1, 2, 3, 4):
            cell = row.cells[index]
            useful_points = cell.width.pt - 10.8  # standard 108-twip cell margins, both sides
            assert all(font.getlength(word) / 100 <= useful_points for word in cell.text.split())
        assert row.cells[2].text == owner
        assert row.cells[3].text == '2026-11-15'
        assert row.cells[4].text == '2026-10-10'
    assert [row.cells[1].text for row in table.rows[1:]] == [
        workflows.label(workflows.STATUSES, status, lang) for status in workflows.STATUSES]
    assert WORKFLOW_COLUMNS['rg_retest'] == [38, 17, 15, 15, 15]
    assert data == original
