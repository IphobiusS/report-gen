"""Authored retest tables paginate inside native Markdown quotes and lists."""
from copy import deepcopy
from pathlib import Path
import re

from lxml import html
import pytest
from weasyprint import HTML

import engine
import validate
import workflows
from report_layout import prepare_html
from safe_markup import md

THEMES = ('serio', 'corporativo', 'htb', 'offsec', 'ejecutivo', 'laboratorio')


@pytest.mark.parametrize('language', ['es', 'en'])
@pytest.mark.parametrize('tested_on', ['', '2026-10-13'])
def test_retest_heading_omits_separator_for_absent_date(language, tested_on):
    source = {'meta': {'lang': language}, 'findings': [{
        'id': 'F01', 'title': 'Synthetic observation', 'retests': [{
            'tested_on': tested_on, 'result': 'not_verified', 'notes_md': 'No execution.'}]}]}
    before = deepcopy(source)
    section = next(s for s in workflows.report_sections(source) if s['key'] == 'rg_retest')
    heading = html.fragment_fromstring(md(section['body']), create_parent='div').xpath('.//h3')[-1].text_content()
    expected = ['F01'] + ([tested_on] if tested_on else [])
    expected.append(workflows.label(workflows.RESULTS, 'not_verified', language))
    assert heading == ' · '.join(expected)
    assert source == before


def report_source(language, context):
    state = 'Pendiente' if language == 'es' else 'Pending'
    table = '\n'.join(['| Row | State |', '|---|---|'] +
                      [f'| Q{i:02d} | {state} |' for i in range(1, 25)])
    if context == 'quote':
        note = '\n'.join('> ' + line for line in table.splitlines())
    elif context == 'list':
        note = '1. Documentary register\n\n' + '\n'.join('    ' + line for line in table.splitlines())
    else:
        note = '1. Documentary register\n\n' + '\n'.join('    > ' + line for line in table.splitlines())
    findings = [{'id': f'F{i:02d}', 'title': f'Synthetic observation {i:02d}',
                 'severity': 'medium', 'owner': 'Lab analyst', 'status': 'open'}
                for i in range(1, 13)]
    findings[0]['retests'] = [{'tested_on': '2026-10-13', 'result': 'open',
                              'assessor': 'Synthetic analyst', 'notes_md': note}]
    return {'meta': {'lang': language}, 'findings': findings,
            'report': {'workflow': {'include_retest': True}}}


@pytest.mark.parametrize('theme', THEMES)
@pytest.mark.parametrize('language', ['es', 'en'])
@pytest.mark.parametrize('context', ['quote', 'list', 'list-quote'])
def test_nested_retest_table_uses_page_and_keeps_native_rows(theme, language, context):
    source = report_source(language, context)
    original = deepcopy(source)
    assert not any(level == 'error' for level, _ in validate.validate(source, lang=language))
    section = next(s for s in workflows.report_sections(source) if s['key'] == 'rg_retest')
    sheet = (Path(engine.THEMES) / f'{theme}.css').as_uri()
    markup = (f'<html><head><link rel="stylesheet" href="{sheet}"></head><body>'
              '<section class="frontsec sect-rg_retest" id="rg_retest">'
              f'<h2>{section["title"]}</h2>{md(section["body"])}</section></body></html>')
    prepared = prepare_html(markup)
    tree = html.document_fromstring(prepared)
    assert tree.text_content() == html.document_fromstring(markup).text_content()
    group = tree.xpath('//*[@id="rg_retest"]/div')[0]
    assert 'retest-table-group' in group.get('class', '').split()
    assert not group.xpath('./table') and len(group.xpath('.//table')) == 1
    document = HTML(string=prepared).render()
    rows, texts, row_pages, headings = [], [], {}, []
    for number, page in enumerate(document.pages):
        root = page._page_box
        top, bottom = root.content_box_y(), root.content_box_y() + root.height
        boxes = list(root.descendants())
        texts.append(' '.join(b.text for b in boxes if type(b).__name__ == 'TextBox'))
        for box in boxes:
            if type(box).__name__ == 'BlockBox' and box.element.tag == 'h3':
                headings.append((number, box.position_y, box.height))
            if type(box).__name__ != 'TableRowBox':
                continue
            text = ''.join(b.text for b in box.descendants() if type(b).__name__ == 'TextBox')
            match = re.match(r'(Q\d{2})', text)
            if match:
                identifier = match[1]
                rows.append(identifier)
                row_pages[identifier] = number
                assert top - .1 <= box.position_y <= box.position_y + box.height <= bottom + .1
                assert ('Pendiente' if language == 'es' else 'Pending') in text
    assert len(document.pages) == 2
    assert 'Q01' in texts[0] and headings[0][0] == 0
    assert rows == [f'Q{i:02d}' for i in range(1, 25)]
    assert all('Row State' in text for text in texts)
    assert row_pages['Q22'] == row_pages['Q23'] == row_pages['Q24']
    assert source == original


@pytest.mark.parametrize('body', ['<blockquote><p>Retest narrative.</p></blockquote>',
                                 '<ol><li>Retest narrative.</li></ol>'])
def test_nested_narrative_without_table_retains_keep_policy(body):
    source = '<section id="rg_retest"><h3>Retest</h3><p>Assessor: Synthetic</p>' + body + '</section>'
    tree = html.document_fromstring(prepare_html(source))
    group = tree.xpath('//*[@id="rg_retest"]/div')[0]
    assert group.get('class') == 'retest-entry'
    assert tree.text_content() == html.document_fromstring(source).text_content()


@pytest.mark.parametrize('retest', [False, True])
def test_word_generated_retest_counts_keep_first_row_only(tmp_path, retest):
    from docx import Document
    from exporters.richtext import write_markdown
    body = 'Open: 24\n\n| Finding | State |\n|---|---|\n' + '\n'.join(
        f'| F{i:02d} | Open |' for i in range(1, 25))
    body += '\n\nLater authored context.\n\n| Key | Value |\n|---|---|\n| Synthetic | Original |'
    document = Document()
    write_markdown(document, body, tmp_path, lambda _: None, keep_tables=False, retest=retest)
    assert bool(document.paragraphs[0].paragraph_format.keep_with_next) is retest
    later = next(p for p in document.paragraphs if p.text == 'Later authored context.')
    assert not later.paragraph_format.keep_with_next
    rows = document.tables[0].rows
    assert all(p.paragraph_format.keep_with_next for cell in rows[0].cells for p in cell.paragraphs)
    # Initial counts do not make the long register one unbreakable keep chain.
    assert all(not p.paragraph_format.keep_with_next for cell in rows[1].cells for p in cell.paragraphs)
    assert [row.cells[0].text for row in rows[1:]] == [f'F{i:02d}' for i in range(1, 25)]
