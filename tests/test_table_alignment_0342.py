"""Authored table alignment survives publication and editable Word export."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import html
import pytest

from exporters.richtext import write_markdown
from safe_markup import md

TABLE = ('| Item | State | Amount | Default |\n'
         '| :--- | :---: | ---: | --- |\n'
         '| Invoice 01 | **Open** | 12.50 | Original |\n'
         '| Invoice 02 | [Closed](https://evidence.example/status) | 1042.25 | Kept |')

@pytest.mark.parametrize('depth', [0, 1, 2])
def test_table_alignment_and_link_content_survive_docx_roundtrip(tmp_path, depth):
    source = TABLE
    for _ in range(depth):
        source = '- Table evidence\n\n' + '\n'.join('    ' + line for line in source.splitlines())
    tree = html.fragment_fromstring(str(md(source)), create_parent='div')
    expected = ['left', 'center', 'right', None]
    for row in tree.xpath('.//tr'):
        assert [cell.get('align') for cell in row] == expected
        assert [cell.get('style') for cell in row] == ['text-align:left', 'text-align:center', 'text-align:right', None]
    document = Document()
    write_markdown(document, source, tmp_path, lambda value: document.add_paragraph(value))
    output = tmp_path / 'alignment.docx'
    document.save(output)
    reopened = Document(output)
    assert len(reopened.tables) == 1
    for row in reopened.tables[0].rows:
        assert [cell.paragraphs[0].alignment for cell in row.cells] == [
            WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, None]
    assert [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] == [
        ['Item', 'State', 'Amount', 'Default'], ['Invoice 01', 'Open', '12.50', 'Original'],
        ['Invoice 02', 'Closed', '1042.25', 'Kept']]
    assert reopened.tables[0].cell(1, 1).paragraphs[0].runs[0].bold is True
    links = [rel.target_ref for rel in reopened.part.rels.values() if rel.reltype.endswith('/hyperlink')]
    assert links == ['https://evidence.example/status']

@pytest.mark.parametrize('raw', [
    '<table><tr><td align="right" style="position:fixed;top:0">AUTHORED</td></tr></table>',
    '<table><tr><td align="right;display:none">AUTHORED</td></tr></table>',
    '<style>td{display:none}</style>',
])
def test_table_alignment_does_not_enable_authored_html_or_css(raw):
    tree = html.fragment_fromstring(str(md(TABLE + '\n\n' + raw)), create_parent='div')
    assert len(tree.xpath('.//table')) == 1
    assert not tree.xpath('.//style')
    assert raw in tree.text_content()
    assert all(value in {'text-align:left', 'text-align:center', 'text-align:right'}
               for value in tree.xpath('.//@style'))

def test_invalid_alignment_cannot_reach_generated_css(monkeypatch):
    import safe_markup
    # Exercise the defense at the sanitizer boundary even though the shared
    # parser currently generates only the three supported alignment values.
    monkeypatch.setattr(safe_markup, '_parsed', lambda text:
                        '<table><tr><td align="right;display:none" style="display:none">Kept</td></tr></table>')
    tree = html.fragment_fromstring(str(safe_markup.md('unused')), create_parent='div')
    cell = tree.xpath('.//td')[0]
    assert cell.text_content() == 'Kept'
    assert cell.get('align') is None and cell.get('style') is None

@pytest.mark.parametrize('theme', ['serio', 'corporativo', 'htb', 'offsec', 'ejecutivo', 'laboratorio'])
def test_pdf_layout_honors_authored_alignment_over_theme_defaults(theme):
    from pathlib import Path
    from weasyprint import HTML
    import engine
    body = str(md(TABLE)).replace('<table>', '<table class="md-table">')
    sheet = (Path(engine.THEMES) / f'{theme}.css').as_uri()
    document = HTML(string=f'<html><head><link rel="stylesheet" href="{sheet}"></head><body>{body}</body></html>').render()
    checked = 0
    for page in document.pages:
        for cell in page._page_box.descendants():
            if type(cell).__name__ != 'TableCellBox' or cell.element.get('align') is None:
                continue
            alignment = cell.element.get('align')
            assert cell.style['text_align_all'] == alignment
            lines = [box for box in cell.descendants() if type(box).__name__ == 'LineBox']
            assert lines
            for line in lines:
                left = cell.content_box_x()
                right = left + cell.width
                if alignment == 'left':
                    assert abs(line.position_x - left) < .5
                elif alignment == 'right':
                    assert abs(line.position_x + line.width - right) < .5
                else:
                    assert abs(line.position_x + line.width / 2 - (left + right) / 2) < .5
            checked += 1
    assert checked == 9
