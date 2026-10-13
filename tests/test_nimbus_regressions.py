"""Regressions exposed by report-content auditing and nearby export edge cases."""
import copy
import zipfile

import pytest
from PIL import Image
from PIL import ImageDraw
from lxml import html
from docx import Document

import engine
import export
import validate
from safe_markup import md
from wizard import dump_yaml


def project(sections=True):
    data = {'meta': {'lang': 'es', 'theme': 'corporativo', 'report_title': 'Demostración',
                     'assessor': 'Autor', 'assessor_title': 'Rol específico',
                     'osid': 'OS-123', 'email': 'autor@example.test'},
            'findings': [{'id': 'F1', 'title': 'Hallazgo', 'severity': 'info', 'cvss': 0,
                          'description_md': 'Contenido del hallazgo sin modo explícito.',
                          'walkthrough': []}]}
    if sections: data['report'] = {'sections': [{'key': 'findings'}]}
    return data


@pytest.mark.parametrize('sections', [True, False])
def test_markdown_preserves_cover_and_special_evidence_labels(tmp_path, sections):
    path = tmp_path / 'proof [1].png'
    Image.new('RGB', (25, 25), 'white').save(path)
    data = project(sections)
    data['meta']['branding'] = dict.fromkeys(('logo', 'client_logo', 'cover_image'), path.name)
    caption = 'Resultado ] y [canario] <script> & "texto"'
    data['findings'][0]['walkthrough'] = [{'figure': {'src': path.name, 'caption': caption}}]
    source = tmp_path / 'engagement.yaml'; dump_yaml(data, source)
    archive = export.export(source, 'md', tmp_path / 'report.zip')
    with zipfile.ZipFile(archive) as z:
        assert z.read(path.name) == path.read_bytes()
        tree = html.fragment_fromstring(str(md(z.read('report.md').decode())), create_parent='div')
    for text in ('Rol específico', 'OS-123', 'autor@example.test'):
        assert text in tree.text_content()
    assert len(tree.xpath('.//img[not(contains(@src,"report-gen-assets/severity-"))]')) == 4
    assert len(tree.xpath('.//img[contains(@src,"report-gen-assets/severity-")]')) == (0 if sections else 1)
    assert tree.xpath('.//img')[-1].get('alt') == caption
    assert 'Figura 1.' in tree.text_content()
    assert not tree.xpath('.//script')


@pytest.mark.parametrize('sections', [True, False])
def test_command_cannot_close_its_generated_markdown_fence(tmp_path, sections):
    data = project(sections)
    command = 'cat <<EOF\n```\n# must remain code\n````\nEOF'
    data['findings'][0]['walkthrough'] = [{'command': command}]
    result = export.to_markdown(data, data['meta'], engine.load_lang('es'), tmp_path)
    tree = html.fragment_fromstring(str(md(result)), create_parent='div')
    assert len(tree.xpath('.//pre')) == 1
    assert tree.xpath('.//pre')[0].text_content().rstrip('\n') == command
    assert not tree.xpath('.//h1[text()="must remain code"]')


@pytest.mark.parametrize('sections', [True, False])
def test_implicit_vulnerability_mode_matches_explicit_mode(tmp_path, sections):
    Image.new('RGB', (20, 20), 'white').save(tmp_path / 'proof.png')
    data = project(sections)
    data['findings'][0]['walkthrough'] = [{'figure': {'src': 'proof.png', 'caption': 'Evidencia'}}]
    validate.require_valid(data, final=True)
    implicit = copy.deepcopy(data)
    explicit = copy.deepcopy(data); explicit['findings'][0]['mode'] = 'vuln'
    for document in (implicit, explicit): engine.number_figures(document['findings'])
    assert implicit['findings'][0]['walkthrough'][0]['figure']['number'] == 1
    a = engine.render_html(engine.build_env(), implicit, tmp_path, {}, False, engine.load_lang('es'))
    b = engine.render_html(engine.build_env(), explicit, tmp_path, {}, False, engine.load_lang('es'))
    for markup in (a, b):
        assert 'Contenido del hallazgo sin modo explícito.' in markup
        assert 'Modo de finding desconocido' not in markup
    assert html.fromstring(a).xpath('//*[@id="F1"]')[0].text_content() == html.fromstring(b).xpath('//*[@id="F1"]')[0].text_content()


@pytest.mark.parametrize('sections', [True, False])
@pytest.mark.parametrize('lang', ['es', 'en'])
def test_word_localizes_chart_and_keeps_list_label_with_items(tmp_path, monkeypatch, sections, lang):
    data = project(sections)
    data['meta']['lang'] = lang
    if sections:
        data['report']['sections'].insert(0, {'key': 'findings_summary'})
    data['findings'][0]['walkthrough'] = [{'text_md': '**Asset list label**\n\n- portal.example\n- api.example'}]
    data['findings'][0]['remediation_summary'] = {'short_md': 'One.', 'medium_md': 'Two.', 'long_md': 'Three.'}
    drawn = []
    original = ImageDraw.ImageDraw.text

    def record(self, xy, text, *args, **kwargs):
        drawn.append(text)
        if any(ord(c) > 127 for c in text):
            font = kwargs['font']
            assert all(bytes(font.getmask(c)) != bytes(font.getmask('\U0010ffff'))
                       for c in text if ord(c) > 127)
        return original(self, xy, text, *args, **kwargs)

    monkeypatch.setattr(ImageDraw.ImageDraw, 'text', record)
    source = tmp_path / 'engagement.yaml'; dump_yaml(data, source)
    out = export.export(source, 'docx', tmp_path / 'report.docx')
    assert set(engine.load_lang(lang)['severity'].values()) <= set(drawn)
    doc = Document(out)
    label = next(p for p in doc.paragraphs if p.text == 'Asset list label')
    assert label.paragraph_format.keep_with_next is True
    for key in ['short_term', 'medium_term', 'long_term']:
        text = engine.load_lang(lang)['labels'][key] + ':'
        label = next(p for p in doc.paragraphs if p.text.strip() == text)
        assert label.paragraph_format.keep_with_next is True
