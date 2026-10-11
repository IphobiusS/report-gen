"""Word long tables retain a small tail without building oversized chains."""
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm
import pytest

from exporters.pagination import keep_table_together


def table_with_rows(count):
    document = Document()
    table = document.add_table(rows=count, cols=3)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Cm(5)
            cell.text = f'Row {i} / value {j}'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return table


@pytest.mark.parametrize('count', [2, 3, 4, 25, 73])
def test_small_tail_keeps_rows_and_repeated_header_without_chaining_body(count):
    table = table_with_rows(count)
    keep_table_together(table, whole=False)
    for i, row in enumerate(table.rows):
        expected = i == 0 or max(1, count-3) <= i < count-1
        assert all(p.paragraph_format.keep_with_next == expected for cell in row.cells for p in cell.paragraphs)
        assert row._tr.xpath('./w:trPr/w:cantSplit')


@pytest.mark.parametrize('kind', ['long_text', 'fixed_height'])
def test_tall_tail_is_not_forced_into_one_page(kind):
    table = table_with_rows(73)
    for row in table.rows[-3:]:
        if kind == 'fixed_height':
            row.height = Cm(11)
        else:
            row.cells[1].text = 'Recorded control and detailed acceptance criterion. ' * 30
    keep_table_together(table, whole=False)
    assert not any(p.paragraph_format.keep_with_next for row in table.rows[1:]
                   for cell in row.cells for p in cell.paragraphs)
