"""Long evidence captions must not consume the figure's entire page budget."""
from copy import deepcopy
import hashlib
from zipfile import ZipFile

from docx import Document
from docx.shared import Cm
from PIL import Image
import pytest

import engine
import export
from exporters.richtext import write_markdown


MULTILINE = '\n'.join(f'Line {i:02d}: synthetic annotation.' for i in range(1, 61))
FLOWING = ' '.join(f'Annotation {i:03d} documents a synthetic observation without live evidence.'
                   for i in range(1, 121))


def image_file(tmp_path):
    path = tmp_path / 'portrait.png'
    Image.new('RGB', (800, 1600), 'navy').save(path)
    return path


def assert_picture(document, caption, expected_width):
    picture = next(s for s in document.inline_shapes if s._inline.docPr.get('descr') == caption)
    assert picture.width >= expected_width
    assert abs(picture.height - picture.width * 2) <= 2
    section = document.sections[-1]
    assert picture.height <= section.page_height - section.top_margin - section.bottom_margin
    return picture


@pytest.mark.parametrize('theme', ['corporativo', 'ejecutivo', 'laboratorio'])
@pytest.mark.parametrize('caption', [MULTILINE, FLOWING])
def test_structured_long_caption_preserves_legible_image_and_can_continue(tmp_path, theme, caption):
    source = image_file(tmp_path)
    data = {'meta': {'lang': 'en', 'theme': theme},
            'report': {'sections': [{'key': 'findings'}]},
            'findings': [{'id': 'F01', 'title': 'Synthetic observation', 'severity': 'medium',
                          'description_md': 'Documentary fixture.',
                          'walkthrough': [{'figure': {'src': source.name, 'caption': caption}}]}]}
    before = deepcopy(data)
    path = tmp_path / 'report.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, path)
    document = Document(path)
    assert_picture(document, caption, Cm(8))
    caption_paragraph = next(p for p in document.paragraphs if p.text.endswith(caption))
    assert caption_paragraph.paragraph_format.keep_together is False
    figure_paragraph = caption_paragraph._p.getprevious()
    assert figure_paragraph.xpath('./w:pPr/w:keepNext[not(@w:val="0")]')
    with ZipFile(path) as archive:
        media = [hashlib.sha256(archive.read(n)).digest() for n in archive.namelist()
                 if n.startswith('word/media/')]
    assert hashlib.sha256(source.read_bytes()).digest() in media
    assert data == before


@pytest.mark.parametrize('context', ['block', 'inline', 'list', 'table'])
def test_markdown_long_caption_flows_in_each_figure_context(tmp_path, context):
    source = image_file(tmp_path)
    image = f'![{FLOWING}]({source.name})'
    bodies = {'block': image, 'inline': 'Before ' + image + ' After.',
              'list': '1. Before ' + image + '\n\n2. After.',
              'table': '| Description | Evidence |\n|---|---|\n| Synthetic | ' + image + ' |'}
    doc = Document()
    write_markdown(doc, bodies[context], tmp_path, lambda _: None)
    assert_picture(doc, FLOWING, Cm(3) if context == 'table' else Cm(5))
    paragraphs = doc.paragraphs if context != 'table' else doc.tables[0].rows[1].cells[1].paragraphs
    caption = next(p for p in paragraphs if FLOWING in p.text)
    assert caption.paragraph_format.keep_together is False
    assert FLOWING in caption.text


def test_extremely_long_caption_does_not_progressively_reduce_picture(tmp_path):
    image_file(tmp_path)
    sizes = []
    for count in [60, 600]:
        caption = '\n'.join(f'Annotation {i}' for i in range(count))
        data = {'meta': {'lang': 'en'}, 'report': {'sections': [{'key': 'findings'}]},
                'findings': [{'id': 'F01', 'title': 'Synthetic', 'severity': 'info',
                              'walkthrough': [{'figure': {'src': 'portrait.png', 'caption': caption}}]}]}
        path = tmp_path / f'report-{count}.docx'
        export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, path)
        doc = Document(path)
        shape = assert_picture(doc, caption, Cm(8))
        sizes.append((shape.width, shape.height))
        assert any(p.text.endswith(caption) for p in doc.paragraphs)
    assert sizes[0] == sizes[1]
