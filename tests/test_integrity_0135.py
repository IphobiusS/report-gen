"""Regression contracts discovered independently on the delivered 0.13.4 ZIP."""
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys

import pytest
from lxml import html
from PIL import Image
from docx import Document

import app as webapp
import engine
import evidence
import export
import project_history as history
import resources
import workflows
from safe_markup import md


def document():
    return {'meta': {'report_title': 'Audit', 'lang': 'es', 'theme': 'corporativo'},
            'findings': [{'id': 'F1', 'title': 'Proof', 'severity': 'info', 'cvss': 0}]}


@pytest.mark.parametrize('literal', [
    '>     cat img/shot.png', '> >     cat img/shot.png',
    r'\\`img/shot.png`', r'\\\\`img/shot.png`',
    '- One\n\n        cat img/shot.png',
    'The original file img/shot.png was captured yesterday.',
    '![img/shot.png]\n\n[img/shot.png]: img/shot.png',
])
def test_redaction_changes_destinations_only(literal):
    source = '![Real](img/shot.png)\n\n'+literal
    changed = evidence.replace_references({'body': source}, 'img/shot.png', 'img/masked.png')['body']
    expected_literal = literal
    if literal.startswith('![img/shot.png]'):
        expected_literal = literal.replace(']: img/shot.png', ']: img/masked.png')
    assert changed == '![Real](img/masked.png)\n\n'+expected_literal
    before = html.fragment_fromstring(str(md(source,img_base="")), create_parent='div')
    after = html.fragment_fromstring(str(md(changed,img_base="")), create_parent='div')
    assert [n.text_content() for n in before.xpath('.//code')] == [n.text_content() for n in after.xpath('.//code')]
    assert set(after.xpath('.//img/@src')) == {'img/masked.png'}


def test_only_resource_fields_change_when_plain_text_is_a_filename():
    value = {'title':'img/shot.png','caption':'img/shot.png','command':'![x](img/shot.png)',
             'src':'img/shot.png', 'description_md':'![Real](img/shot.png)'}
    result = evidence.replace_references(value, 'img/shot.png', 'img/new.png')
    for key in ('title','caption','command'): assert result[key] == value[key]
    assert result['src'] == 'img/new.png'


@pytest.mark.parametrize('text,missing', [
    ('![Required][proof]\n\n[proof]: img/missing.png', True),
    ('![Required][]\n\n[Required]: img/missing.png', True),
    ('![Required]\n\n[Required]: img/missing.png', True),
    ('`![Example](img/missing.png)`', False),
    ('>     ![Example](img/missing.png)', False),
    ('![Remote](https://invalid.example/shot.png)', True),
])
def test_delivery_review_uses_the_renderers_parser(tmp_path, text, missing):
    data = document(); data['findings'][0]['description_md'] = text
    errors = [row for row in workflows.review(data,tmp_path) if row['level']=='error']
    assert bool(errors) is missing


def test_command_image_syntax_is_not_a_report_resource(tmp_path):
    data = document();data['findings'][0]['walkthrough'] = [{'command':'echo ![x](img/missing.png)'}]
    assert resources.referenced_images(data) == set()
    assert not [x for x in workflows.review(data,tmp_path) if x['level']=='error']


@pytest.mark.parametrize('field,value', [
    ('overview',None), ('overview','text'), ('overview',[]),
    ('confidentiality_md',['text']), ('email',[]), ('date',{}),
])
def test_invalid_export_fields_rejected_before_save(tmp_path,field,value):
    root=tmp_path/'project';data=document();history.save(root,data)
    before=(root/'engagement.yaml').read_bytes();data['meta'][field]=value
    with pytest.raises(ValueError): history.save(root,data)
    assert (root/'engagement.yaml').read_bytes()==before


@pytest.mark.parametrize('sections', [True,False])
@pytest.mark.parametrize('case', ['assessor','caption','branding','version','host'])
def test_optional_and_numeric_display_values_export(tmp_path,sections,case):
    data=document()
    if sections:data['report']={'sections':[{'key':'methodology'},{'key':'findings'}]}
    if case=='assessor':data['meta']['assessor']=None
    elif case=='caption':
        (tmp_path/'img').mkdir();Image.new('RGB',(12,12),'blue').save(tmp_path/'img/shot.png')
        data['findings'][0]['walkthrough']=[{'figure':{'src':'img/shot.png','caption':None}}]
    elif case=='branding':data['meta']['branding']={'wordmark':123,'byline':456,'confidential_text':None}
    elif case=='version':data['meta'].update(version=2,date=2026,osid=987)
    else:data['findings'][0].update(mode='machine',host={'name':123,'ip':123,'os':123})
    history.save(tmp_path,data)
    for fmt in ('md','docx'):
        output=export.export(tmp_path/'engagement.yaml',fmt,tmp_path.parent/f'output.{fmt}')
        assert output.stat().st_size>100
        if fmt=='docx':assert Document(output)


CRASH = r'''
import os, sys
from pathlib import Path
import project_history as h
root=Path(sys.argv[1]);point=sys.argv[2];rename=Path.rename

def crash(self,target):
    if point=='before_move' and self==root: os._exit(91)
    result=rename(self,target)
    if point=='after_move' and self==root: os._exit(91)
    if point=='after_install' and Path(target)==root: os._exit(91)
    return result
Path.rename=crash
h.save(root,{'meta':{'report_title':'NEW'},'findings':[]})
'''


@pytest.mark.parametrize('point,title', [('before_move','OLD'),('after_move','OLD'),('after_install','NEW')])
def test_process_exit_at_commit_recovers_via_api(tmp_path,monkeypatch,point,title):
    root=tmp_path/'audit';history.save(root,{'meta':{'report_title':'OLD'},'findings':[]})
    (root/'img').mkdir();Image.new('RGB',(8,8),'blue').save(root/'img/proof.png')
    image=(root/'img/proof.png').read_bytes()
    result=subprocess.run([sys.executable,'-c',CRASH,str(root),point],cwd=Path(history.__file__).parent)
    assert result.returncode==91
    monkeypatch.setattr(webapp,'PROJECTS',tmp_path)
    client=webapp.app.test_client();listed=client.get('/api/projects').get_json()
    assert len(listed)==1
    assert client.get('/api/projects/audit').get_json()['meta']['report_title']==title
    assert (root/'img/proof.png').read_bytes()==image
    assert not [p for p in tmp_path.iterdir() if p.name.startswith('.')]


def test_repeated_rollback_failure_preserves_recovery_material(tmp_path,monkeypatch):
    root=tmp_path/'audit';history.save(root,document());before=(root/'engagement.yaml').read_bytes()
    original=Path.rename
    def fail(source,target):
        if source.name.startswith(('.transaction-','.previous-')):raise OSError('disk unavailable')
        return original(source,target)
    with monkeypatch.context() as m:
        m.setattr(Path,'rename',fail)
        with pytest.raises(OSError):history.save(root,document())
    assert list(tmp_path.glob('.transaction-*'))
    assert list(tmp_path.glob('.previous-*'))
    assert list(tmp_path.glob('.commit-*.json'))
    history.recover_transactions(tmp_path)
    assert (root/'engagement.yaml').read_bytes()==before
    assert not list(tmp_path.glob('.*'))


def test_recovery_rejects_traversal_without_deleting_files(tmp_path):
    sentinel=tmp_path/'keep';sentinel.mkdir();(sentinel/'proof').write_text('keep')
    parent=tmp_path/'projects';parent.mkdir()
    record={'format':'report-gen-commit-1','target':'../keep','previous':'.previous-a','stage':'.transaction-a','had_target':True}
    (parent/'.commit-transaction-a.json').write_text(json.dumps(record))
    with pytest.raises(ValueError):history.recover_transactions(parent)
    assert (sentinel/'proof').read_text()=='keep'


def test_creation_cleanup_can_retry_after_journal_removal_failure(tmp_path,monkeypatch):
    stage=tmp_path/'.transaction-create';stage.mkdir()
    journal=tmp_path/'.commit-transaction-create.json'
    journal.write_text(json.dumps({'format':'report-gen-commit-1','target':'new','previous':'.previous-new','stage':stage.name,'had_target':False}))
    original=Path.unlink
    def fail(path,*a,**k):
        if path==journal:raise OSError('cleanup unavailable')
        return original(path,*a,**k)
    with monkeypatch.context() as m:
        m.setattr(Path,'unlink',fail);history.recover_transactions(tmp_path)
    assert journal.exists() and not stage.exists()
    history.recover_transactions(tmp_path)
    assert not list(tmp_path.iterdir())


def test_redaction_api_preserves_nested_commands_and_export_pixels(tmp_path,monkeypatch):
    monkeypatch.setattr(webapp,'PROJECTS',tmp_path);client=webapp.app.test_client()
    client.environ_base['HTTP_X_CSRF_TOKEN']=client.get('/api/session').get_json()['csrf']
    assert client.post('/api/projects',json={'slug':'audit'}).status_code==200
    root=tmp_path/'audit';Image.new('RGB',(20,20),'red').save(root/'img/shot.png')
    data=document();uid=workflows.new_uid()
    data['evidence']=[{'uid':uid,'src':'img/shot.png','title':'Proof'}]
    data['findings'][0]['description_md']='![Real](img/shot.png)\n\n>     cat img/shot.png'
    history.save(root,data);etag=client.get('/api/projects/audit').headers['ETag']
    response=client.post('/api/projects/audit/evidence/'+uid+'/edit',json={'operations':[{'type':'redact','x':0,'y':0,'w':1,'h':1}]},headers={'If-Match':etag})
    assert response.status_code==200
    value=engine.load_engagement(root/'engagement.yaml')
    assert value['findings'][0]['description_md'].endswith('>     cat img/shot.png')
    with Image.open(root/value['evidence'][0]['src']) as image:assert image.getextrema()==((0,0),)*3
    assert not (root/'img/shot.png').exists()


@pytest.mark.parametrize('sections',[True,False])
def test_markdown_plain_metadata_and_table_values_are_not_markup(tmp_path, sections):
    from exporters.markdown import to_markdown
    data=document();literal='Access | control with `code` and *stars*'
    data['meta']['report_title']='Literal *asterisks*'
    data['findings'][0]['title']=literal
    data['meta']['contacts']={'client':[{'name':'A | B','title':'T *literal*','email':'a@example.test'}]}
    data['meta']['scope']=[{'target':'a|b', 'description':0}]
    if sections:data['report']={'sections':[{'key':'contacts'},{'key':'scope'},{'key':'findings_summary'},{'key':'findings'}]}
    source=to_markdown(deepcopy(data),data['meta'],engine.load_lang('es'),tmp_path)
    tree=html.fragment_fromstring(str(md(source,img_base="")),create_parent='div')
    assert tree.xpath('.//h1')[0].text_content()=='Literal *asterisks*'
    assert not tree.xpath('.//h1/em')
    rows=[[cell.text_content() for cell in row.xpath('./td')] for row in tree.xpath('.//tbody/tr')]
    assert ['1','Informativa (0)',literal] in rows
    assert ['A | B','T *literal*','a@example.test'] in rows
    assert ['a|b','0'] in rows
    from exporters.docx import to_docx
    target=tmp_path/'plain.docx'
    to_docx(deepcopy(data),data['meta'],engine.load_lang('es'),tmp_path,target)
    word_rows=[[cell.text for cell in row.cells] for table in Document(target).tables for row in table.rows]
    assert ['1','Informativa (0)',literal] in word_rows
    assert ['A | B','T *literal*','a@example.test'] in word_rows
    assert ['a|b','0'] in word_rows


@pytest.mark.parametrize('point,expected',[('before_move','CURRENT'),('after_move','CURRENT'),('after_install','OLD')])
def test_restore_process_interruption_recovers_complete_version(tmp_path,monkeypatch,point,expected):
    root=tmp_path/'audit';data=document();data['meta']['report_title']='OLD';history.save(root,data)
    sid=history.history(root)[0]['id']
    data['meta']['report_title']='CURRENT';history.save(root,data)
    code=CRASH.replace("h.save(root,{'meta':{'report_title':'NEW'},'findings':[]})",'h.restore(root,sys.argv[3])')
    result=subprocess.run([sys.executable,'-c',code,str(root),point,sid],cwd=Path(history.__file__).parent)
    assert result.returncode==91
    monkeypatch.setattr(webapp,'PROJECTS',tmp_path)
    response=webapp.app.test_client().get('/api/projects/audit')
    assert response.status_code==200
    assert response.get_json()['meta']['report_title']==expected
    restored=tmp_path/'imported';history.import_bundle(history.export_bundle(root),restored)
    assert engine.load_engagement(restored/'engagement.yaml')==engine.load_engagement(root/'engagement.yaml')


@pytest.mark.parametrize('point,exists',[('before_install',False),('after_install',True)])
def test_creation_interruption_never_reserves_incomplete_project(tmp_path,monkeypatch,point,exists):
    root=tmp_path/'audit'
    code=CRASH.replace("if point=='before_move' and self==root:","if point=='before_install' and Path(target)==root:")
    result=subprocess.run([sys.executable,'-c',code,str(root),point],cwd=Path(history.__file__).parent)
    assert result.returncode==91
    monkeypatch.setattr(webapp,'PROJECTS',tmp_path)
    response=webapp.app.test_client().get('/api/projects')
    assert bool(response.get_json()) is exists
    assert root.exists() is exists
    assert not list(tmp_path.glob('.*'))
