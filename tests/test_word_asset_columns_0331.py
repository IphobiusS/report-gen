"""Actual Word inventory exports retain readable identifiers and table data."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from PIL import ImageFont
import pytest

import engine
import export
from report_layout import WORKFLOW_COLUMNS


@pytest.mark.parametrize('lang', ['en', 'es'])
@pytest.mark.parametrize('density', ['comfortable', 'compact'])
@pytest.mark.parametrize('theme', ['serio', 'corporativo', 'ejecutivo', 'laboratorio', 'htb', 'offsec'])
def test_word_asset_columns_fit_common_identifiers_and_keep_other_formats(tmp_path, lang, density, theme):
    assets = [
        dict(uid='A01', name='DIRECTORY-05', target='198.51.100.207', kind='server', notes='Synthetic identity directory.'),
        dict(uid='A02', name='RECOVERY-12', target='2001:db8:ffff::207', kind='workstation', notes='Synthetic recovery record.'),
        dict(uid='A03', name='GATEWAY-04', target='api.albor.example', kind='application', notes='Synthetic application boundary.'),
        dict(uid='A04', name='LONG-SYNTHETIC-IDENTIFIER-THAT-MAY-WRAP', target='https://api.albor.example/long/administrative/resource', kind='server', notes='Long cells may wrap but must preserve all text.'),
    ]
    data = dict(meta=dict(lang=lang, theme=theme, report_density=density),
                report=dict(sections=[], workflow=dict(include_assets=True)), assets=assets, findings=[])
    original = deepcopy(data)
    output = tmp_path / f'assets-{theme}-{density}-{lang}.docx'
    export.to_docx(data, data['meta'], engine.load_lang(lang), tmp_path, output)
    document = Document(output)
    table = next(t for t in document.tables if len(t.columns) == 4 and
                 t.cell(0, 0).text == ('Asset' if lang == 'en' else 'Activo'))
    assert table._tbl.xpath('./w:tblPr/w:tblLayout[@w:type="fixed"]')
    assert table._tbl.xpath('./w:tr[1]/w:trPr/w:tblHeader')
    grid = [int(v) for v in table._tbl.xpath('./w:tblGrid/w:gridCol/@w:w')]
    assert len(grid) == 4
    assert all([int(v) for v in row._tr.xpath('./w:tc/w:tcPr/w:tcW/@w:w')] == grid
               for row in table.rows)
    assert len(table.rows) == len(assets) + 1
    assert [[c.text for c in row.cells] for row in table.rows[1:]] == [
        [a[k] for k in ('name', 'target', 'kind', 'notes')] for a in assets]

    # Measure actual OOXML cell widths with the conservative bundled fallback
    # face. The former 14% name share forced DIRECTORY-05 to split its last
    # digit even though the adjacent IPv4 target cell had spare room.
    font = ImageFont.truetype(str(Path(export.__file__).parent / 'exporters/fonts/DejaVuSans.ttf'), 1050)
    for row in table.rows[1:4]:
        for cell in row.cells[:3]:
            assert font.getlength(cell.text) / 100 <= cell.width.pt - 10.8
    section = document.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    for row in table.rows:
        assert abs(sum(c.width for c in row.cells) - available) < 3000
    assert WORKFLOW_COLUMNS['rg_assets'] == [14, 32, 16, 38]
    assert data == original
