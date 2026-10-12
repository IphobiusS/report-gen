"""A report inventory must not claim that unverified records were assessed."""
from copy import deepcopy
import re
import zipfile

from docx import Document
from lxml import html
from pypdf import PdfReader
import pytest

import engine
import export
import workflows
from wizard import dump_yaml


SUMMARY = {
    'en': 'The report contains 2 findings with the following recorded severity classification: 1 high, 1 medium.',
    'es': 'El informe contiene 2 hallazgos con la siguiente clasificación de severidad registrada: 1 alta, 1 media.',
}
INVENTORY = {'en': 'Asset inventory', 'es': 'Inventario de activos'}
FORBIDDEN = ('The assessment identified', 'La evaluación identificó',
             'machine assessed', 'machines assessed', 'máquina evaluada', 'máquinas evaluadas',
             'Assessed assets', 'Activos evaluados',
             'performed a security assessment', 'realizó una evaluación de seguridad',
             'Each finding in this report was verified manually',
             'Cada hallazgo de este informe se verificó manualmente')
APPROACH = {'en': 'Inclusion in this report does not establish',
            'es': 'La inclusión en este informe no acredita'}
AUTHORED_APPROACH = 'AUTHOR_APPROACH_CANARY: Only the explicitly documented activities took place.'


def document(lang):
    data = {
        'meta': {'lang': lang, 'theme': 'corporativo', 'report_title': 'Unverified documentary records'},
        'report': {'sections': [{'key': 'findings_summary'}, {'key': 'findings'}]},
        'assets': [{'name': 'Inventoried only', 'target': '192.0.2.13', 'kind': 'server',
                    'notes': 'No contact or assessment took place.'}],
        'findings': [
            {'id': 'F01', 'mode': 'vuln', 'title': 'First documentary hypothesis', 'severity': 'high',
             'status': 'open', 'verification': 'unverified',
             'description_md': 'This is an authored hypothesis, not an observed vulnerability.'},
            {'id': 'F02', 'mode': 'vuln', 'title': 'Second documentary hypothesis', 'severity': 'medium',
             'status': 'open', 'verification': 'unverified',
             'description_md': 'Recorded classification does not establish technical verification.'},
        ],
    }
    workflows.normalize(data)
    data['coverage'] = [{'uid': '20000000-0000-4000-8000-000000000013',
                         'asset_uid': data['assets'][0]['uid'], 'test': 'Behavioral comparison',
                         'status': 'pending', 'notes': 'No test was executed.'}]
    return data


@pytest.mark.parametrize('lang', ['es', 'en'])
@pytest.mark.parametrize('verification', ['unverified', 'confirmed', 'inconclusive', 'not_reproduced'])
def test_recorded_classification_never_infers_verification_or_changes_state(lang, verification):
    data = document(lang)
    data['findings'][1].update(verification=verification, status='in_progress')
    before = deepcopy(data)
    assert engine.summary_text(data['findings'], engine.load_lang(lang)) == SUMMARY[lang]
    assert data == before
    sections = workflows.report_sections(data)
    assert next(s['title'] for s in sections if s['key'] == 'rg_assets') == INVENTORY[lang]
    verification_text = next(s['body'] for s in sections if s['key'] == 'rg_retest')
    assert workflows.label(workflows.VERIFICATIONS, verification, lang) in verification_text
    assert data == before


@pytest.mark.parametrize('lang', ['es', 'en'])
@pytest.mark.parametrize('machines', [0, 1, 2])
def test_pending_and_machine_records_are_counted_without_claiming_assessment(lang, machines):
    findings = [{'id': 'F01', 'title': 'Pending', 'status': 'open', 'verification': 'unverified'}]
    findings += [{'id': f'M{i}', 'mode': 'machine', 'title': f'Machine {i}',
                  'status': 'open', 'verification': 'unverified'} for i in range(machines)]
    before = deepcopy(findings)
    text = engine.summary_text(findings, engine.load_lang(lang))
    assert text.startswith('The report contains 1 finding' if lang == 'en' else 'El informe contiene 1 hallazgo')
    assert ('1 pending assessment' if lang == 'en' else '1 pendiente de valoración') in text
    if machines:
        expected = (f'{machines} machine record' if lang == 'en' else
                    f'{machines} registro de máquina' if machines == 1 else f'{machines} registros de máquinas')
        assert expected in text
    assert not any(phrase in text for phrase in FORBIDDEN)
    assert findings == before


@pytest.mark.parametrize('lang', ['es', 'en'])
@pytest.mark.parametrize('fmt', ['pdf', 'docx', 'md', 'html'])
@pytest.mark.parametrize('layout', ['sections', 'legacy', 'legacy-authored'])
def test_public_exports_keep_inventory_and_summary_neutral(tmp_path, lang, fmt, layout):
    source = tmp_path / 'source'
    source.mkdir()
    data = document(lang)
    if layout != 'sections':
        data.pop('report')
    if layout == 'legacy-authored':
        data['meta']['overview'] = {'approach_md': AUTHORED_APPROACH}
    dump_yaml(data, source / 'engagement.yaml')
    original = (source / 'engagement.yaml').read_bytes()
    path = export.export(source / 'engagement.yaml', fmt, tmp_path / f'report.{fmt}')
    if fmt == 'pdf':
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)
    elif fmt == 'docx':
        doc = Document(path)
        text = '\n'.join([p.text for p in doc.paragraphs] +
                         [c.text for table in doc.tables for row in table.rows for c in row.cells])
    elif fmt == 'md':
        with zipfile.ZipFile(path) as archive:
            text = archive.read('report.md').decode('utf-8')
    else:
        text = html.document_fromstring(path.read_text(encoding='utf-8')).text_content()
    text = re.sub(r'\s+', ' ', text)
    assert SUMMARY[lang] in text
    assert INVENTORY[lang] in text
    assert 'This is an authored hypothesis, not an observed vulnerability.' in text
    assert ('Unverified' if lang == 'en' else 'Sin verificar') in text
    assert ('Pending' if lang == 'en' else 'Pendiente') in text
    assert not any(phrase in text for phrase in FORBIDDEN)
    if layout == 'legacy':
        assert APPROACH[lang] in text
        assert 'AUTHOR_APPROACH_CANARY' not in text
    elif layout == 'legacy-authored':
        assert AUTHORED_APPROACH in text
        assert APPROACH[lang] not in text
    assert (source / 'engagement.yaml').read_bytes() == original
