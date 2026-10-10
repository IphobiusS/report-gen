"""October audit: resource integrity, export parity and rendered report layout."""
from pathlib import Path
import zipfile

import pytest
from PIL import Image
from docx import Document
from docx.shared import Cm
from lxml import html

import app as webapp
import engine
import evidence
import export
import project_history
import validate
from exporters.docx import _palette
from exporters.richtext import write_markdown
from _support import auth_client
from safe_markup import md
from workflows import md_text

ROOT = Path(__file__).resolve().parents[1]


def test_retest_caption_is_plain_text_with_safe_escaping():
    caption = 'RT01. Evidence [TLS] & <script>alert("x")</script>'
    tree = html.fragment_fromstring(str(md('![' + md_text(caption) + '](img/a.png)')), create_parent='div')
    assert tree.xpath('.//figcaption')[0].text_content() == caption
    assert tree.xpath('.//img')[0].get('alt') == caption
    assert not tree.xpath('.//script')


def data():
    return {'meta': {'lang': 'es', 'theme': 'corporativo', 'report_title': 'Audit'},
            'report': {'sections': [{'key': 'findings'}]},
            'findings': [{'id': 'F1', 'mode': 'vuln', 'title': 'Evidence', 'cvss': 0,
                          'severity': 'info', 'description_md': 'Synthetic example.'}]}


def test_redaction_updates_only_complete_paths():
    original = {'src': 'img/a.png', 'original': 'originals/a.png',
                'body': '![A](img/a.png) ![B](img/a.png.webp) `img/a.png` https://x/img/a.png'}
    result = evidence.replace_references(original, 'img/a.png', 'img/redacted.png')
    assert result['src'] == 'img/redacted.png'
    assert result['original'] == original['original']
    # Censorship changes image destinations, not literal code evidence.
    assert result['body'] == '![A](img/redacted.png) ![B](img/a.png.webp) `img/a.png` https://x/img/a.png'
    assert original['src'] == 'img/a.png'


def test_missing_id_is_a_draft_but_not_a_final_report():
    doc = data(); del doc['findings'][0]['id']
    validate.require_valid(doc)
    with pytest.raises(ValueError, match='sin id'):
        validate.require_valid(doc, final=True)


@pytest.mark.parametrize('version', ['4.0', '2.0', 3.1, {}])
def test_cvss_version_matches_the_vector(version):
    doc = data(); doc['findings'][0].update(cvss=9.8, severity='critical', cvss_version=version,
        cvss_vector='CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H')
    assert any(level == 'error' for level, _ in validate.validate(doc))


def test_failed_markdown_export_preserves_previous_file(tmp_path):
    from wizard import dump_yaml
    doc = data(); doc['findings'][0]['walkthrough'] = [{'figure': {'src': 'img/missing.png'}}]
    source = tmp_path / 'engagement.yaml'; dump_yaml(doc, source)
    previous = tmp_path / 'delivery.zip'; previous.write_bytes(b'previous delivery')
    with pytest.raises(ValueError, match='Imagen'):
        export.export(source, 'md', previous)
    assert previous.read_bytes() == b'previous delivery'
    assert not list(tmp_path.glob('.export-*'))


def test_export_snapshots_images_and_yaml_together(tmp_path, monkeypatch):
    monkeypatch.setattr(webapp, 'PROJECTS', tmp_path)
    project = tmp_path / 'snapshot'; (project / 'img').mkdir(parents=True)
    image = project / 'img/e.png'; Image.new('RGB', (40, 20), 'red').save(image)
    original = image.read_bytes()
    doc = data(); doc['findings'][0]['walkthrough'] = [{'figure': {'src': 'img/e.png'}}]
    project_history.save(project, doc)
    client = auth_client(webapp.app)
    real_export = webapp.export.export
    def changed_during_render(source, fmt, target):
        assert source.parent != project
        # Models a concurrent redaction/save after the export snapshot is captured.
        image.unlink()
        assert (source.parent / 'img/e.png').read_bytes() == original
        return real_export(source, fmt, target)
    monkeypatch.setattr(webapp.export, 'export', changed_during_render)
    response = client.post('/api/projects/snapshot/export/md', json=doc)
    assert response.status_code == 200, response.get_json()
    from io import BytesIO
    with zipfile.ZipFile(BytesIO(response.data)) as archive:
        assert archive.read('img/e.png') == original
        assert not any('originals/' in n or 'history/' in n for n in archive.namelist())
    assert not list(project.glob('.render-*'))


@pytest.mark.parametrize('sections', [True, False])
def test_word_supports_webp_and_cover_resources(tmp_path, sections):
    image = tmp_path / 'proof.webp'; Image.new('RGB', (60, 40), 'white').save(image)
    doc = data()
    if not sections: doc.pop('report')
    doc['meta'].update(osid='DEMO-123', email='demo@example.test', branding={
        'logo': 'proof.webp', 'client_logo': 'proof.webp', 'cover_image': 'proof.webp'})
    doc['findings'][0]['walkthrough'] = [{'text_md': '![WebP](proof.webp)',
                                       'figure': {'src': 'proof.webp', 'caption': 'WebP figure'}}]
    out = tmp_path / 'webp.docx'
    export.to_docx(doc, doc['meta'], engine.load_lang('es'), tmp_path, out)
    result = Document(out)
    assert len(result.inline_shapes) >= 5
    # Metadata remains on the cover after separating it into three layout zones.
    cover_text = '\n'.join(result.tables[0]._tbl.xpath('.//w:t/text()'))
    assert 'DEMO-123' in cover_text and 'demo@example.test' in cover_text


def test_word_portrait_figure_fits_page_and_keeps_caption(tmp_path):
    Image.new('RGB', (100, 1600), 'white').save(tmp_path / 'portrait.png')
    doc = data(); doc['findings'][0]['walkthrough'] = [{'figure': {'src': 'portrait.png', 'caption': 'Portrait'}}]
    out = tmp_path / 'portrait.docx'
    export.to_docx(doc, doc['meta'], engine.load_lang('es'), tmp_path, out)
    result = Document(out); shape = result.inline_shapes[-1]; sec = result.sections[-1]
    assert shape.height <= sec.page_height - sec.top_margin - sec.bottom_margin - Cm(1.9)
    figure = next(p for p in result.paragraphs if p._p.xpath('.//wp:inline'))
    assert figure.paragraph_format.keep_with_next is True


def test_workflow_word_tables_can_span_pages():
    doc = Document()
    write_markdown(doc, '| A | B |\n| --- | --- |\n| first | row |\n| second | row |',
                   ROOT, lambda _: None, keep_tables=False)
    table = doc.tables[0]
    assert not table.cell(1, 0).paragraphs[0].paragraph_format.keep_with_next
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)


def luminance(color):
    values = [int(color[i:i+2], 16)/255 for i in (0, 2, 4)]
    return sum((v/12.92 if v <= .04045 else ((v+.055)/1.055)**2.4)*w
               for v, w in zip(values, [.2126, .7152, .0722]))


def contrast(a, b):
    light, dark = sorted([luminance(a), luminance(b)], reverse=True)
    return (light+.05)/(dark+.05)


@pytest.mark.parametrize('theme', ['corporativo', 'serio', 'offsec', 'htb', 'ejecutivo', 'laboratorio'])
def test_severity_text_contrast_in_all_themes(theme):
    palette = _palette({'theme': theme})
    for severity in engine.SEVERITY_ORDER:
        assert contrast(palette('sev-chip-ink', 'ffffff'), palette('sev-'+severity, 'ffffff')) >= 4.5


@pytest.mark.parametrize('theme', ['corporativo', 'serio', 'offsec', 'htb', 'ejecutivo', 'laboratorio'])
def test_rendered_machine_chips_have_legible_contrast(theme):
    from weasyprint import CSS, HTML
    rendered = HTML(string='<span class="host-chip">192.0.2.10</span>'
                           '<span class="sev-chip sev-machine">Máquina</span>').render(
        stylesheets=[CSS(filename=str(engine.THEMES / f'{theme}.css'))])
    chips = [box for box in rendered.pages[0]._page_box.descendants()
             if type(box).__name__ == 'InlineBox' and box.element_tag == 'span']
    assert len(chips) == 2
    for box in chips:
        colors = [box.style[name] for name in ('color', 'background_color')]
        assert all(c.space == 'srgb' and c.alpha == 1 for c in colors)
        hexes = [''.join(f'{round(v * 255):02x}' for v in c.coordinates) for c in colors]
        assert contrast(*hexes) >= 4.5, (theme, box.element.get('class'), hexes)


def test_word_dark_theme_uses_legible_body_colors(tmp_path):
    palette = _palette({'theme': 'htb'}, for_word=True)
    for token in ('ink', 'ink-soft', 'muted', 'accent-ink'):
        assert contrast(palette(token, 'ffffff'), 'ffffff') >= 4.5
    assert palette('cover-bg', 'ffffff') != 'ffffff'
    doc = data(); doc['meta']['theme'] = 'htb'
    out = tmp_path / 'contrast.docx'
    export.to_docx(doc, doc['meta'], engine.load_lang('es'), tmp_path, out)
    rendered = Document(out)
    for style in ('Caption', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
        assert contrast(str(rendered.styles[style].font.color.rgb), 'ffffff') >= 4.5


@pytest.mark.parametrize('sections', [True, False])
def test_css_headers_quote_text_without_html_entities(tmp_path, sections):
    doc = data()
    if not sections: doc.pop('report')
    doc['meta'].update(report_title='A & B "Quoted"', branding={'wordmark': 'x</style><script>test</script>'})
    source = engine.render_html(engine.build_env(), doc, tmp_path, {}, False, engine.load_lang('es'))
    tree = html.fromstring(source)
    assert not tree.xpath('//script')
    css = tree.xpath('//style')[0].text
    assert 'A & B' in css and '&amp;' not in css and '&#34;' not in css
    assert r'\22 ' in css and r'\3c ' in css


@pytest.mark.parametrize('sections', [True, False])
def test_workflow_html_has_table_styles_and_retest_groups(tmp_path, sections):
    doc = data()
    if not sections: doc.pop('report')
    doc['findings'][0]['retests'] = [dict(uid='00000000-0000-0000-0000-000000000001',
        tested_on='2026-10-02', assessor='Demo', result='open', notes_md='Observed', evidence_uids=[])]
    source = engine.render_html(engine.build_env(), doc, tmp_path, {}, False, engine.load_lang('es'))
    tree = html.fromstring(source)
    assert tree.xpath('//*[@id="rg_retest"]/table[contains(@class,"md-table")]')
    assert tree.xpath('//*[@id="rg_retest"]/div[@class="retest-entry"]/h3')


def test_markdown_table_has_real_padding_and_repeats_header(tmp_path):
    from weasyprint import HTML
    doc = data(); doc['report']['sections'].insert(0, dict(key='custom', title='Long table', body=
        '| Name | Result |\n| --- | --- |\n'+'\n'.join(f'| ROW{i:03} | Example result |' for i in range(100))))
    source = engine.render_html(engine.build_env(), doc, tmp_path, {}, False, engine.load_lang('es'))
    rendered = HTML(string=source, base_url=str(tmp_path)).render()
    cells=[]; header_pages=set()
    for number, page in enumerate(rendered.pages):
        for box in page._page_box.descendants():
            if type(box).__name__ == 'TableCellBox' and box.element.tag in ('td', 'th'):
                if box.element.text in ('Name', 'Result'):
                    cells.append(box); header_pages.add(number)
    assert len(header_pages) > 1
    assert cells and all(box.padding_left >= 7 for box in cells)


def test_theme_default_is_not_overridden_by_a_green_fallback(tmp_path):
    doc = data()
    source = engine.render_html(engine.build_env(), doc, tmp_path, {}, False, engine.load_lang('es'))
    assert '--accent: #39ff14' not in source
