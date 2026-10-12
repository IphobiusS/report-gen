"""A full-page final table must not strand its spacer before a new chapter."""
import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

import engine
import export
from design_library import profile
from design_samples import sample_data


@pytest.mark.parametrize('ident', ['enterprise-assurance', 'centri-btl2', 'offsec-oscp'])
@pytest.mark.parametrize('tail', ['', '\n\nA final note.', '\n\n---'])
def test_table_to_chapter_has_no_standalone_empty_spacer(tmp_path, ident, tail):
    data = sample_data(profile(ident), 'quiet', 'en')
    data['report']['sections'] = [
        {'key': 'one', 'title': 'Table chapter',
         'body': '| Marker | Detail |\n|---|---|\n| A | Preserved |' + tail},
        {'key': 'two', 'title': 'Following chapter', 'body': 'The next chapter.'},
    ]
    output = tmp_path / 'boundary.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, output)
    document = Document(output)
    heading = next(p for p in document.paragraphs
                   if p.style.name == 'Heading 1' and 'Following chapter' in p.text)
    assert heading.paragraph_format.page_break_before is True
    previous = heading._p.getprevious()
    if not tail:
        assert previous.tag == qn('w:tbl')
    else:
        # Keep the authored prose and explicit horizontal-rule placeholder.
        assert previous.tag == qn('w:p')
        assert ''.join(previous.itertext())
        assert previous.xpath('.//w:t')[0].text == ('A final note.' if 'note' in tail else ' ')
    assert ['A', 'Preserved'] in [[cell.text for cell in row.cells]
                                 for table in document.tables for row in table.rows]


@pytest.mark.parametrize('protected', ['properties', 'bookmark'])
def test_table_boundary_preserves_nonempty_xml_paragraph(tmp_path, monkeypatch, protected):
    from exporters import richtext
    original = richtext.write_markdown

    def write_with_protected_boundary(document, text, *args, **kwargs):
        original(document, text, *args, **kwargs)
        if str(text).endswith('| A | Preserved |'):
            paragraph = document.paragraphs[-1]
            if protected == 'properties':
                paragraph.paragraph_format.space_after = Pt(7)
            else:
                bookmark = OxmlElement('w:bookmarkStart')
                bookmark.set(qn('w:id'), '9000')
                bookmark.set(qn('w:name'), 'preserved_boundary')
                paragraph._p.append(bookmark)
                end = OxmlElement('w:bookmarkEnd')
                end.set(qn('w:id'), '9000')
                paragraph._p.append(end)

    monkeypatch.setattr(richtext, 'write_markdown', write_with_protected_boundary)
    data = sample_data(profile('enterprise-assurance'), 'quiet', 'en')
    data['report']['sections'] = [
        {'key': 'one', 'title': 'Table chapter',
         'body': '| Marker | Detail |\n|---|---|\n| A | Preserved |'},
        {'key': 'two', 'title': 'Following chapter', 'body': 'Next.'},
    ]
    output = tmp_path / 'protected.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, output)
    document = Document(output)
    heading = next(p for p in document.paragraphs
                   if p.style.name == 'Heading 1' and 'Following chapter' in p.text)
    previous = heading._p.getprevious()
    assert previous.tag == qn('w:p')
    assert previous.find(qn('w:pPr' if protected == 'properties' else 'w:bookmarkStart')) is not None
