"""Portrait evidence must leave room for its full, editable caption."""
from copy import deepcopy

from docx import Document
from docx.shared import Cm
from PIL import Image
import pytest

import engine
import export
from exporters.richtext import write_markdown


CAPTION = (
    'SYNTHETIC portrait evidence. This image preserves the full request chronology, '
    'comparison identity and negative-control result. The first rows establish the '
    'intended tenant and service, the middle rows describe the modeled policy decision, '
    'and the final rows preserve the rejection expected for an unrelated tenant. Every '
    'identifier belongs to the generated fixture. The evidence does not represent a '
    'live capture, does not contain an operational credential and does not prove that '
    'a production bank has been tested or corrected. Keep this context attached to the '
    'complete portrait when exporting or reviewing the report.'
)


@pytest.mark.parametrize('theme', ['corporativo', 'ejecutivo', 'laboratorio'])
@pytest.mark.parametrize('context', ['walkthrough', 'section', 'inline'])
def test_portrait_reserves_caption_and_preceding_labels(tmp_path, theme, context):
    Image.new('RGB', (800, 1600), 'white').save(tmp_path / 'portrait.png')
    data = {'meta': {'lang': 'en', 'theme': theme, 'report_density': 'compact'},
            'report': {'sections': [{'key': 'findings'}]},
            'findings': [{'id': 'F01', 'title': 'Synthetic portrait evidence',
                          'severity': 'medium', 'description_md': 'Observed control.'}]}
    if context == 'walkthrough':
        data['findings'][0]['walkthrough'] = [
            {'lead': 'Compare the complete chronology',
             'figure': {'src': 'portrait.png', 'caption': CAPTION}}]
    else:
        body = f'![{CAPTION}](portrait.png)'
        if context == 'inline':
            body = 'Preceding inline context ' + body + ' Following context.'
        data['report']['sections'].append({'key': 'notes', 'title': 'Evidence', 'body': body})
    original = deepcopy(data)
    path = tmp_path / 'portrait.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, path)
    document = Document(path)
    shape = document.inline_shapes[-1]
    page = document.sections[-1]
    available = page.page_height - page.top_margin - page.bottom_margin

    # A roughly hundred-word caption plus labels needs materially more than
    # the old fixed 2 cm. Keep the portrait readable and its aspect ratio.
    assert Cm(18) <= shape.height <= available - Cm(4)
    assert abs(shape.height - shape.width * 2) <= 2
    paragraph = next(p for p in document.paragraphs if p._p.xpath('.//wp:inline'))
    assert paragraph.paragraph_format.keep_with_next is True
    caption = next(p for p in document.paragraphs if CAPTION in p.text)
    assert caption.paragraph_format.keep_together is True
    assert data == original


def test_horizontal_image_does_not_shrink_when_caption_fits(tmp_path):
    Image.new('RGB', (1600, 800), 'white').save(tmp_path / 'horizontal.png')
    data = {'meta': {'lang': 'es', 'theme': 'corporativo'},
            'report': {'sections': [{'key': 'notes', 'title': 'Evidencia',
                                    'body': f'![{CAPTION}](horizontal.png)'}]}, 'findings': []}
    target = tmp_path / 'horizontal.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, target)
    shape = Document(target).inline_shapes[-1]
    assert shape.width == Cm(15)
    assert abs(shape.width - shape.height * 2) <= 2


@pytest.mark.parametrize('in_list', [False, True])
def test_inline_markdown_figure_retains_prose_as_separate_blocks(tmp_path, in_list):
    Image.new('RGB', (800, 1600), 'white').save(tmp_path / 'portrait.png')
    document = Document()
    markdown = ('**Before** [control](https://example.test/control) '
                '![Retained caption](portrait.png) After the figure.')
    if in_list:
        markdown = '1. ' + markdown + '\n\n2. Next control.'
    write_markdown(document, markdown, tmp_path, lambda _: None)
    paragraphs = document.paragraphs
    picture_index = next(i for i, p in enumerate(paragraphs) if p._p.xpath('.//wp:inline'))
    assert picture_index == 1
    assert paragraphs[0].text == 'Before control '
    assert any(run.bold and run.text == 'Before' for run in paragraphs[0].runs)
    assert paragraphs[0]._p.xpath('.//w:hyperlink')
    assert paragraphs[picture_index].text == ''
    assert paragraphs[picture_index + 1].text == 'Retained caption'
    assert paragraphs[picture_index + 2].text == ' After the figure.'
    if in_list:
        assert paragraphs[0]._p.xpath('./w:pPr/w:numPr')
        assert paragraphs[-1].text == 'Next control.'
