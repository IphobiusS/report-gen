"""Wide ordinals keep their value and available content width in Word."""
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from PIL import Image
import pytest

from exporters.richtext import write_markdown


def _render(tmp_path, body):
    document = Document()
    document.styles['Normal'].font.size = Pt(11)
    def code(text):
        paragraph = document.add_paragraph(text)
        return paragraph
    write_markdown(document, body, tmp_path, code)
    return document


def _numbers(document):
    return [p for p in document.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]


@pytest.mark.parametrize('start,count,sequence', [(0, 2, False), (42, 2, False),
    (65534, 4, False), (65535, 3, False), (65536, 3, True),
    (1234567890, 3, True), (2147483647, 1, True)])
def test_only_unsupported_native_starts_use_cached_recalculable_sequences(tmp_path, start, count, sequence):
    document = _render(tmp_path, '\n'.join(f'{start + i}. Item {i} stays exact.' for i in range(count)))
    paragraphs = _numbers(document)
    assert len(paragraphs) == count
    ids = {p._p.pPr.numPr.numId.val for p in paragraphs}
    assert len(ids) == 1
    num_id = ids.pop()
    root = document.part.numbering_part.element
    numbers = [n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == str(num_id)]
    abstract_id = numbers[0].find(qn('w:abstractNumId')).get(qn('w:val'))
    definition = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == abstract_id)
    level = definition.find(qn('w:lvl'))
    assert level.find(qn('w:numFmt')).get(qn('w:val')) == ('none' if sequence else 'decimal')
    assert level.find(qn('w:start')).get(qn('w:val')) == str(1 if sequence else start)
    expected_tags = ['start', 'numFmt', 'suff', 'lvlText'] if sequence else ['start', 'numFmt', 'lvlText']
    assert [child.tag for child in level] == [qn('w:' + tag) for tag in expected_tags]
    for index, paragraph in enumerate(paragraphs):
        fields = paragraph._p.findall(qn('w:fldSimple'))
        assert len(fields) == int(sequence)
        if sequence:
            switch = f'\\r {start}' if index == 0 else '\\n'
            assert fields[0].get(qn('w:instr')) == f' SEQ RGList{num_id} {switch} \\* ARABIC '
            assert [t.text for t in fields[0].iter(qn('w:t'))] == [str(start + index)]
            assert fields[0].get(qn('w:fldLock')) is None
            following = fields[0].getnext()
            assert following.xpath('./w:t/text()') == ['.']
            assert len(following.findall(qn('w:tab'))) == 1
        else:
            assert paragraph.text == f'Item {index} stays exact.'


def test_independent_and_nested_sequences_do_not_share_counters(tmp_path):
    document = _render(tmp_path, '65536. Outer first\n\n    1234567890. Inner first\n    1. Inner second\n\n65537. Outer second\n\nParagraph boundary.\n\n65536. Independent first')
    fields = [p._p.find(qn('w:fldSimple')) for p in _numbers(document)]
    assert [next(f.iter(qn('w:t'))).text for f in fields] == ['65536', '1234567890', '1234567891', '65537', '65536']
    instructions = [f.get(qn('w:instr')).split() for f in fields]
    assert instructions[0][1] == instructions[3][1]
    assert instructions[1][1] == instructions[2][1]
    assert len({parts[1] for parts in instructions}) == 3


@pytest.mark.parametrize('size', [9, 11, 14])
def test_wide_labels_reserve_room_at_body_size_without_font_shrinking(tmp_path, size):
    document = Document(); document.styles['Normal'].font.size = Pt(size)
    write_markdown(document, '1234567890. Long narrative body.', tmp_path, lambda _: None)
    paragraph = _numbers(document)[0]
    assert abs(paragraph.paragraph_format.left_indent + paragraph.paragraph_format.first_line_indent - Cm(.25)) < Pt(.05)
    assert -paragraph.paragraph_format.first_line_indent > Pt(size * 7)
    assert paragraph.paragraph_format.tab_stops[0].position == paragraph.paragraph_format.left_indent
    assert not paragraph._p.xpath('.//w:rPr/w:sz')


def test_nested_blocks_inherit_the_extra_space_reserved_for_wide_marker(tmp_path):
    body = ('1234567890. Parent\n\n    Continuation paragraph.\n\n'
            '    42. Nested ordinary item.\n\n'
            '    | Evidence | Result |\n    |---|---|\n    | row | retained |\n\n'
            '        Code evidence.\n\n1234567891. Next parent')
    document = _render(tmp_path, body)
    first, nested, final = _numbers(document)
    continuation = next(p for p in document.paragraphs if p.text == 'Continuation paragraph.')
    code = next(p for p in document.paragraphs if p.text == 'Code evidence.')
    assert continuation.paragraph_format.left_indent == first.paragraph_format.left_indent
    assert code.paragraph_format.left_indent == first.paragraph_format.left_indent
    assert nested.paragraph_format.left_indent > first.paragraph_format.left_indent
    assert nested.paragraph_format.left_indent + nested.paragraph_format.first_line_indent > first.paragraph_format.left_indent
    assert final.paragraph_format.left_indent == first.paragraph_format.left_indent
    table = document.tables[0]
    offset = int(table._tbl.tblPr.find(qn('w:tblInd')).get(qn('w:w'))) / 20
    assert abs(offset - first.paragraph_format.left_indent.pt) < .1
    section = document.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    assert abs(sum(c.width for c in table.columns) + first.paragraph_format.left_indent - available) < Pt(.1)


def test_figure_first_keeps_the_list_marker_indent_and_native_drawing(tmp_path):
    Image.new('RGB', (500, 100), '#123456').save(tmp_path / 'evidence.png')
    document = _render(tmp_path, '65536. ![Evidence caption](evidence.png)\n\n65537. Next item')
    first, second = _numbers(document)
    assert len(first._p.findall(qn('w:fldSimple'))) == 1
    assert len(first._p.xpath('.//w:drawing')) == 1
    assert first.paragraph_format.left_indent == second.paragraph_format.left_indent
    assert first.paragraph_format.keep_with_next is True
    caption = next(p for p in document.paragraphs if p.style.name == 'Caption')
    assert caption.paragraph_format.left_indent == first.paragraph_format.left_indent


def test_three_nested_ten_digit_markers_leave_printable_text_width(tmp_path):
    document = _render(tmp_path, '1234567890. Outer\n\n    1234567890. Middle\n\n        1234567890. Inner content remains readable.\n\n        1. Second inner\n\n    1. Second middle\n\n1. Second outer')
    paragraphs = _numbers(document)
    section = document.sections[-1]
    usable = section.page_width - section.left_margin - section.right_margin
    assert len(paragraphs) == 6
    assert all(usable - p.paragraph_format.left_indent > Cm(5) for p in paragraphs)
    assert all(p.paragraph_format.left_indent + p.paragraph_format.first_line_indent >= 0 for p in paragraphs)
    assert [next(p._p.find(qn('w:fldSimple')).iter(qn('w:t'))).text for p in paragraphs] == ['1234567890', '1234567890', '1234567890', '1234567891', '1234567891', '1234567891']


@pytest.mark.parametrize('start', [1, 42, 999])
def test_short_native_ordinals_reserve_a_readable_gap_and_consistent_wrapping(tmp_path, start):
    document = _render(tmp_path, f'{start}. A long narrative that may wrap.\n\n    Continuation text.')
    first = _numbers(document)[0]
    continuation = next(p for p in document.paragraphs if p.text == 'Continuation text.')
    from PIL import ImageFont
    from pathlib import Path
    font = ImageFont.truetype(str(Path(__file__).parents[1] / 'exporters/fonts/DejaVuSans.ttf'), 1100)
    actual_marker_width = font.getlength(str(start) + '.') / 100
    assert -first.paragraph_format.first_line_indent.pt - actual_marker_width >= 3.5
    assert continuation.paragraph_format.left_indent == first.paragraph_format.left_indent
    assert first.paragraph_format.tab_stops[0].position == first.paragraph_format.left_indent
    assert first._p.find(qn('w:fldSimple')) is None


def test_bullet_geometry_does_not_acquire_numeric_marker_padding(tmp_path):
    document = _render(tmp_path, '- First bullet\n- Second bullet')
    assert all(abs(p.paragraph_format.left_indent - Cm(.65)) < Pt(.05) for p in _numbers(document))
    assert all(abs(p.paragraph_format.first_line_indent + Cm(.4)) < Pt(.05) for p in _numbers(document))


@pytest.mark.parametrize('start', [1, 65536])
def test_code_first_list_marker_stays_with_its_first_code_line(tmp_path, start):
    document = _render(tmp_path, f'{start}. \n\n        CODE FIRST\n\n{start + 1}. Next item')
    marker, following = _numbers(document)
    code = next(p for p in document.paragraphs if p.text == 'CODE FIRST')
    assert marker.paragraph_format.keep_with_next is True
    assert code.paragraph_format.left_indent == marker.paragraph_format.left_indent
    assert code.paragraph_format.keep_with_next is None
    assert following.paragraph_format.keep_with_next is None


@pytest.mark.parametrize('content', ['    | Header | Value |\n    |---|---|\n    | ID | retained |',
                                      '    > A quoted first paragraph.'])
def test_other_first_blocks_keep_only_the_ordinal_boundary(tmp_path, content):
    document = _render(tmp_path, '65536. \n\n' + content + '\n\n65537. Next item')
    marker, following = _numbers(document)
    assert marker.paragraph_format.keep_with_next is True
    assert following.paragraph_format.keep_with_next is None
