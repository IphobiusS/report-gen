"""Long finding prose and authored tables flow without losing content."""
from copy import deepcopy
from docx import Document
from lxml import html
import pdfplumber
import pytest

import delivery
import engine
import export
from wizard import dump_yaml


def document(language):
    narrative = ('Evidence retained without omission.' if language == 'en'
                 else 'Evidencia conservada sin omisiones.')
    finding = {'id': 'F01', 'mode': 'vuln', 'title': 'LONG_NARRATIVE_TITLE',
               'severity': 'medium', 'cwe': 'CWE-200', 'affected': 'srv.example.test'}
    for key, token in [('description_md', 'DESCRIPTION'), ('impact_md', 'IMPACT'),
                       ('remediation_md', 'REMEDIATION')]:
        finding[key] = '\n\n'.join(
            f'{token}_{i:02d} ' + (narrative + ' ') * 9 for i in range(16))
    finding['references'] = ['https://example.test/reference-one', 'https://example.test/reference-two']
    return {'meta': {'lang': language, 'theme': 'corporativo', 'report_title': 'Narrative regression'},
            'report': {'sections': [{'key': 'findings'}]}, 'findings': [finding]}


@pytest.mark.parametrize('language', ['en', 'es'])
@pytest.mark.parametrize('fmt', ['pdf', 'html', 'docx', 'md'])
def test_long_narrative_survives_all_formats_and_remains_in_finding(tmp_path, language, fmt):
    source = tmp_path / 'engagement.yaml'
    data = document(language)
    original = deepcopy(data)
    dump_yaml(data, source)
    before = source.read_bytes()
    captured = delivery.capture(tmp_path, tmp_path / 'snapshot')
    output = export.export(source, fmt, tmp_path / ('report.' + delivery.get(fmt).extension))
    checks = delivery.inspect_export(output, fmt, captured)
    assert checks['paragraphs_checked'] >= 48
    assert source.read_bytes() == before and data == original
    if fmt == 'pdf':
        with pdfplumber.open(output) as pdf:
            pages = [page.extract_text() or '' for page in pdf.pages]
        assert len(pages) > 3
        text = '\n'.join(pages)
        for kind in ('DESCRIPTION', 'IMPACT', 'REMEDIATION'):
            for i in range(16):
                assert text.count(f'{kind}_{i:02d}') == 1
        title_pages = [i for i, text in enumerate(pages) if 'LONG_NARRATIVE_TITLE' in text]
        first_body = next(i for i, text in enumerate(pages) if 'DESCRIPTION_00' in text)
        assert first_body in title_pages
    elif fmt == 'html':
        tree = html.parse(str(output))
        intro = tree.xpath('//*[contains(concat(" ",normalize-space(@class)," ")," finding-intro ")]')[0]
        assert not any(token in intro.text_content() for token in ('DESCRIPTION_', 'IMPACT_', 'REMEDIATION_'))
        assert len(intro.xpath('.//table')) == 1
    elif fmt == 'docx':
        doc = Document(output)
        assert not any('DESCRIPTION_' in table._tbl.xml for table in doc.tables)
        assert any('DESCRIPTION_15' in paragraph.text for paragraph in doc.paragraphs)


@pytest.mark.parametrize('language', ['en', 'es'])
@pytest.mark.parametrize('field', ['description_md', 'impact_md', 'remediation_md'])
def test_real_multipage_markdown_table_preserves_every_row_without_mutation(tmp_path, language, field):
    data = document(language)
    data['findings'][0][field] = '| Evidence | Result |\n| --- | --- |\n' + '\n'.join(
        f'| TABLE_ROW_{i:03d} | Preserved in its original finding |' for i in range(120))
    source = tmp_path / 'engagement.yaml'
    dump_yaml(data, source)
    before = source.read_bytes()
    output = tmp_path / 'report.pdf'
    output.write_bytes(b'previous verified output')
    export.export(source, 'pdf', output)
    assert source.read_bytes() == before
    with pdfplumber.open(output) as pdf:
        pages = [page.extract_text() or '' for page in pdf.pages]
    text = '\n'.join(pages)
    for i in range(120):
        assert text.count(f'TABLE_ROW_{i:03d}') == 1
    occupied = [page for page in pages if 'TABLE_ROW_' in page]
    assert len(occupied) >= 3
    assert all('Evidence' in page and 'Result' in page for page in occupied)


@pytest.mark.parametrize('language', ['en', 'es'])
def test_narrative_page_index_points_to_actual_finding_header(tmp_path, language):
    data = document(language)
    source = engine.render_html(engine.build_env(), data, tmp_path, {}, True, engine.load_lang(language))
    output = tmp_path / 'markers.pdf'
    engine.render_pdf_weasyprint(source, output, tmp_path)
    index = engine.read_page_map(output, ['F01'])
    with pdfplumber.open(output) as pdf:
        page = pdf.pages[index['F01'] - 1].extract_text()
    assert 'LONG_NARRATIVE_TITLE' in page and 'DESCRIPTION_00' in page
