"""Native DOCX preview trust boundary, provenance and owned process lifecycle."""
from hashlib import sha256
from io import BytesIO
import json
import sys
import threading
import time
from types import SimpleNamespace
import zipfile

from docx import Document
import pytest
from pypdf import PdfWriter

import docx_preview as preview
from export_jobs import _start_process


def docx():
    document = Document()
    document.add_heading('Synthetic Meridiano native preview', 0)
    document.add_paragraph('Fictional logistics report; exact DOCX source stays unchanged.')
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf(pages=2):
    stream = BytesIO()
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    writer.write(stream)
    return stream.getvalue()


def repack(raw, changed):
    stream = BytesIO()
    with zipfile.ZipFile(BytesIO(raw)) as source, zipfile.ZipFile(stream, 'w') as destination:
        for name in source.namelist():
            destination.writestr(name, changed.pop(name, source.read(name)))
        for name, content in changed.items():
            destination.writestr(name, content)
    return stream.getvalue()


def verified(raw, name='case.docx', revision='frozen-revision'):
    return {'source_revision': revision, 'files': {name: {'bytes': len(raw), 'sha256': sha256(raw).hexdigest()}}}


def eventually(predicate, seconds=10):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        result = predicate()
        if result:
            return result
        time.sleep(.02)
    pytest.fail('Preview did not reach its expected state')


def test_docx_individual_and_delivery_exact_bytes():
    raw = docx()
    preview.validate_docx(raw)
    assert preview.extract_docx(raw, 'case.docx', verified(raw), 'frozen-revision') == raw
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w') as archive:
        archive.writestr('informe.docx', raw)
        archive.writestr('attachments/not-the-report.docx', b'unrelated')
    assert preview.extract_docx(stream.getvalue(), 'case.zip', verified(raw, 'informe.docx'), 'frozen-revision') == raw


@pytest.mark.parametrize('change', ['revision', 'bytes', 'hash', 'member'])
def test_source_receipt_mismatch_refused(change):
    raw = docx()
    record = verified(raw)
    if change == 'revision':
        record['source_revision'] = 'other'
    elif change == 'bytes':
        record['files']['case.docx']['bytes'] += 1
    elif change == 'hash':
        record['files']['case.docx']['sha256'] = '0' * 64
    else:
        record['files'] = {'other.docx': record['files']['case.docx']}
    with pytest.raises(preview.Rejected) as failure:
        preview.extract_docx(raw, 'case.docx', record, 'frozen-revision')
    assert failure.value.code == 'source'


@pytest.mark.parametrize('name,content', [
    ('word/vbaProject.bin', b'macro'),
    ('word/embeddings/oleObject1.bin', b'object'),
    ('word/activeX/activeX1.xml', b'<x/>'),
    ('../escape.xml', b'<x/>'),
    ('word/evil.xml', b'<!DOCTYPE x [<!ENTITY payload "x">]><x>&payload;</x>'),
    ('word/evil.xml', ('<w:altChunk xmlns:w="' + preview._W + '"/>').encode()),
    ('word/evil.xml', ('<w:instrText xmlns:w="' + preview._W + '">INCLUDETEXT https://example.invalid/private</w:instrText>').encode()),
    ('word/evil.xml', ('<w:fldSimple xmlns:w="' + preview._W + '" w:instr="DDEAUTO cmd /c x"/>').encode()),
])
def test_active_or_invalid_docx_never_reaches_engine(name, content):
    with pytest.raises(preview.Rejected):
        preview.validate_docx(repack(docx(), {name: content}))


@pytest.mark.parametrize('kind,target,allowed', [
    ('hyperlink', 'https://example.invalid/report#part', True),
    ('hyperlink', '#rg_ref_123', True),
    ('hyperlink', 'mailto:synthetic@example.invalid', True),
    ('image', 'https://example.invalid/image.png', False),
    ('attachedTemplate', 'https://example.invalid/template.dotm', False),
    ('hyperlink', 'file:///private/local.txt', False),
    ('hyperlink', '//example.invalid/share', False),
])
def test_external_relationship_policy(kind, target, allowed):
    data = ('<Relationships xmlns="' + preview._REL + '"><Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + kind + '" '
            'Target="' + target + '" TargetMode="External"/></Relationships>').encode()
    raw = repack(docx(), {'word/_rels/extra.xml.rels': data})
    if allowed:
        preview.validate_docx(raw)
    else:
        with pytest.raises(preview.Rejected, match='extern'):
            preview.validate_docx(raw)


@pytest.mark.parametrize('instruction', ['PAGE', 'NUMPAGES', ' PAGEREF rg_toc_49 \\h ',
                                         ' SEQ RGList12 \\r 100001 \\* ARABIC ',
                                         ' SEQ RGList12 \\n \\* ARABIC '])
def test_real_report_field_vocabulary(instruction):
    data = ('<w:instrText xmlns:w="' + preview._W + '">' + instruction + '</w:instrText>').encode()
    preview.validate_docx(repack(docx(), {'word/field.xml': data}))


def test_duplicate_members_and_expanded_budget(monkeypatch):
    raw = docx()
    stream = BytesIO(raw)
    with zipfile.ZipFile(stream, 'a') as archive:
        with pytest.warns(UserWarning):
            archive.writestr('word/document.xml', '<x/>')
    with pytest.raises(preview.Rejected):
        preview.validate_docx(stream.getvalue())
    monkeypatch.setattr(preview, 'MAX_EXPANDED', 10)
    with pytest.raises(preview.Rejected):
        preview.validate_docx(raw)


@pytest.mark.parametrize('target', ['file:///etc/passwd', 'https://example.invalid/image.png',
                                   '../../../private.png', '//example.invalid/share', '%2e%2e/%2e%2e/private.png'])
def test_missing_external_mode_cannot_disguise_external_resource(target):
    data = ('<Relationships xmlns="' + preview._REL + '"><Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            'Target="' + target + '"/></Relationships>').encode()
    with pytest.raises(preview.Rejected):
        preview.validate_docx(repack(docx(), {'word/_rels/document.xml.rels': data}))


def test_utf16_doctype_cannot_bypass_entity_rejection():
    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY a "not-expanded">]><x>&a;</x>'
    with pytest.raises(preview.Rejected):
        preview.validate_docx(repack(docx(), {'word/encoded.xml': xml.encode('utf-16')}))


def test_unavailable_capability_and_no_automatic_install(monkeypatch):
    monkeypatch.setattr(preview, 'executable', lambda: None)
    value = preview.capabilities()
    assert value['available'] is False and value['original_word_verified'] is False
    assert 'LibreOffice' in value['unavailable_message']
    store = preview.Store()
    try:
        with pytest.raises(preview.Rejected) as failure:
            store.start('owner', docx(), 'revision', 'job')
        assert failure.value.code == 'unavailable' and not list(store.root.iterdir())
    finally:
        store.close()


def configure_worker(tmp_path):
    raw = docx()
    (tmp_path / 'source.docx').write_bytes(raw)
    config = {'executable': 'controlled-engine', 'source_docx_sha256': sha256(raw).hexdigest(),
              'revision': 'frozen-revision', 'job_id': 'own-job'}
    (tmp_path / 'config.json').write_text(json.dumps(config))
    return raw, config


def test_worker_receipt_identifies_derived_layout_and_unchanged_source(tmp_path, monkeypatch):
    raw, config = configure_worker(tmp_path)
    calls = []
    def run(command, **kwargs):
        calls.append(command)
        if '--version' in command:
            return SimpleNamespace(stdout=b'LibreOffice 24.2.7.2\n', returncode=0)
        (tmp_path / 'rendered/source.pdf').write_bytes(pdf(3))
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(preview.subprocess, 'run', run)
    # Resource limits belong to a disposable subprocess, not the pytest runner.
    monkeypatch.setattr(preview, '_worker_limits', lambda: None)
    assert preview.worker(tmp_path) == 0
    receipt = json.loads((tmp_path / 'receipt.json').read_text())
    assert receipt['pages'] == 3 and receipt['source_docx_sha256'] == config['source_docx_sha256']
    assert receipt['engine']['version'] == 'LibreOffice 24.2.7.2' and receipt['word_verified'] is False
    assert receipt['visual_acceptance'] == 'not_assessed' and (tmp_path / 'source.docx').read_bytes() == raw
    assert all('-env:UserInstallation=' + (tmp_path / 'profile').as_uri() in call for call in calls)
    assert 'MacroSecurityLevel' in (tmp_path / 'profile/user/registrymodifications.xcu').read_text()


@pytest.mark.parametrize('outcome', ['no_output', 'too_many_pages', 'changed_source', 'bad_engine'])
def test_worker_rejects_incomplete_or_unbound_result(tmp_path, monkeypatch, outcome):
    configure_worker(tmp_path)
    def run(command, **kwargs):
        if '--version' in command:
            return SimpleNamespace(stdout=b'Other editor' if outcome == 'bad_engine' else b'LibreOffice 24.2.7.2', returncode=0)
        if outcome != 'no_output':
            (tmp_path / 'rendered/source.pdf').write_bytes(pdf(3))
        if outcome == 'changed_source':
            (tmp_path / 'source.docx').write_bytes(b'changed')
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(preview, '_worker_limits', lambda: None)
    monkeypatch.setattr(preview.subprocess, 'run', run)
    if outcome == 'too_many_pages':
        monkeypatch.setattr(preview, 'MAX_PAGES', 2)
    assert preview.worker(tmp_path) == 1
    assert not (tmp_path / 'preview.pdf').exists() and not (tmp_path / 'receipt.json').exists()


@pytest.mark.parametrize('mode', ['cancel', 'timeout', 'close'])
def test_real_owned_worker_stops_and_deletes_source(tmp_path, monkeypatch, mode):
    controlled = tmp_path / 'controlled.py'
    controlled.write_text('import sys,time\nif sys.stdin.buffer.read(1)!=b"1":raise SystemExit(99)\ntime.sleep(60)\n')
    started = threading.Event()
    processes = []
    def launch(_command):
        process = _start_process([sys.executable, str(controlled)])
        processes.append(process)
        started.set()
        return process
    monkeypatch.setattr(preview, '_start_process', launch)
    monkeypatch.setattr(preview, 'executable', lambda: sys.executable)
    store = preview.Store(timeout=.2 if mode == 'timeout' else 10)
    try:
        row = store.start('owner', docx(), 'revision', 'job')
        assert started.wait(5)
        with pytest.raises(KeyError):
            store.status('other-owner', row['id'])
        with pytest.raises(preview.Rejected) as failure:
            store.start('owner', docx(), 'revision', 'other-job')
        assert failure.value.code == 'busy'
        if mode == 'cancel':
            store.cancel('owner', row['id'])
        elif mode == 'close':
            store.close()
        eventually(lambda: not store.records[row['id']]['thread'].is_alive())
        final = store.records[row['id']]
        assert final['stage'] == ('error' if mode == 'timeout' else 'cancelled')
        assert final['error_code'] == ('timeout' if mode == 'timeout' else None)
        assert not (final['root'] / 'source.docx').exists()
        assert processes[0].returncode is not None
    finally:
        store.close()


def test_native_libreoffice_actual_docx_conversion():
    if not preview.executable():
        pytest.skip('LibreOffice Writer is optional and not installed')
    raw = docx()
    store = preview.Store()
    try:
        row = store.start('owner', raw, 'native-revision', 'native-job')
        def finished():
            value = store.status('owner', row['id'])
            return value if value['stage'] in ('done', 'error', 'cancelled') else None
        final = eventually(finished, seconds=190)
        assert final['stage'] == 'done', final
        receipt = final['receipt']
        assert receipt['source_docx_sha256'] == sha256(raw).hexdigest()
        assert receipt['source_revision'] == 'native-revision' and receipt['word_verified'] is False
        assert receipt['pages'] >= 1
        assert store.result('owner', row['id']).startswith(b'%PDF-')
        assert {p.name for p in store.records[row['id']]['root'].iterdir()} == {'preview.pdf', 'receipt.json'}
    finally:
        store.close()
