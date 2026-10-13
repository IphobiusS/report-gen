"""Fragment links keep their complete click range through native DOCX saves."""
import os
import re
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
import pytest

from exporters import richtext

URL = 'https://archive.albor.example/café/東京?label=Δ&revision=ALB%2D015#revisión'
TARGET = 'https://archive.albor.example/caf%C3%A9/%E6%9D%B1%E4%BA%AC?label=%CE%94&revision=ALB%2D015#revisi%C3%B3n'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def events(element):
    result = []
    for run in element.iter(qn('w:r')):
        properties = run.find(qn('w:rPr'))
        style = etree.tostring(properties) if properties is not None else None
        for child in run:
            if child.tag == qn('w:rPr'):
                continue
            if child.tag == qn('w:t'):
                result.extend(('text', char, style) for char in child.text or '')
            else:
                result.append(('object', etree.tostring(child), style))
    return result


def document(text, tmp_path):
    result = Document()
    richtext.write_markdown(result, text, tmp_path, result.add_paragraph)
    return result


@pytest.mark.parametrize('label', [URL, 'ASCII **bold** and *italic* label',
    'Latin東京Latin', 'cafe\u0301東京\ufe00tail', 'a𠀀\U000e0100z',
    'before👩\u200d💻東京after', '한글LatinかなLatinカナ'])
def test_fragment_link_segmentation_preserves_every_character_style_and_target(tmp_path, monkeypatch, label):
    markup = f'Before [{label}]({URL}) after.'
    original = richtext._separate_hyperlink_runs
    with monkeypatch.context() as patch:
        patch.setattr(richtext, '_separate_hyperlink_runs', lambda link: None)
        before = document(markup, tmp_path)
    after = document(markup, tmp_path)
    assert events(before.element) == events(after.element)
    expected = before.part.rels[before.element.xpath('//w:hyperlink')[0].get(qn('r:id'))].target_ref
    assert expected == TARGET
    links = after.element.xpath('//w:hyperlink')
    assert len(links) >= 2
    assert all(after.part.rels[link.get(qn('r:id'))].target_ref == expected for link in links)
    assert not after.element.xpath('//w:lang')
    for link in links:
        before_xml = etree.tostring(link)
        original(link)
        assert etree.tostring(link) == before_xml
    # The logical label is copyable without hidden break hints or extra spaces.
    assert ''.join(after.element.xpath('//w:t/text()')) == ''.join(before.element.xpath('//w:t/text()'))


@pytest.mark.parametrize('target', ['https://example.test/no-fragment', '#heading'])
def test_links_outside_the_affected_external_fragment_path_are_xml_identical(tmp_path, monkeypatch, target):
    markup = f'[Latin**東京**Latin]({target})'
    with monkeypatch.context() as patch:
        patch.setattr(richtext, '_separate_hyperlink_runs', lambda link: None)
        before = document(markup, tmp_path)
    after = document(markup, tmp_path)
    assert etree.tostring(after.element) == etree.tostring(before.element)


@pytest.mark.parametrize('parts', [('東', '\u0301', 'x'), ('𠀀', '\U000e0100', 'x'),
                                   ('👩', '\u200d', '💻', '東'), ('e', '\u0301', '東')])
def test_existing_style_boundaries_do_not_create_new_cluster_run_boundaries(parts):
    doc = Document();p = doc.add_paragraph();link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), 'rId42');link.set(qn('w:history'), '1');p._p.append(link)
    for index, part in enumerate(parts):
        run = p.add_run(part);run.bold = bool(index % 2);link.append(run._r)
    before = events(p._p)
    richtext._separate_hyperlink_runs(link)
    assert events(p._p) == before
    texts = [''.join(n.xpath('.//w:t/text()')) for n in p._p if n.tag == qn('w:hyperlink')]
    assert texts == list(parts)
    assert all(n.get(qn('w:history')) == '1' for n in p._p)


@pytest.mark.parametrize('kind', ['br', 'tab', 'drawing', 'empty'])
def test_nontext_run_nodes_are_preserved_exactly_and_not_duplicated(kind):
    doc=Document();p=doc.add_paragraph();link=OxmlElement('w:hyperlink');p._p.append(link)
    link.set(qn('r:id'),'rId42')
    for text in ['Latin東京', 'tail']:
        run=p.add_run(text);link.append(run._r)
        if text.startswith('Latin'):
            item=OxmlElement('w:r')
            if kind!='empty':item.append(OxmlElement('w:'+kind))
            link.append(item)
    before=events(p._p)
    richtext._separate_hyperlink_runs(link)
    assert events(p._p)==before
    if kind!='empty':assert len(p._p.xpath('.//w:'+kind))==1
    else:assert len(p._p.xpath('.//w:r[not(*)]'))==1


def associations(path):
    """Native fields and source links compared by actual visible character."""
    with ZipFile(path) as archive:
        root=etree.fromstring(archive.read('word/document.xml'))
        rels=etree.fromstring(archive.read('word/_rels/document.xml.rels'))
    targets={n.get('Id'):n.get('Target') for n in rels if n.get('Type','').endswith('/hyperlink')}
    rows=[]
    for paragraph in root.xpath('//w:body//w:p', namespaces={'w':W[1:-1]}):
        output=[];fields=[]
        def target(code):
            match=re.search(r'\bHYPERLINK\s+"([^"\r\n]*)"(?:\s+\\l\s+"([^"\r\n]*)")?',code)
            return match[1]+('#'+match[2] if match[2] is not None else '') if match else None
        def visit(node,uri=None):
            if node.tag in (W+'pPr',W+'rPr'):return
            if node.tag==W+'hyperlink':uri=targets.get(node.get(qn('r:id')))
            if node.tag==W+'fldSimple':uri=target(node.get(W+'instr',''))
            if node.tag==W+'fldChar':
                kind=node.get(W+'fldCharType')
                if kind=='begin':fields.append({'code':'','result':False})
                elif kind=='separate':fields[-1]['result']=True
                elif kind=='end':fields.pop()
                return
            if node.tag==W+'instrText':
                if fields:fields[-1]['code']+=node.text or ''
                return
            if node.tag in (W+'t',W+'tab',W+'br',W+'cr'):
                if fields and not all(field['result'] for field in fields):return
                uri=next((target(field['code']) for field in reversed(fields) if target(field['code'])),uri)
                text=node.text or '' if node.tag==W+'t' else '\t' if node.tag==W+'tab' else '\n'
                output.extend((char,uri) for char in text);return
            for child in node:visit(child,uri)
        visit(paragraph)
        assert not fields
        if output:rows.append(output)
    return rows


def test_character_association_guard_rejects_partial_link_with_unchanged_text(tmp_path):
    doc=document(f'[{URL}]({URL})',tmp_path);before=tmp_path/'before.docx';doc.save(before)
    link=doc.element.xpath('//w:hyperlink')[1]
    parent=link.getparent();position=parent.index(link)
    for child in list(link):parent.insert(position,child);position+=1
    parent.remove(link);after=tmp_path/'after.docx';doc.save(after)
    assert [[c for c,_ in row] for row in associations(before)]==[[c for c,_ in row] for row in associations(after)]
    assert associations(before)!=associations(after)


@pytest.mark.skipif(not os.environ.get('REPORT_GEN_TEST_SOFFICE'), reason='Set REPORT_GEN_TEST_SOFFICE for native DOCX link acceptance')
@pytest.mark.parametrize('label', [URL, 'ASCII **bold** and *italic* label',
                                   'LatinالعربيةLatin', 'LatinदेवनागरीLatin', 'Latin각Latin',
                                   'e**́**東京', 'before👩\u200d💻東京after'])
def test_native_fragment_links_retain_whole_click_ranges_after_two_saves(tmp_path, label):
    doc=document(f'[{label}]({URL})',tmp_path)
    source=tmp_path/'source.docx';doc.save(source);expected=associations(source)
    for number in range(2):
        output=tmp_path/str(number);output.mkdir();profile=tmp_path/f'profile-{number}'
        result=subprocess.run([os.environ['REPORT_GEN_TEST_SOFFICE'],'-env:UserInstallation='+profile.as_uri(),
            '--headless','--nologo','--norestore','--convert-to','docx:Office Open XML Text',
            '--outdir',str(output),str(source)],capture_output=True,text=True,timeout=90)
        assert result.returncode==0,result.stdout+result.stderr
        source=output/source.name
        assert associations(source)==expected
