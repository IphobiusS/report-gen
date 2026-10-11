"""Nested evidence images must fit the containing text area without clipping."""
import hashlib
from io import BytesIO
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image
import pytest

import engine
import export
from exporters.images import paragraph_width
from exporters.richtext import write_markdown


def make_image(tmp_path):
    path = tmp_path / 'evidence.png'
    Image.new('RGB', (1800, 250), 'navy').save(path)
    return path


@pytest.mark.parametrize('width', ['100%', '90cm', '9000px', '50%'])
@pytest.mark.parametrize('context', ['list', 'nested_list', 'table'])
def test_figure_width_respects_list_indent_or_cell_padding(tmp_path, width, context):
    source = make_image(tmp_path)
    image = f'![Exact caption](evidence.png){{width={width}}}'
    if context == 'list':
        body = '1. Before ' + image + ' After.'
    elif context == 'nested_list':
        body = '1. Outer\n    1. Before ' + image + ' After.'
    else:
        body = '| Evidence | Figure |\n|---|---|\n| Reference | ' + image + ' |'
    doc = Document()
    write_markdown(doc, body, tmp_path, lambda _: None)
    paragraph = next(p for p in (doc.paragraphs if context != 'table' else doc.tables[0].rows[1].cells[1].paragraphs)
                     if p._p.xpath('.//wp:inline'))
    shape = doc.inline_shapes[0]
    if context == 'table':
        expected = doc.tables[0].rows[1].cells[1].width - Pt(10.8)
    else:
        section = doc.sections[-1]
        expected = section.page_width - section.left_margin - section.right_margin - paragraph.paragraph_format.left_indent
    expected = int(expected * (.5 if width == '50%' else 1))
    assert abs(shape.width - expected) <= 2
    assert abs(shape.width / shape.height - 1800 / 250) < .001
    assert shape._inline.docPr.get('descr') == 'Exact caption'
    assert all(shape._inline.get(edge) == '0' for edge in ('distL', 'distR', 'distT', 'distB'))
    target = BytesIO(); doc.save(target)
    with zipfile.ZipFile(target) as archive:
        media = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
    assert [hashlib.sha256(blob).digest() for blob in media] == [hashlib.sha256(source.read_bytes()).digest()]


def test_containing_width_honors_explicit_cell_margins_and_paragraph_indents():
    doc = Document(); table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0); cell.width = Cm(10)
    margins = OxmlElement('w:tcMar')
    for edge, value in [('start', 180), ('end', 240)]:
        margin = OxmlElement('w:' + edge); margin.set(qn('w:w'), str(value)); margins.append(margin)
    cell._tc.get_or_add_tcPr().append(margins)
    p = cell.paragraphs[0]
    p.paragraph_format.left_indent = Cm(.5)
    p.paragraph_format.right_indent = Cm(.4)
    p.paragraph_format.first_line_indent = Cm(.3)
    expected = cell.width - Pt(21) - p.paragraph_format.left_indent - p.paragraph_format.right_indent - p.paragraph_format.first_line_indent
    assert paragraph_width(doc, p) == expected


def test_structured_evidence_and_severity_chart_have_no_implicit_side_offset(tmp_path):
    make_image(tmp_path)
    data = {'meta': {'lang': 'en', 'theme': 'ejecutivo', 'report_title': 'Figure bounds', 'date': '2026-10-11', 'version': '1',
                     'assessor': 'Example analyst', 'client': 'Example bank', 'osid': 'LAB', 'email': 'test@example.test'},
            'findings': [{'id': 'F-001', 'title': 'Synthetic finding', 'severity': 'high',
                          'description_md': 'Synthetic evidence.', 'walkthrough': [{'figure': {'src': 'evidence.png', 'caption': 'Evidence caption'}}]}]}
    target = tmp_path / 'report.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, target)
    doc = Document(target)
    figures = [shape for shape in doc.inline_shapes if shape._inline.docPr.get('descr') in {'Evidence caption', 'report-gen:severity-chart'}]
    assert len(figures) == 2
    for shape in figures:
        assert all(shape._inline.get(edge) == '0' for edge in ('distL', 'distR', 'distT', 'distB'))
