"""A native contents index must recalculate actual page references after edits."""
import importlib.util
import os
from pathlib import Path
import re

import fitz
from docx import Document
from docx.shared import RGBColor
import pytest

from exporters.navigation import add_toc


def acceptance():
    path = Path(__file__).resolve().parents[1] / 'scripts/verify-libreoffice.py'
    spec = importlib.util.spec_from_file_location('contents_acceptance', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_document(path, count=3, *, long_titles=False, body=True):
    document = Document()
    document.add_paragraph('Contents field regression fixture')
    document.add_page_break()
    entries = []
    for number in range(1, count + 1):
        title = f'Chapter {number} distinct target'
        if long_titles:
            title += ': ' + 'a wrapped title whose reference stays aligned and editable ' * 3
        heading = document.add_heading(title, 1)
        heading.paragraph_format.page_break_before = True
        entries.append((heading, heading.text, 1))
        for paragraph in range(number * 4 if body else 0):
            document.add_paragraph(f'Body {number}-{paragraph}. ' + 'Deterministic documentary text. ' * 7)
    add_toc(document, entries, 'Contents', RGBColor(0, 0, 0))
    document.save(path)
    return document


def test_contents_keep_only_independent_live_fields(tmp_path):
    document = make_document(tmp_path / 'contents.docx', count=5)
    fields = document.element.xpath('//w:instrText')
    assert len(fields) == 5
    assert all(re.fullmatch(r' PAGEREF rg_toc_\d+ \\h ', field.text) for field in fields)
    # A wrapper TOC field makes LibreOffice treat all nested page references
    # as stale cached text. Every contents field has depth exactly one.
    depth = 0
    for node in document.element.xpath('//w:fldChar'):
        kind = node.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType')
        if kind == 'begin':
            depth += 1
            assert depth == 1
        elif kind == 'end':
            depth -= 1
        assert depth >= 0
    assert depth == 0
    assert len(document.element.xpath('//w:sdt/w:sdtContent/w:p/w:hyperlink')) == 5
    assert not document.element.xpath('//w:sdtPr/w:lock')


@pytest.mark.parametrize('mutation', ['none', 'missing-number', 'wrong-number', 'wrong-destination', 'wrong-heading',
                                     'swapped-number-rows'])
def test_contents_guard_checks_printed_numbers_and_heading_destinations(tmp_path, mutation):
    source = tmp_path / 'source.docx'
    make_document(source)
    rendered = tmp_path / 'synthetic-guard-fixture.pdf'
    # This is deliberately a synthetic PDF for negative guard tests, not
    # evidence of a successful LibreOffice rendering.
    with fitz.open() as pdf:
        for _ in range(4):
            pdf.new_page()
        toc = pdf[0]
        toc.insert_text((60, 50), 'Synthetic native-contents guard fixture')
        for number in range(1, 4):
            title = f'Chapter {number} distinct target'
            heading = 'Wrong heading' if mutation == 'wrong-heading' and number == 2 else title
            pdf[number].insert_text((60, 100), heading)
            y = 100 + number * 25
            toc.insert_text((60, y), title)
            title_box = toc.search_for(title)[0]
            target = 3 if mutation == 'wrong-destination' and number == 2 else number
            toc.insert_link({'kind': fitz.LINK_GOTO, 'from': title_box, 'page': target, 'to': fitz.Point(60, 85)})
            if mutation == 'missing-number' and number == 2:
                continue
            shown = '9' if mutation == 'wrong-number' and number == 2 else str(number + 1)
            number_y = 100 + (3 - number) * 25 if mutation == 'swapped-number-rows' and number < 3 else y
            toc.insert_text((500, number_y), shown)
            box = fitz.Rect(498, number_y - 12, 515, number_y + 3)
            toc.insert_link({'kind': fitz.LINK_GOTO, 'from': box, 'page': number, 'to': fitz.Point(60, 85)})
        pdf.save(rendered)
    if mutation == 'none':
        assert acceptance().inspect_contents(source, rendered)['entries_checked'] == 3
    else:
        with pytest.raises(ValueError, match='Contents|contents'):
            acceptance().inspect_contents(source, rendered)


@pytest.mark.skipif(not os.environ.get('REPORT_GEN_TEST_SOFFICE'),
                    reason='Set REPORT_GEN_TEST_SOFFICE for native LibreOffice field acceptance')
def test_native_page_references_recalculate_after_reflow(tmp_path):
    native = acceptance()
    source = tmp_path / 'original.docx'
    document = make_document(source)
    executable = os.environ['REPORT_GEN_TEST_SOFFICE']
    pdf = native.convert(executable, source, tmp_path / 'original-render')
    before = native.inspect_contents(source, pdf)
    assert before['entries_checked'] == 3
    second = next(p for p in document.paragraphs if p.text == 'Chapter 2 distinct target')
    # Ordinary Word edit: insert body content without rebuilding the index,
    # changing its fields, or writing any page-number cache.
    for number in range(30):
        second.insert_paragraph_before(f'Inserted paragraph {number}. ' + 'Extra editable text. ' * 15)
    edited = tmp_path / 'edited.docx'
    document.save(edited)
    edited_pdf = native.convert(executable, edited, tmp_path / 'edited-render')
    after = native.inspect_contents(edited, edited_pdf)
    assert after['entries_checked'] == 3
    assert after['rows'][0]['shown'] == before['rows'][0]['shown']
    assert all(new['shown'] > old['shown'] for new, old in zip(after['rows'][1:], before['rows'][1:]))


@pytest.mark.skipif(not os.environ.get('REPORT_GEN_TEST_SOFFICE'),
                    reason='Set REPORT_GEN_TEST_SOFFICE for native LibreOffice field acceptance')
def test_native_contents_with_wrapped_titles_spanning_multiple_pages(tmp_path):
    source = tmp_path / 'long-index.docx'
    make_document(source, count=20, long_titles=True, body=False)
    native = acceptance()
    pdf = native.convert(os.environ['REPORT_GEN_TEST_SOFFICE'], source, tmp_path / 'render')
    result = native.inspect_contents(source, pdf)
    assert result['entries_checked'] == 20
    assert len({row['contents_page'] for row in result['rows']}) > 1
