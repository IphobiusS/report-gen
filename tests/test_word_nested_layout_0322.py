"""Keep chains and widths use the space of the actual Word content block."""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import pytest

import engine
import export
from design_library import profile
from design_samples import sample_data
from exporters.pagination import keep_table_together, code_fits_page
from exporters.richtext import write_markdown


@pytest.mark.parametrize('explicit_sections', [True, False])
def test_section_boundary_does_not_need_a_separate_break_paragraph(tmp_path, explicit_sections):
    data = sample_data(profile('enterprise-assurance'), 'quiet', 'en')
    if explicit_sections:
        data['report']['sections'] = [
            {'key': 'one', 'title': 'First section', 'body': '**Affected assets**\n\n- 192.0.2.1\n- 192.0.2.2'},
            {'key': 'two', 'title': 'Next section', 'body': 'Independent next section.'},
        ]
    else:
        data.pop('report', None)
    output = tmp_path / 'section-boundary.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, output)
    document = Document(output)
    headings = [p for p in document.paragraphs if p.style.name == 'Heading 1']
    assert len(headings) >= 2
    for paragraph in headings:
        assert paragraph.paragraph_format.page_break_before is True
        previous = paragraph._p.getprevious()
        assert not any(node.get(qn('w:type')) == 'page'
                       for node in previous.iter(qn('w:br')))
    assert all(p.paragraph_format.page_break_before is not True
               for p in document.paragraphs if p.style.name != 'Heading 1')


@pytest.mark.parametrize('row_count', [4, 13, 24])
def test_few_tall_rows_do_not_force_an_oversized_whole_table(row_count):
    document = Document()
    table = document.add_table(rows=row_count, cols=2)
    for index, row in enumerate(table.rows):
        row.cells[0].text = f'ROW-{index:02d}'
        row.cells[1].text = 'Detailed evidence with original content and context. ' * 40
    keep_table_together(table)
    assert all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows)
    assert not any(p.paragraph_format.keep_with_next
                   for row in table.rows for cell in row.cells for p in cell.paragraphs)


def test_short_table_still_moves_as_a_whole_when_it_fits():
    document = Document()
    table = document.add_table(rows=3, cols=2)
    for index, row in enumerate(table.rows):
        row.cells[0].text = str(index)
        row.cells[1].text = 'Verified'
    keep_table_together(table)
    for index, row in enumerate(table.rows):
        assert all(p.paragraph_format.keep_with_next is (index < 2)
                   for cell in row.cells for p in cell.paragraphs)


def test_compact_body_spacing_does_not_expand_grid_table_rows():
    document = Document()
    normal = document.styles['Normal'].paragraph_format
    normal.space_after = Pt(4)
    normal.line_spacing = 1.05
    body = document.add_paragraph('Body keeps its compact spacing.')
    table = document.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    explicit = table.cell(1, 1).paragraphs[0].paragraph_format
    explicit.space_after = Pt(7)
    explicit.line_spacing = 1.2
    metadata = document.add_table(rows=1, cols=1)
    keep_table_together(table)
    keep_table_together(metadata)
    for paragraph in (table.cell(0, 0).paragraphs[0], table.cell(1, 0).paragraphs[0]):
        assert paragraph.paragraph_format.space_after == Pt(0)
        assert paragraph.paragraph_format.line_spacing == 1
    assert explicit.space_after == Pt(7)
    assert explicit.line_spacing == 1.2
    assert metadata.cell(0, 0).paragraphs[0].paragraph_format.space_after is None
    assert metadata.cell(0, 0).paragraphs[0].paragraph_format.line_spacing is None
    assert body.paragraph_format.space_after is None
    assert body.paragraph_format.line_spacing is None
    assert normal.space_after == Pt(4)
    assert normal.line_spacing == 1.05


@pytest.mark.parametrize('count,text,kept', [(3, '192.0.2.1', True), (5, '192.0.2.1', False), (3, 'Long evidence ' * 30, False)])
def test_short_labelled_list_does_not_orphan_last_item(tmp_path, count, text, kept):
    document = Document()
    body = '**Affected assets**\n\n' + '\n'.join('- ' + text for _ in range(count)) + '\n\nAfter the list.'
    write_markdown(document, body, tmp_path, lambda _: None)
    label, *items, after = document.paragraphs
    assert label.paragraph_format.keep_with_next is True
    assert len(items) == count
    if kept:
        assert all(p.paragraph_format.keep_together for p in items)
        assert all(p.paragraph_format.keep_with_next for p in items[:-1])
        assert items[-1].paragraph_format.keep_with_next is False
    else:
        assert all(p.paragraph_format.keep_with_next is None for p in items)
    assert after.paragraph_format.keep_with_next is None


@pytest.mark.parametrize('content', ['- Parent\n    - Nested\n- Second', '- Parent\n\n    ![Evidence](evidence.png)\n\n- Second'])
def test_nested_or_illustrated_labelled_list_is_not_chained(tmp_path, content):
    from PIL import Image
    Image.new('RGB', (20, 20), 'white').save(tmp_path / 'evidence.png')
    document = Document()
    write_markdown(document, '**Affected assets**\n\n' + content, tmp_path, lambda _: None)
    first_item = next(p for p in document.paragraphs if p.text.strip() == 'Parent')
    assert first_item.paragraph_format.keep_with_next is None
    assert first_item.paragraph_format.keep_together is None


@pytest.mark.parametrize('depth', [0, 1, 2, 4])
def test_markdown_table_width_and_position_follow_list_depth(tmp_path, depth):
    table = '| Marker | Evidence |\n|---|---|\n| ROW | Content |'
    for _ in range(depth):
        table = '1. Parent\n\n' + '\n'.join('    ' + line for line in table.splitlines())
    document = Document()
    write_markdown(document, table, tmp_path, lambda _: None)
    result = document.tables[0]
    section = document.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    parents = [p for p in document.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    content_indent = parents[-1].paragraph_format.left_indent if parents else 0
    width = available - content_indent
    indent = result._tbl.tblPr.find(qn('w:tblInd'))
    if depth:
        assert int(indent.get(qn('w:w'))) == round(content_indent / 635)
    else:
        assert indent is None
    assert abs(sum(column.width for column in result.columns) - width) < Pt(.1)
    assert abs(sum(cell.width for cell in result.rows[0].cells) - width) < Pt(.1)
    assert [[cell.text for cell in row.cells] for row in result.rows] == [['Marker', 'Evidence'], ['ROW', 'Content']]


def test_nested_code_uses_its_narrower_area_before_requesting_keep_lines(tmp_path):
    data = sample_data(profile('enterprise-assurance'), 'quiet', 'en')
    # Fits in one line at page width but wraps at the fourth list level.
    transcript = '\n'.join(f'ROW{i:02d}-' + 'x' * 76 for i in range(30))
    body = '\n'.join('    ' + line for line in transcript.splitlines())
    for _ in range(4):
        body = '1. Collection\n\n' + '\n'.join('    ' + line for line in body.splitlines())
    data['report']['sections'] = [{'key': 'nested', 'title': 'Nested evidence', 'body': body}]
    output = tmp_path / 'nested.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, output)
    document = Document(output)
    code = next(p for p in document.paragraphs if p.style.name == 'ReportGenCode')
    assert code.text == transcript
    parents = [p for p in document.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert code.paragraph_format.left_indent == parents[-1].paragraph_format.left_indent
    assert code_fits_page(document, transcript)
    assert code.paragraph_format.keep_together is False
