"""Fenced code preserves authored literals, containers and Word indentation."""
from docx import Document
from docx.oxml.ns import qn
from lxml import html
import pytest

from exporters.richtext import write_markdown
from safe_markup import md, raw_image_sources, replace_image_destinations


LITERAL = (
    "reference=000000000000000000000042\n"
    "\tstate= pending  \n"
    "\n"
    "  indented = value\n"
    "literal <script>alert(1)</script> & <img src='local.png'>\n"
    "![example](img/private.png) [private](file:///tmp/example)\n"
    "# Heading is code\n"
    "1. list-looking code\n"
    "| a | b |\n"
    "|---|---|"
)


def tree(source):
    return html.fragment_fromstring(str(md(source, img_base='')), create_parent='div')


def fenced_list(start=713, fence='```', language='text', payload=LITERAL, depth=0, spaces=4):
    # The installed parser uses four spaces per list level independently of
    # the ordinal's width; aligning a ten-digit marker with twelve spaces
    # instead denotes intentionally indented literal code.
    indent = ' ' * spaces
    source = f'{start}. BEGIN\n\n' + '\n'.join(
        indent + line for line in (fence + language + '\n' + payload + '\n' + fence).split('\n')
    ) + f'\n\n{indent}CONTINUATION\n\n{start+1}. NEXT\n\nOUTSIDE'
    for _ in range(depth):
        source = '7. PARENT\n\n' + '\n'.join('    ' + line for line in source.split('\n'))
    return source


@pytest.mark.parametrize('start', [1, 713, 65536, 2147483646])
@pytest.mark.parametrize('fence', ['```', '~~~~'])
def test_nested_fence_is_a_literal_block_with_continuation(start, fence):
    rendered = tree(fenced_list(start, fence))
    code = rendered.xpath('.//ol/li/pre/code')
    assert len(code) == 1
    assert code[0].text == LITERAL + '\n'
    assert not code[0].xpath('.//*')
    assert code[0].get('class') == 'language-text'
    assert not rendered.xpath('.//script|.//img|.//table|.//h1')
    ordered = rendered.xpath('./ol')
    assert len(ordered) == 1
    assert int(ordered[0].get('start', '1')) == start
    assert len(ordered[0].findall('li')) == 2
    assert ordered[0][0].xpath('./p')[1].text == 'CONTINUATION'
    assert ordered[0][1].text_content().strip() == 'NEXT'
    assert rendered.xpath('./p')[-1].text == 'OUTSIDE'


@pytest.mark.parametrize('language', ['', 'python', 'bash', 'text'])
@pytest.mark.parametrize('depth', [1, 3])
def test_deep_nested_fences_exclude_the_info_string(language, depth):
    rendered = tree(fenced_list(language=language, depth=depth))
    code = rendered.xpath('.//pre/code')
    assert len(code) == 1
    assert code[0].text == LITERAL + '\n'
    assert len(code[0].xpath('ancestor::li')) == depth + 1
    if language:
        assert code[0].get('class') == 'language-' + language


@pytest.mark.parametrize('fence,payload', [
    ('````', 'inner fence example\n```python\nprint(42)\n```'),
    ('~~~~', 'mixed delimiter\n```\n~~~\n~~~text'),
])
def test_shorter_and_mixed_fences_are_literal_content(fence, payload):
    code = tree(fenced_list(fence=fence, payload=payload)).xpath('.//pre/code')
    assert len(code) == 1
    assert code[0].text == payload + '\n'


def test_two_different_code_blocks_stay_in_their_own_list_items():
    source = '5. ONE\n\n    ```text\n    first\n    ```\n\n6. TWO\n\n    ~~~python\n    second\n    ~~~\n\n7. THREE'
    root = tree(source)
    assert [c.text for c in root.xpath('.//pre/code')] == ['first\n', 'second\n']
    items = root.xpath('./ol/li')
    assert len(items) == 3
    assert items[0].xpath('./pre/code')[0].text == 'first\n'
    assert items[1].xpath('./pre/code')[0].text == 'second\n'
    assert items[2].text_content().strip() == 'THREE'


def test_report_reproducer_with_five_spaces_under_item_713():
    code = tree(fenced_list(spaces=5)).xpath('.//ol/li/pre/code')
    assert len(code) == 1
    assert code[0].text == LITERAL + '\n'


@pytest.mark.parametrize('spaces', [8, 12])
def test_extra_indentation_preserves_authored_fences_as_literal_code(spaces):
    source = fenced_list(start=2147483646, spaces=spaces, payload='literal')
    code = tree(source).xpath('.//pre/code')
    assert len(code) == 1
    # Do not reinterpret or strip an intentionally indented code example.
    prefix = ' ' * (spaces - 8)
    assert code[0].text == '\n'.join(prefix + line if line else '' for line in [
        '```text', 'literal', '```', '', 'CONTINUATION',
    ]) + '\n'


def test_code_literals_do_not_become_editable_image_destinations():
    source = fenced_list(payload='![literal](img/private.png)\n[local](file:///tmp/private)')
    assert raw_image_sources(source) == set()
    assert replace_image_destinations(source, 'img/private.png', 'img/changed.png') == source


@pytest.mark.parametrize('start', [713, 65536, 2147483646])
def test_docx_code_callback_receives_exact_payload_and_keeps_list_indent(tmp_path, start):
    document = Document()
    calls = []

    def code_block(text):
        calls.append(text)
        return document.add_paragraph(text)

    write_markdown(document, fenced_list(start), tmp_path, code_block)
    assert calls == [LITERAL]
    paragraphs = document.paragraphs
    first = next(p for p in paragraphs if p.text.endswith('BEGIN'))
    code = next(p for p in paragraphs if p.text == LITERAL)
    continuation = next(p for p in paragraphs if p.text == 'CONTINUATION')
    assert code.paragraph_format.left_indent == first.paragraph_format.left_indent
    assert continuation.paragraph_format.left_indent == first.paragraph_format.left_indent
    assert code._p.pPr.find(qn('w:numPr')) is None
    assert any(p.text.endswith('NEXT') for p in paragraphs)
    assert paragraphs[-1].text == 'OUTSIDE'


@pytest.mark.parametrize('source,expected', [
    ('```text\n\talpha  \n\n beta\n```', '\talpha  \n\n beta\n'),
    ('    alpha\n    beta', 'alpha\nbeta\n'),
])
def test_existing_top_level_and_indented_code_are_unchanged(source, expected):
    nodes = tree(source).xpath('.//pre/code')
    assert len(nodes) == 1
    assert nodes[0].text == expected


def test_inline_backticks_remain_inline_code():
    root = tree('1. Prose with ``literal ` value`` then normal text.\n2. NEXT')
    assert not root.xpath('.//pre')
    assert root.xpath('.//code')[0].text == 'literal ` value'


def test_quote_fence_preserves_literal_content_and_quote_scope():
    source = '> Before\n>\n> ```text\n> alpha\n> beta\n> ```\n>\n> After\n\nOutside'
    root = tree(source)
    nodes = root.xpath('./blockquote/pre/code')
    assert len(nodes) == 1
    assert nodes[0].text == 'alpha\nbeta\n'
    assert root.xpath('./p')[-1].text == 'Outside'


def test_reusable_sanitization_keeps_fenced_examples_literal():
    from reusable_markup import sanitize_markdown
    source = fenced_list(payload='![literal](img/private.png)\n[private](file:///tmp/example)')
    projected = sanitize_markdown(source)
    assert isinstance(projected, str)
    assert tree(projected).xpath('.//pre/code')[0].text == '![literal](img/private.png)\n[private](file:///tmp/example)\n'


def test_reusable_projection_removes_actual_resources_but_preserves_code():
    from reusable_markup import sanitize_markdown
    payload = '![literal](img/private.png)\n[private](file:///tmp/example)'
    source = fenced_list(payload=payload) + '\n\n![ACTUAL](img/private.png)\n\n[PRIVATE](file:///tmp/example)'
    root = tree(sanitize_markdown(source))
    assert root.xpath('.//pre/code')[0].text == payload + '\n'
    assert not root.xpath('.//img|.//a[@href]')
    assert 'PRIVATE' in root.text_content()


def test_image_edit_changes_only_prose_destination_not_fenced_literal():
    payload = '![literal](img/private.png)'
    source = fenced_list(payload=payload) + '\n\n![ACTUAL](img/private.png)'
    changed = replace_image_destinations(source, 'img/private.png', 'img/changed.png')
    root = tree(changed)
    assert root.xpath('.//pre/code')[0].text == payload + '\n'
    assert root.xpath('.//img/@src') == ['img/changed.png']


EMPTY_AND_WHITESPACE_LINES = [
    pytest.param([], id='zero-lines'),
    pytest.param([''], id='one-empty-line'),
    pytest.param(['', '', ''], id='three-empty-lines'),
    pytest.param(['   '], id='spaces-only'),
    pytest.param(['\t'], id='tab-only'),
    pytest.param([' \t  '], id='mixed-whitespace'),
    pytest.param(['\t', '  ', ''], id='whitespace-and-final-empty-line'),
    pytest.param(['', '\t ', '', '  ', ''], id='interleaved-empty-lines'),
]


def whitespace_fence(lines, context):
    source_lines = ['```text', *lines, '```']
    if context == 'list':
        return '713. BEGIN\n\n' + '\n'.join('    ' + line for line in source_lines) + '\n\n714. NEXT'
    if context == 'quote':
        return '\n'.join('> ' + line for line in source_lines) + '\n\nOUTSIDE'
    return '\n'.join(source_lines) + '\n\nOUTSIDE'


@pytest.mark.parametrize('context', ['root', 'list', 'quote'])
@pytest.mark.parametrize('lines', EMPTY_AND_WHITESPACE_LINES)
def test_empty_and_whitespace_fences_preserve_exact_line_count(context, lines):
    root = tree(whitespace_fence(lines, context))
    code = root.xpath('.//pre/code')
    assert len(code) == 1
    expected = ''.join(line + '\n' for line in lines)
    assert code[0].text_content() == expected
    assert code[0].get('class') == 'language-text'
    if context == 'list':
        assert len(code[0].xpath('ancestor::li')) == 1
        assert root.xpath('./ol/li')[-1].text_content().strip() == 'NEXT'
    elif context == 'quote':
        assert len(code[0].xpath('ancestor::blockquote')) == 1
        assert root.xpath('./p')[-1].text == 'OUTSIDE'
    else:
        assert not code[0].xpath('ancestor::li|ancestor::blockquote')
        assert root.xpath('./p')[-1].text == 'OUTSIDE'


@pytest.mark.parametrize('lines', EMPTY_AND_WHITESPACE_LINES)
def test_docx_whitespace_only_fence_keeps_payload_after_one_terminal_newline(tmp_path, lines):
    document = Document()
    calls = []

    def code_block(text):
        calls.append(text)
        return document.add_paragraph(text)

    write_markdown(document, whitespace_fence(lines, 'list'), tmp_path, code_block)
    # Word removes the parser's terminal line separator only. Authored empty
    # lines inside the block and whitespace characters must remain intact.
    assert calls == [''.join(line + '\n' for line in lines).removesuffix('\n')]
    assert document.paragraphs[-1].text == 'NEXT'

