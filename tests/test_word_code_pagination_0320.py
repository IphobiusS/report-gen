"""Word may split oversized transcripts without losing a heading or code line."""
from docx import Document
import pytest

import engine
import export
from design_library import profile
from design_samples import sample_data


@pytest.mark.parametrize('ident', ['enterprise-assurance', 'centri-btl2'])
@pytest.mark.parametrize('transcript,keep', [
    ('find ./evidence -type f -print\nsha256sum ./evidence/session.txt', True),
    ('\n'.join(f'LOG-{i:03d} synthetic-evidence-{i:03d} VERIFIED' for i in range(90)), False),
    ('\n'.join(f'LONG-{i:03d} ' + 'bounded-synthetic-token ' * 50 for i in range(20)), False),
])
def test_code_can_split_if_hard_or_wrapped_lines_exceed_page(tmp_path, ident, transcript, keep):
    data = sample_data(profile(ident), 'quiet', 'en')
    data['report']['sections'] = [{'key': 'transcript', 'title': 'Evidence',
        'body': '## Collection transcript\n\n```text\n' + transcript + '\n```\n\nEND-OF-TRANSCRIPT'}]
    target = tmp_path / 'report.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, target)
    doc = Document(target)
    code = next(p for p in doc.paragraphs if p.style.name == 'ReportGenCode')
    assert code.text.rstrip('\n') == transcript
    assert code.paragraph_format.keep_together is keep
    heading = next(p for p in doc.paragraphs if p.text == 'Collection transcript')
    assert heading.style.paragraph_format.keep_with_next or heading.paragraph_format.keep_with_next
    assert any(p.text == 'END-OF-TRANSCRIPT' for p in doc.paragraphs)
