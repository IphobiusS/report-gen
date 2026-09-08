"""Regression coverage for the September 2026 audit. No security bypass in TESTING."""
import copy
import io
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import zipfile

import pytest
from lxml import html
import app as webapp
import engine
import export
import validate
import storage
from wizard import dump_yaml
from _support import auth_client

VECTOR = "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H"
ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(webapp, "PROJECTS", tmp_path)
    webapp.app.config["TESTING"] = True
    c = auth_client(webapp.app)
    assert c.post("/api/projects", json={"slug": "review", "title": "Original"}).status_code == 200
    return c


def test_markdown_attributes_and_urls_are_sanitized():
    for payload in ['![x" onerror="document.documentElement.dataset.audit=1](missing.png)', '[click](javascript:alert%281%29)', '[click](file:///etc/passwd)', '![x](https://example.invalid/x.png)', '![x](../outside.png)', '![x](img/a.png){width="1px;position:fixed"}']:
        tree = html.fragment_fromstring(str(engine.md(payload, img_base="")), create_parent=True)
        assert not tree.xpath('//@onerror|//@onclick')
        for node in tree.xpath('.//*[@href or @src]'):
            assert not (node.get("href") or node.get("src")).lower().startswith(("javascript:", "file:", "https://example.invalid"))
        assert not tree.xpath('.//*[contains(@style,"position")]')


def test_raw_payload_and_code_remain_literal():
    result = str(engine.md('Payload: <script>alert(1)</script>\n\n```\n<img onerror="a()">\n```\n\n`![literal](img/x.png)`', img_base=""))
    tree = html.fragment_fromstring(result, create_parent=True)
    assert '<script>alert(1)</script>' in tree.text_content()
    assert '<img onerror="a()">' in tree.text_content()
    assert not tree.xpath('.//script|.//img')


def test_markdown_bases_are_request_local():
    def render(i):
        base = f"/api/projects/p{i}/"
        return base, str(engine.md('![x](img/x.png)', img_base=base))
    with ThreadPoolExecutor(max_workers=8) as pool:
        for base, value in pool.map(render, range(32)):
            assert f'src="{base}img/x.png"' in value


def test_csrf_origin_host_and_content_type(client):
    raw = webapp.app.test_client()
    assert raw.post('/api/projects', json={}).status_code == 403
    assert client.post('/api/projects', json={}, headers={'Origin':'https://example.invalid'}).status_code == 403
    assert client.get('/api/projects', headers={'Host':'example.invalid'}).status_code == 400
    assert client.post('/api/projects', data='{}', content_type='text/plain').status_code == 415
    assert client.get('/api/session').headers['Cache-Control'] == 'no-store'
    assert "script-src 'self'" in client.get('/').headers['Content-Security-Policy']


@pytest.mark.parametrize('data', [{'meta': []}, {'meta':{},'findings':[None]}, {'findings':[{'id':'F1','phases':[{'steps':[1]}]}]}, {'report':{'sections':[None]}}, {'meta':{'contacts':{'client':[None]}}}])
def test_invalid_structure_does_not_replace_project(client, data):
    before = client.get('/api/projects/review').get_json()
    assert client.put('/api/projects/review', json=data).status_code == 422
    assert client.get('/api/projects/review').get_json() == before


def test_compare_and_swap_backup_and_restore(client):
    before = client.get('/api/projects/review')
    data = before.get_json(); data['meta']['report_title'] = 'Updated'
    assert client.put('/api/projects/review', json=data, headers={'If-Match':before.headers['ETag']}).status_code == 200
    stale = copy.deepcopy(data); stale['meta']['report_title'] = 'Stale tab'
    assert client.put('/api/projects/review', json=stale, headers={'If-Match':before.headers['ETag']}).status_code == 409
    assert client.get('/api/projects/review').get_json()['meta']['report_title'] == 'Updated'
    latest = client.get('/api/projects/review')
    assert client.post('/api/projects/review/restore', json={}, headers={'If-Match':latest.headers['ETag']}).status_code == 200
    assert client.get('/api/projects/review').get_json()['meta']['report_title'] == 'Original'


def test_atomic_write_preserves_previous_on_error(tmp_path, monkeypatch):
    target = tmp_path/'engagement.yaml'; dump_yaml({'meta':{'report_title':'Original'}}, target)
    before = target.read_bytes()
    def fail_replace(*args): raise OSError('simulated disk failure')
    monkeypatch.setattr(storage.os, 'replace', fail_replace)
    with pytest.raises(OSError): dump_yaml({'meta':{'report_title':'Lost'}}, target)
    assert target.read_bytes() == before


@pytest.mark.parametrize('score,severity', [(1,'low'),(-1,'low'),(11,'critical'),('NaN','critical'),('inf','critical')])
def test_score_and_vector_disagreement_is_error(score, severity):
    data={'meta':{'report_title':'T'},'findings':[{'id':'F1','title':'T','severity':severity,'cvss':score,'cvss_vector':VECTOR}]}
    assert any(level == 'error' for level,_ in validate.validate(data))


def test_zero_and_invalid_language():
    data={'meta':{'report_title':'T'},'findings':[{'id':'F1','title':'T','severity':'info','cvss':0}]}
    assert validate.validate(data) == []
    data['meta']['lang']='fr'
    assert any(level == 'error' for level,_ in validate.validate(data))


def test_export_rejects_invalid_score_without_saving(client):
    before = client.get('/api/projects/review')
    data=before.get_json(); data['findings']=[{'id':'F1','title':'T','severity':'low','cvss':'1.0','cvss_vector':VECTOR}]
    for fmt in ['md','pdf','docx']:
        assert client.post('/api/projects/review/export/'+fmt,json=data).status_code == 422
        assert client.get('/api/projects/review').headers['ETag'] == before.headers['ETag']


def test_markdown_download_is_self_contained_snapshot(client):
    from PIL import Image
    stream=io.BytesIO(); Image.new('RGB',(20,20),'red').save(stream,'PNG');stream.seek(0)
    upload=client.post('/api/projects/review/image',data={'file':(stream,'evidence.png')},content_type='multipart/form-data').get_json()['src']
    before=client.get('/api/projects/review')
    data=before.get_json();data['report']['sections']=[{'key':'executive_summary','summary':f'![Proof]({upload})'}]
    response=client.post('/api/projects/review/export/md',json=data)
    assert response.status_code == 200 and response.mimetype == 'application/zip'
    with zipfile.ZipFile(io.BytesIO(response.data)) as package:
        assert upload in package.namelist()
        assert f']({upload})' in package.read('report.md').decode()
    assert client.get('/api/projects/review').headers['ETag'] == before.headers['ETag']


def test_docx_markdown_preserves_evidence_and_code(tmp_path):
    from docx import Document
    source=ROOT/'reports/canon'
    text='![Evidence](img/fig1.png)\n\n```bash\n  first\n    second\n```\n\n|A|B|\n|---|---|\n|1|2|\n\n[Reference](https://example.com/)'
    data={'meta':{'lang':'es','report_title':'T','theme':'corporativo'},'report':{'sections':[{'key':'executive_summary','summary':text}]},'findings':[]}
    out=tmp_path/'rich.docx';export.to_docx(data,data['meta'],engine.load_lang('es'),source,out)
    doc=Document(out)
    assert len(doc.inline_shapes)==1
    assert any(p.text=='  first\n    second' for p in doc.paragraphs)
    assert any(table.cell(0,0).text=='A' for table in doc.tables)
    assert any(rel.reltype.endswith('/hyperlink') for rel in doc.part.rels.values())


def test_resource_policy(tmp_path):
    from resources import local_image, restricted_fetcher
    with pytest.raises(ValueError): local_image('../secret.png', tmp_path)
    with pytest.raises(ValueError): restricted_fetcher(tmp_path, ROOT/'themes')('https://example.invalid/pixel.png')
    with pytest.raises(ValueError): restricted_fetcher(tmp_path, ROOT/'themes')((ROOT/'README.md').as_uri())


def test_actual_pdf_has_expected_content_and_explicit_chart(tmp_path):
    import subprocess, sys
    import pdfplumber
    out=tmp_path/'report.pdf'
    subprocess.run([sys.executable,str(ROOT/'engine.py'),str(ROOT/'reports/canon/engagement.yaml'),'-o',str(out),'--html'],check=True,capture_output=True,timeout=120)
    with pdfplumber.open(out) as pdf:
        assert len(pdf.pages) >= 5
        text='\n'.join(page.extract_text() or '' for page in pdf.pages)
        assert 'PhantomKernel' in text and 'Informativa' in text
    source=out.with_suffix('.html').read_text()
    assert 'fill="#b3001b"' in source and 'text-anchor="middle"' in source


@pytest.mark.parametrize('vector', ["CVSS:3.1/AV:N", VECTOR.replace('/S:U/', '/S:Z/'), VECTOR + '/AV:N'])
def test_invalid_cvss31_vector_rejected(client, vector):
    assert client.post('/api/cvss', json={'vector': vector}).status_code == 400


def test_zero_score_is_preserved_in_all_exports(tmp_path):
    from docx import Document
    data = {'meta': {'report_title': 'Zero'}, 'findings': [{'id': 'F1', 'title': 'Information', 'severity': 'info', 'cvss': 0}], 'report': {'sections': [{'key': 'findings', 'enabled': True}]}}
    path = tmp_path / 'engagement.yaml'; dump_yaml(data, path)
    data, meta, labels = export.prepare(path)
    tree = html.fromstring(engine.render_html(engine.build_env(), data, tmp_path, {}, False, labels))
    assert any('0' in n.text_content() for n in tree.xpath('//span[contains(@class,"sev-info")]'))
    assert '(0)' in export.to_markdown(data, meta, labels, tmp_path)
    doc = Document(export.export(path, 'docx', tmp_path / 'zero.docx'))
    assert any(cell.text.strip() == '0' for table in doc.tables for row in table.rows for cell in row.cells)
