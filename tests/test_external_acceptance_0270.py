"""The external acceptance kit contains real documents, never fabricated observations."""
import csv
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import zipfile

from lxml import etree
from docx import Document
from docx.oxml.ns import qn
import pytest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("prepare_acceptance", ROOT / "scripts/prepare-external-acceptance.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def test_prepared_documents_and_blank_sessions_remain_unaccepted(tmp_path):
    result = module.prepare(tmp_path / "kit")
    assert result["word_native"] == result["human_sessions"] == "not_run"
    assert result["human_observations_recorded"] == 0
    assert result["prepared_docx_content"] == "verified"
    assert {row["case"] for row in result["documents"]} == {"quality-lab", "pending-severity", "nimbus", "long-content", "url-layout-es", "url-layout-en", "url-layout-pt"}
    for row in result["documents"]:
        path = tmp_path / "kit" / row["docx"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == row["docx_sha256"]
        assert row["content_checks"]["editable_document"] is True
        assert row["content_checks"]["archive_integrity"] is True
        with zipfile.ZipFile(path) as package:
            document = etree.fromstring(package.read("word/document.xml"))
            assert document.xpath('//*[local-name()="p"]')
    with zipfile.ZipFile(tmp_path / "kit/word-input/long-content.docx") as package:
        text = package.read("word/document.xml").decode()
        assert "ASSET-001" in text and "ASSET-080" in text
        assert "COMANDO_LITERAL_WORD_0270" in text
    with (tmp_path / "kit/human-sessions.csv").open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 15
    assert all(not row["outcome"] and not row["seconds"] for row in rows)
    assert all(row["commit"] == result["commit"] for row in rows)
    summary = module.summarize(tmp_path / "kit")
    assert summary["word_native"] == summary["word_visual"] == summary["human_sessions"] == "not_run"
    assert summary["word_acceptance"] == "pending"
    assert summary["usability_acceptance"] == "not_assessed"
    assert summary["prepared_docx_content"] == "verified"


def test_preparation_preserves_existing_evidence(tmp_path):
    marker = tmp_path / "existing.json"
    marker.write_text('{"human_observations_recorded": 1}')
    with pytest.raises(ValueError, match="empty output"):
        module.prepare(tmp_path)
    assert marker.read_text() == '{"human_observations_recorded": 1}'


def test_complex_trilingual_inputs_are_copied_and_exported_without_claiming_word(tmp_path):
    source = tmp_path / 'projects'
    before = {}
    for language in ('es', 'en', 'pt'):
        folder = source / language
        folder.mkdir(parents=True)
        module.dump_yaml({'meta': {'report_title': f'SYNTHETIC-{language}', 'lang': language},
                          'report': {'sections': [{'key': 'findings'}]},
                          'findings': [{'id': 'F001', 'mode': 'vuln', 'title': f'COMPLEX-{language}',
                                        'description_md': 'Synthetic input, not a Word observation.'}]},
                         folder / 'engagement.yaml')
        before[language] = (folder / 'engagement.yaml').read_bytes()
    output = tmp_path / 'kit'
    result = module.prepare(output, source)
    assert len(result['documents']) == 10
    for language, original in before.items():
        assert (source / language / 'engagement.yaml').read_bytes() == original
        target = output / f'word-input/complex-{language}.docx'
        document = Document(target)
        assert any(f'COMPLEX-{language}' in paragraph.text for paragraph in document.paragraphs)
    assert module.summarize(output)['word_acceptance'] == 'pending'


@pytest.fixture(scope="module")
def blank_kit(tmp_path_factory):
    path = tmp_path_factory.mktemp("acceptance") / "kit"
    module.prepare(path)
    return path


@pytest.fixture
def kit(blank_kit, tmp_path):
    return Path(shutil.copytree(blank_kit, tmp_path / "kit"))


def edit_rows(path, mutate):
    with path.open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        fields, rows = reader.fieldnames, list(reader)
    mutate(rows)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fields)
        writer.writeheader()
        writer.writerows(rows)


def observed(row, outcome="independent"):
    # Synthetic records exercise the validator; they are not real usability evidence.
    row.update(os_browser_viewer="TEST FIXTURE", seconds="120", setup_seconds="0", interventions="0",
               outcome=outcome, observer="TEST FIXTURE", observed_at="2026-10-10T12:00:00Z")


def test_complete_human_record_is_not_automatic_usability_approval(kit):
    edit_rows(kit / "human-sessions.csv", lambda rows: [observed(row) for row in rows])
    result = module.summarize(kit)
    assert result["participants_complete"] == 5
    assert result["human_observations_recorded"] == 15
    assert result["human_sessions"] == "recorded"
    assert result["usability_acceptance"] == "not_assessed"
    assert result["word_native"] == "not_run"


@pytest.mark.parametrize("change", [
    {"commit": "0" * 40}, {"seconds": "NaN"}, {"interventions": "1"}, {"interventions": "0.5"},
    {"outcome": "blocked"}, {"observer": ""}, {"outcome": "assisted"}, {"seconds": "-1"},
])
def test_incomplete_or_contradictory_observations_are_rejected(kit, change):
    def mutate(rows):
        observed(rows[0])
        rows[0].update(change)
    edit_rows(kit / "human-sessions.csv", mutate)
    with pytest.raises(ValueError):
        module.summarize(kit)


def test_duplicate_human_task_cannot_count_as_another_participant(kit):
    edit_rows(kit / "human-sessions.csv", lambda rows: rows.__setitem__(1, rows[0]))
    with pytest.raises(ValueError, match="fifteen"):
        module.summarize(kit)


def test_visual_approval_without_native_word_evidence_is_rejected(kit):
    def mutate(rows):
        rows[0].update(reviewer="TEST FIXTURE", reviewed_at="2026-10-10", pdf_sha256="0" * 64,
                       **dict.fromkeys(module.VISUAL_CHECKS, "pass"))
    edit_rows(kit / "word-review.csv", mutate)
    with pytest.raises(ValueError, match="Visual review"):
        module.summarize(kit)


def native_fixture(kit):
    """Fake bytes for consistency-test coverage, never a native Word run."""
    manifest = json.loads((kit / "acceptance-preparation.json").read_text())
    folder = kit / "word-verified"
    folder.mkdir()
    rows = []
    for case in manifest["documents"]:
        pdf, docx = f'{case["case"]}.pdf', f'{case["case"]}.docx'
        (folder / pdf).write_bytes(b"TEST FIXTURE, NOT A WORD RENDER")
        (folder / docx).write_bytes(b"TEST FIXTURE, NOT A WORD DOCUMENT")
        rows.append({"file": f'{case["case"]}.docx', "word_version": "TEST FIXTURE", "pages": 1,
                     "status": "fields_updated_and_rendered", "source_sha256": case["docx_sha256"],
                     "pdf": pdf, "pdf_sha256": module.sha256(folder / pdf),
                     "updated_docx": docx, "updated_sha256": module.sha256(folder / docx)})
    (folder / "word-validation.json").write_text(json.dumps({"results": rows}))
    return rows


def test_native_automation_still_requires_visual_review(kit):
    rows = native_fixture(kit)
    summary = module.summarize(kit)
    assert summary["word_native"] == "completed"
    assert summary["word_acceptance"] == "pending"
    assert summary["word_visual"] == "not_run"
    (kit / "word-verified" / rows[0]["pdf"]).write_bytes(b"changed")
    with pytest.raises(ValueError, match="does not match"):
        module.summarize(kit)


def test_visual_failure_remains_an_issue(kit):
    native = native_fixture(kit)
    def mutate(rows):
        for row in rows:
            row.update(reviewer="TEST FIXTURE", reviewed_at="2026-10-10", pdf_sha256=native[0]["pdf_sha256"],
                       **dict.fromkeys(module.VISUAL_CHECKS, "pass"))
        rows[0].update(tables_images="fail", notes="TEST FIXTURE: clipped synthetic table")
    edit_rows(kit / "word-review.csv", mutate)
    summary = module.summarize(kit)
    assert summary["word_visual"] == "issues_found"
    assert summary["word_acceptance"] == "pending"
    assert summary["human_sessions"] == "not_run"


def test_changed_input_rejected(kit):
    (kit / "word-input/quality-lab.docx").write_bytes(b"changed")
    with pytest.raises(ValueError, match="input changed"):
        module.summarize(kit)


def test_long_markdown_table_flows_by_rows_and_preserves_header_and_command(blank_kit):
    document = Document(blank_kit / "word-input/long-content.docx")
    table = next(table for table in document.tables if table.cell(0, 0).text == "Activo")
    assert len(table.rows) == 81
    assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
    for index, row in enumerate(table.rows):
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        assert all(p.paragraph_format.keep_together for cell in row.cells for p in cell.paragraphs)
        assert all(bool(p.paragraph_format.keep_with_next) == (index == 0 or len(table.rows)-3 <= index < len(table.rows)-1)
                   for cell in row.cells for p in cell.paragraphs)
    command = next(p for p in document.paragraphs if "COMANDO_LITERAL_WORD_0270" in p.text)
    assert command.text == "COMANDO_LITERAL_WORD_0270 --documento ejemplo --salida local"
    assert command.paragraph_format.keep_together
    assert not command.paragraph_format.keep_with_next


@pytest.mark.parametrize('lang', ['es', 'en', 'pt'])
def test_literal_url_and_wide_ordinal_kit_is_an_actual_unaccepted_document(blank_kit, lang):
    document = Document(blank_kit / f'word-input/url-layout-{lang}.docx')
    targets = {rel.rId: rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith('/hyperlink')}
    links = [(link.get(qn('r:id')), ''.join(link.xpath('.//w:t/text()')))
             for link in document.element.xpath('//w:hyperlink[@r:id]')]
    assert [text for rid, text in links if targets[rid] == module.LITERAL_URL] == [module.LITERAL_URL] * 3
    assert [text for rid, text in links if targets[rid] == module.ENCODED_URL] == [module.ENCODED_URL]
    fields = document.element.xpath('//w:fldSimple')
    assert [''.join(node.text for node in field.findall('.//' + qn('w:t'))) for field in fields if field.get(qn('w:instr')).strip().startswith('SEQ ')] == ['2147483646', '2147483647']
    expectation = json.loads((blank_kit / 'URL-LAYOUT-EXPECTATIONS.json').read_text())
    assert expectation['literal_url'] == module.LITERAL_URL
    assert expectation['encoded_url'] == module.ENCODED_URL
    assert expectation['maximum_nonfinal_unused_width_points'] == 10.5
    assert expectation['status'] == 'pending_native_render_and_human_measurement'
    assert module.summarize(blank_kit)['word_acceptance'] == 'pending'


def test_native_word_guard_functions_with_actual_sequence_export_and_mock_links(blank_kit, tmp_path):
    """Run real PowerShell guards; mock links are not a Microsoft Word session."""
    executable = os.environ.get('REPORT_GEN_PWSH') or shutil.which('pwsh') or shutil.which('powershell')
    if not executable:
        pytest.skip('PowerShell is required for native-acceptance guard tests; Word is not required')
    with zipfile.ZipFile(blank_kit / 'word-input/url-layout-en.docx') as archive:
        xml = tmp_path / 'document.xml'
        xml.write_bytes(archive.read('word/document.xml'))
    check = tmp_path / 'guards.ps1'
    check.write_text(r'''
param([string]$Source, [string]$DocumentXml)
$ErrorActionPreference = 'Stop'
$tokens = $null; $parseErrors = $null
$ast = [System.Management.Automation.Language.Parser]::ParseFile($Source, [ref]$tokens, [ref]$parseErrors)
if ($parseErrors.Count) { throw 'PowerShell parser rejected the native verification script.' }
$wanted = @('Test-ReportGenWordFieldInstruction', 'Assert-ReportGenWordFields', 'Assert-ReportGenWordInternalLinks', 'Get-ReportGenWordContentsMetrics')
$functions = @($ast.FindAll({ param($node) $node -is [System.Management.Automation.Language.FunctionDefinitionAst] }, $false))
foreach ($name in $wanted) {
    $definition = @($functions | Where-Object { $_.Name -eq $name })
    if ($definition.Count -ne 1) { throw "Missing guard: $name" }
    Invoke-Expression $definition[0].Extent.Text
}
$xml = New-Object System.Xml.XmlDocument
$xml.XmlResolver = $null
$xml.Load($DocumentXml)
$ns = New-Object System.Xml.XmlNamespaceManager($xml.NameTable)
$ns.AddNamespace('w', 'http://schemas.openxmlformats.org/wordprocessingml/2006/main')
Assert-ReportGenWordFields $xml $ns
foreach ($instruction in @(' SEQ RGList10 \r 65536 \* ARABIC ', 'SEQ RGList9 \r 2147483647 \* ARABIC', 'SEQ RGList10 \n \* ARABIC')) {
    if (-not (Test-ReportGenWordFieldInstruction $instruction)) { throw "Valid sequence rejected: $instruction" }
}
foreach ($instruction in @('SEQ Arbitrary \n \* ARABIC', 'SEQ RGList10 \r 65535 \* ARABIC', 'SEQ RGList10 \r 2147483648 \* ARABIC',
    'SEQ RGList10 \r 9999999999999999999999999999 \* ARABIC', 'SEQ RGList0 \n \* ARABIC', 'SEQ RGList10 \n \* ROMAN',
    'SEQ RGList10 \n \* ARABIC DDEAUTO', 'SEQ RGList10 \c \* ARABIC', 'SEQ RGList10 \n \h \* ARABIC',
    'DDEAUTO app topic', 'INCLUDETEXT "remote.docx"', 'LINK Excel.Sheet.12 source')) {
    if (Test-ReportGenWordFieldInstruction $instruction) { throw "Unsupported field accepted: $instruction" }
}
$sequence = @($xml.SelectNodes('//w:fldSimple', $ns) | Where-Object { $_.GetAttribute('instr', $ns.LookupNamespace('w')).Trim().StartsWith('SEQ ') })
$saved = $sequence[1].SelectSingleNode('.//w:t', $ns).InnerText
$sequence[1].SelectSingleNode('.//w:t', $ns).InnerText = '2147483646'
$rejected = $false
try { Assert-ReportGenWordFields $xml $ns } catch { $rejected = $true }
if (-not $rejected) { throw 'Changed cached ordinal was accepted.' }
$sequence[1].SelectSingleNode('.//w:t', $ns).InnerText = $saved
$sequence[0].ParentNode.RemoveChild($sequence[0]) | Out-Null
$rejected = $false
try { Assert-ReportGenWordFields $xml $ns } catch { $rejected = $true }
if (-not $rejected) { throw 'Orphan sequence continuation was accepted.' }
$bookmarks = [pscustomobject]@{ calls = 0 }
$bookmarks | Add-Member ScriptMethod Exists { param($name) $this.calls++; return $name -eq 'present' }
$document = [pscustomobject]@{ Hyperlinks = @([pscustomobject]@{Address='https://example.test/report'; SubAddress='external-fragment'}); Bookmarks=$bookmarks }
Assert-ReportGenWordInternalLinks $document
if ($bookmarks.calls -ne 0) { throw 'External fragment incorrectly resolved as a local bookmark.' }
$document.Hyperlinks = @([pscustomobject]@{Address=''; SubAddress='present'})
Assert-ReportGenWordInternalLinks $document
if ($bookmarks.calls -ne 1) { throw 'Local bookmark was not checked.' }
$document.Hyperlinks = @([pscustomobject]@{Address=''; SubAddress='missing'})
$rejected = $false
try { Assert-ReportGenWordInternalLinks $document } catch { $rejected = $true }
if (-not $rejected) { throw 'Missing local bookmark was accepted.' }
$document = [pscustomobject]@{
    Fields = @(
        [pscustomobject]@{Code=[pscustomobject]@{Text=' PAGEREF rg_toc_1 \h '}},
        [pscustomobject]@{Code=[pscustomobject]@{Text=' PAGEREF "rg_toc_2" \h '}},
        [pscustomobject]@{Code=[pscustomobject]@{Text=' PAGE '}}
    )
    TablesOfContents = [pscustomobject]@{Count=0}
}
$metrics = Get-ReportGenWordContentsMetrics $document @('rg_toc_1', 'rg_toc_2')
if ($metrics.native_toc_count -ne 0 -or $metrics.contents_entry_count -ne 2 -or $metrics.page_reference_field_count -ne 2 -or $metrics.contents_page_reference_count -ne 2) {
    throw 'Native TOC objects were confused with live contents entries.'
}
$rejected = $false
try { Get-ReportGenWordContentsMetrics $document @('rg_toc_1', 'rg_toc_missing') } catch { $rejected = $true }
if (-not $rejected) { throw 'Missing native page-reference field was accepted.' }
$document.Fields += [pscustomobject]@{Code=[pscustomobject]@{Text=' PAGEREF rg_toc_2 \h '}}
$rejected = $false
try { Get-ReportGenWordContentsMetrics $document @('rg_toc_1', 'rg_toc_2') } catch { $rejected = $true }
if (-not $rejected) { throw 'Duplicate native page-reference field was accepted.' }
@{guard_checks='passed'; native_word_tested=$false} | ConvertTo-Json -Compress
''', encoding='utf-8')
    result = subprocess.run([executable, '-NoProfile', '-NonInteractive', '-File', str(check),
                             '-Source', str(ROOT / 'scripts/verify-microsoft-word.ps1'),
                             '-DocumentXml', str(xml)], capture_output=True, text=True, timeout=45)
    assert result.returncode == 0, result.stdout + result.stderr
    assert json.loads(result.stdout) == {'guard_checks': 'passed', 'native_word_tested': False}
