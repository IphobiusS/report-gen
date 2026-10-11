"""Long cover identities retain their own alignment when Word wraps text."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pytest

import engine
import export
from design_library import profile
from design_samples import sample_data


@pytest.mark.parametrize('ident', ['enterprise-assurance', 'centri-btl2'])
@pytest.mark.parametrize('wordmark,classification', [
    ('AURORA INTERNATIONAL BANKING AND TREASURY GROUP',
     'CONFIDENTIAL - RESTRICTED TO ENGAGEMENT STAKEHOLDERS'),
    ('Aurora', 'CONFIDENTIAL\nRestricted to the investigation team'),
    ('AURORA\nINTERNATIONAL BANKING', 'CONFIDENTIAL'),
])
def test_long_cover_classification_cannot_continue_beneath_identity(tmp_path, ident, wordmark, classification):
    data = sample_data(profile(ident), 'brand', 'en')
    data['meta']['branding'].update(wordmark=wordmark, confidential_text=classification)
    target = tmp_path / 'report.docx'
    export.to_docx(data, data['meta'], engine.load_lang('en'), tmp_path, target)
    document = Document(target)
    header = document.tables[0].rows[0].cells[-1]
    # Both strings must survive in isolated wrapping regions: a shared tabbed
    # paragraph incorrectly sends classification continuation lines to the left.
    identity = header.tables[0]
    left, right = identity.rows[0].cells
    assert left.text == wordmark
    assert right.text == classification
    assert right.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert left.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    cover_text = ''.join(document.tables[0]._tbl.xpath('.//w:t/text()'))
    for key in ('report_title', 'report_subtitle', 'client', 'email', 'date', 'version'):
        assert str(data['meta'][key]) in cover_text
    # The lower grid still owns all metadata; the new top grid cannot replace it.
    bottom = document.tables[0].rows[-1].cells[-1]
    assert data['meta']['email'] in ''.join(bottom.tables[0]._tbl.xpath('.//w:t/text()'))
