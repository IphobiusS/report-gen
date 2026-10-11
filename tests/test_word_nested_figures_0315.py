"""Figures retain paragraph boundaries through inline formatting and links."""
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import pytest

from exporters.richtext import write_markdown


@pytest.mark.parametrize('wrapper', ['link', 'bold', 'bold_link', 'emphasis'])
@pytest.mark.parametrize('context', ['paragraph', 'quote', 'tight_list', 'loose_list'])
def test_wrapped_figure_is_a_block_with_caption_and_surrounding_prose(tmp_path, wrapper, context):
    Image.new('RGB', (1000, 1500), 'white').save(tmp_path / 'decision.png')
    image = '![Independent caption](decision.png){width=50mm}'
    if wrapper in {'link', 'bold_link'}:
        image = f'[{image}](https://example.test/decision)'
    if wrapper in {'bold', 'bold_link'}:
        image = '**' + image + '**'
    if wrapper == 'emphasis':
        image = '*' + image + '*'
    body = 'Before the figure ' + image + ' After the figure.'
    if context == 'quote': body = '> ' + body
    if context == 'tight_list': body = '1. ' + body + '\n2. Next case.'
    if context == 'loose_list': body = '1. ' + body + '\n\n2. Next case.'
    document = Document()
    write_markdown(document, body, tmp_path, lambda _: None)
    paragraphs = document.paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p._p.xpath('.//wp:inline'))
    assert paragraphs[index - 1].text == 'Before the figure '
    assert paragraphs[index].text == ''
    caption = paragraphs[index + 1]
    assert ''.join(caption._p.xpath('.//w:t/text()')) == 'Independent caption'
    assert not caption._p.xpath('.//w:vanish')
    assert caption.style.name == 'Caption'
    assert caption.paragraph_format.keep_together is True
    assert paragraphs[index].paragraph_format.keep_with_next is True
    assert paragraphs[index + 2].text == ' After the figure.'
    assert abs(document.inline_shapes[0].width.cm - 5) < .001
    assert abs(document.inline_shapes[0].height.cm - 7.5) < .001
    assert document.inline_shapes[0]._inline.docPr.get('descr') == 'Independent caption'
    if wrapper in {'link', 'bold_link'}:
        links = caption._p.xpath('./w:hyperlink')
        assert len(links) == 1
        relationship = document.part.rels[links[0].get(qn('r:id'))]
        assert relationship.target_ref == 'https://example.test/decision'
        graphic_links = document.inline_shapes[0]._inline.docPr.xpath('./a:hlinkClick')
        assert len(graphic_links) == 1
        assert graphic_links[0].get(qn('r:id')) == links[0].get(qn('r:id'))
    if wrapper in {'bold', 'bold_link'}:
        assert caption._p.xpath('.//w:r[w:t]/w:rPr/w:b[not(@w:val="0")]')
    if wrapper == 'emphasis':
        assert caption._p.xpath('.//w:r[w:t]/w:rPr/w:i[not(@w:val="0")]')
    if context.endswith('list'):
        assert paragraphs[0]._p.xpath('./w:pPr/w:numPr')
        assert paragraphs[-1].text == 'Next case.'


@pytest.mark.parametrize('linked', [False, True])
@pytest.mark.parametrize('loose', [False, True])
@pytest.mark.parametrize('marker', ['1.', '-'])
def test_list_starting_with_figure_keeps_ordinal_on_the_drawing(tmp_path, linked, loose, marker):
    Image.new('RGB', (1000, 1500), 'white').save(tmp_path / 'decision.png')
    image = '![Independent caption](decision.png){width=50mm}'
    if linked: image = f'[{image}](https://example.test/decision)'
    separator = '\n\n' if loose else '\n'
    next_marker = '2.' if marker == '1.' else '-'
    document = Document()
    write_markdown(document, f'{marker} {image}{separator}{next_marker} Next case.',
                   tmp_path, lambda _: None)
    paragraphs = document.paragraphs
    assert paragraphs[0]._p.xpath('./w:pPr/w:numPr')
    assert paragraphs[0]._p.xpath('.//wp:inline')
    assert paragraphs[0].paragraph_format.keep_with_next is True
    assert paragraphs[1].style.name == 'Caption'
    assert ''.join(paragraphs[1]._p.xpath('.//w:t/text()')) == 'Independent caption'
    assert paragraphs[1].paragraph_format.left_indent == paragraphs[0].paragraph_format.left_indent
    numbered = [p for p in paragraphs if p._p.xpath('./w:pPr/w:numPr')]
    assert len(numbered) == 2
    assert numbered[-1].text == 'Next case.'
    assert all(p.text or p._p.xpath('.//wp:inline') for p in numbered)
