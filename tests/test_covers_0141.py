"""Cover geometry, data preservation and native Word styling contracts."""
from copy import deepcopy
from pathlib import Path

import pdfplumber
from PIL import Image
import pytest
from docx import Document

import engine
import export
from report_designs import THEMES

ROOT=Path(__file__).resolve().parents[1]


def source(theme, artwork, directory):
    data=deepcopy(engine.load_engagement(ROOT/'examples/design-sample/engagement.yaml'))
    data['meta'].update(theme=theme,report_title='Evaluación de seguridad de aplicaciones y servicios de identidad',
                        report_subtitle='Entorno de laboratorio y análisis de controles de autorización',
                        client='Organización de demostración para servicios corporativos',
                        assessor='Sebastián Latorre Muñoz',assessor_title='Evaluador técnico',
                        osid='DEMO-123456',email='evaluador@organizacion.example')
    data['findings']=[]
    data['report']={'sections':[{'key':'scope','objectives':'CONTENIDO-DEL-CUERPO'}]}
    if artwork:
        Image.new('RGB',(700,200),'#294c65').save(directory/'art.png')
        data['meta']['branding'].update(logo='art.png',client_logo='art.png',cover_image='art.png')
    return data


@pytest.mark.parametrize('theme',sorted(THEMES))
@pytest.mark.parametrize('artwork',[False,True])
def test_cover_keeps_full_metadata_low_and_title_separate(tmp_path,theme,artwork):
    data=source(theme,artwork,tmp_path)
    html=engine.render_html(engine.build_env(),data,tmp_path,{},False,engine.load_lang('es'))
    out=tmp_path/'cover.pdf';engine.render_pdf_weasyprint(html,out,tmp_path)
    with pdfplumber.open(out) as pdf:
        assert len(pdf.pages)==3  # cover, generated contents page, scope
        # Read drawing order so a wrapped field is not interleaved with the
        # adjacent metadata column. Geometry is checked independently below.
        page=pdf.pages[0];text=''.join(page.extract_text(use_text_flow=True).split())
        for field in ['report_title','report_subtitle','client','assessor','osid','email','date','version']:
            assert ''.join(str(data['meta'][field]).split()) in text,(theme,field)
        assert 'CONTENIDO-DEL-CUERPO' not in page.extract_text()
        assert 'CONTENIDO-DEL-CUERPO' in pdf.pages[2].extract_text()
        words=page.extract_words(extra_attrs=['fontname','size'])
        title=[w for w in words if w['size']>=27]
        assert title and all('ReportCover' in w['fontname'].replace('-','') for w in title)
        dates=page.search(str(data['meta']['date']))
        assert dates and dates[0]['top']>page.height*.69
        client=page.search('Organización')[0]
        assert client['top']>page.height*.64
        assert max(w['bottom'] for w in title)+25<client['top']
        assert all(-1<=w['x0']<w['x1']<=page.width+1 and
                   -1<=w['top']<w['bottom']<=page.height+1 for w in words)
        assert len(page.images)==(3 if artwork else 0)



@pytest.mark.parametrize('theme',sorted(THEMES))
def test_word_cover_preserves_all_data_and_has_no_inherited_title_rule(tmp_path,theme):
    data=source(theme,True,tmp_path)
    path=tmp_path/'cover.docx'
    export.to_docx(data,data['meta'],engine.load_lang('es'),tmp_path,path)
    doc=Document(path)
    text=''.join(doc.tables[0]._tbl.xpath('.//w:t/text()'))
    for field in ['report_title','report_subtitle','client','assessor','osid','email','date','version']:
        assert str(data['meta'][field]) in text
    assert len(doc.inline_shapes)==3
    assert len(doc.tables[0].rows)==3
    assert not doc.styles['Title'].element.xpath('./w:pPr/w:pBdr')
    assert not doc.tables[0]._tbl.xpath('.//w:trHeight[@w:hRule="exact"]')
