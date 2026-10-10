"""Export regressions reproduced in the second, isolated report audit."""
from copy import deepcopy
import json

from docx import Document
from docx.oxml.ns import qn
from lxml import html
import pytest

import delivery
import engine
import export
from export_jobs import worker
from safe_markup import md
from wizard import dump_yaml


def sample():
    return {'meta': {'report_title': 'Synthetic audit', 'lang': 'es', 'theme': 'corporativo'},
            'report': {'sections': [{'key': 'findings_summary'}, {'key': 'findings'}]},
            'findings': [{'id': 'F01', 'mode': 'vuln', 'title': 'Synthetic finding', 'severity': 'low',
                          'description_md': 'Description', 'impact_md': 'Impact', 'remediation_md': 'Fix'}]}


def test_individual_html_worker_generates_downloadable_result(tmp_path):
    source = tmp_path / 'source'; source.mkdir()
    dump_yaml(sample(), source / 'engagement.yaml')
    (tmp_path / 'job.json').write_text(json.dumps({'kind': 'export', 'formats': ['html'],
        'captured': {'image_placements': 0, 'images': {}}}))
    assert worker(tmp_path) == 0
    assert (tmp_path / 'result').read_text().startswith('<!doctype html>')
    assert not (tmp_path / 'error.txt').exists()


@pytest.mark.parametrize('field', ['description_md', 'impact_md', 'remediation_md'])
@pytest.mark.parametrize('source,tag', [
    ('| Header | Value |\n| --- | --- |\n| Field | SYNTHETIC |', 'table'),
    ('```text\nSYNTHETIC\n```', 'pre'),
    ('1. SYNTHETIC\n2. Second item', 'ol'),
    ('- SYNTHETIC\n- Second item', 'ul'),
])
def test_markdown_field_start_preserves_block_structure(tmp_path, field, source, tag):
    data = sample(); data['findings'][0][field] = source
    rendered = export.to_markdown(data, data['meta'], engine.load_lang('es'), tmp_path)
    tree = html.fragment_fromstring(str(md(rendered)), create_parent='div')
    assert any('SYNTHETIC' in node.text_content() for node in tree.xpath('.//' + tag))


@pytest.mark.parametrize('kind', ['walkthrough', 'remediation_summary', 'machine_summary'])
def test_other_labelled_markdown_keeps_code_block(tmp_path, kind):
    data = sample(); finding = data['findings'][0]; code = '```text\nSYNTHETIC\n```'
    if kind == 'walkthrough':
        finding['walkthrough'] = [{'lead': 'Control', 'text_md': code}]
    elif kind == 'remediation_summary':
        finding['remediation_summary'] = {'short_md': code}
    else:
        finding.update(mode='machine', host={'name': 'Synthetic'}, summary_md=code)
    rendered = export.to_markdown(data, data['meta'], engine.load_lang('es'), tmp_path)
    tree = html.fragment_fromstring(str(md(rendered)), create_parent='div')
    assert any('SYNTHETIC' in node.text_content() for node in tree.xpath('.//pre'))


@pytest.mark.parametrize('fmt', ['html', 'pdf', 'docx', 'md'])
def test_duplicate_ids_rejected_before_replacing_any_export(tmp_path, fmt):
    data = sample(); data['findings'].append(deepcopy(data['findings'][0]))
    dump_yaml(data, tmp_path / 'engagement.yaml')
    destination = tmp_path / ('protected.' + ('zip' if fmt == 'md' else fmt))
    destination.write_bytes(b'previous export')
    with pytest.raises(ValueError, match='duplicad'):
        export.export(tmp_path / 'engagement.yaml', fmt, destination)
    assert destination.read_bytes() == b'previous export'


def test_word_complex_list_preserves_block_order_and_continuation_indent(tmp_path):
    data = sample()
    data['findings'][0]['description_md'] = (
        '1. FIRST_ITEM\n\n    CONTINUATION\n\n        CODE_IN_LIST\n\n'
        '    - NESTED_ITEM\n\n    AFTER_NESTED\n\n2. LAST_ITEM')
    path = tmp_path / 'list.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, path)
    paragraphs = Document(path).paragraphs
    markers = ['FIRST_ITEM', 'CONTINUATION', 'CODE_IN_LIST', 'NESTED_ITEM', 'AFTER_NESTED', 'LAST_ITEM']
    relevant = [p for p in paragraphs if any(m in p.text for m in markers)]
    assert [p.text.strip() for p in relevant] == markers
    first = relevant[0].paragraph_format.left_indent
    for p in (relevant[1], relevant[2], relevant[4]):
        assert p.paragraph_format.left_indent == first
        assert not p._p.xpath('./w:pPr/w:numPr')
    assert all(r.font.name == 'DejaVu Sans Mono' for r in relevant[2].runs if r.text)
    assert relevant[3].paragraph_format.left_indent > first
    assert relevant[-1]._p.xpath('./w:pPr/w:numPr/w:numId')[0].get(qn('w:val')) == relevant[0]._p.xpath('./w:pPr/w:numPr/w:numId')[0].get(qn('w:val'))


def test_word_list_preserves_spaces_between_inline_styles(tmp_path):
    data = sample()
    data['findings'][0]['description_md'] = '- **First** `second` **third**'
    path = tmp_path / 'inline-list.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, path)
    assert any(p.text == 'First second third' for p in Document(path).paragraphs)


@pytest.mark.parametrize('section_based', [True, False])
def test_word_in_document_links_target_existing_bookmarks(tmp_path, section_based):
    data = sample()
    target = 'findings_summary'
    if not section_based:
        del data['report']; target = 'sec-summary'
    data['findings'][0]['description_md'] = f'[Summary](#{target}) and [Finding](#F01). [External](https://example.com/docs).'
    path = tmp_path / 'links.docx'
    export.to_docx(data, data['meta'], engine.load_lang('es'), tmp_path, path)
    doc = Document(path)
    names = {n.get(qn('w:name')) for n in doc.element.xpath('//w:bookmarkStart')}
    links = doc.element.xpath('//w:hyperlink')
    for label in ('Summary', 'Finding'):
        link = next(n for n in links if ''.join(n.xpath('.//w:t/text()')) == label)
        assert link.get(qn('w:anchor')) in names
        assert link.get(qn('r:id')) is None
    external = [r.target_ref for r in doc.part.rels.values() if 'hyperlink' in r.reltype]
    assert external == ['https://example.com/docs']


@pytest.mark.parametrize('kind', ['external_image', 'missing_anchor', 'missing_image'])
def test_html_inspection_detects_broken_deliverable(tmp_path, kind):
    dump_yaml(sample(), tmp_path / 'engagement.yaml')
    path = export.export(tmp_path / 'engagement.yaml', 'html', tmp_path / 'report.html')
    tree = html.document_fromstring(path.read_text())
    expected = 0
    if kind == 'external_image':
        tree.find('body').append(html.Element('img', src='https://example.com/image.png'))
    elif kind == 'missing_anchor':
        tree.xpath('//a[starts-with(@href,"#")]')[0].set('href', '#no-such-section')
    else:
        expected = 1
    path.write_text(html.tostring(tree, encoding='unicode'))
    with pytest.raises(ValueError):
        delivery.inspect_html(path, expected)
