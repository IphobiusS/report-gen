"""Bounded local diagnostics. Invoked in a disposable subprocess by the editor."""
from languages import language_argument, language_index
import json
from pathlib import Path
import tempfile


def run(lang="es"):
    w = lambda es, en, pt: (es, en, pt)[language_index(lang)]
    results = []
    with tempfile.TemporaryDirectory(prefix='report-gen-doctor-') as directory:
        root = Path(directory)
        for name, label, action in [
            ('weasyprint', 'PDF / WeasyPrint', lambda: weasy(root)),
            ('chromium', w('PDF / Chromium (alternativa)','PDF / Chromium (alternative)', 'PDF / Chromium (alternativa)'), lambda: chromium(root)),
            ('docx', w('Word editable','Editable Word', 'Word editável'), lambda: word(root)),
        ]:
            try:
                action()
                results.append({'id': name, 'label': label, 'ready': True, 'required': name == 'docx',
                                'detail': w('Prueba de exportación correcta','Export test passed', 'Teste de exportação aprovado')})
            except Exception as exc:
                instructions = {'weasyprint': w('Motor PDF opcional para abrir el editor. Revisa las bibliotecas nativas de WeasyPrint en README.md.',
                                               'PDF engine is optional for opening the editor. Check native WeasyPrint libraries in README.en.md.', 'O mecanismo PDF é opcional para abrir o editor. Consulte as bibliotecas nativas do WeasyPrint em README.pt.md.'),
                                'chromium': w('Motor PDF alternativo opcional. En el entorno .venv, instala playwright y ejecuta: python -m playwright install chromium',
                                              'Optional alternative PDF engine. In the .venv environment, install playwright and run: python -m playwright install chromium', 'Mecanismo PDF alternativo opcional. No ambiente .venv, instale playwright e execute: python -m playwright install chromium'),
                                'docx': w('Dependencia necesaria. Ejecuta: python start.py --setup',
                                          'Required dependency. Run: python start.py --setup', 'Dependência necessária. Execute: python start.py --setup')}
                results.append({'id': name, 'label': label, 'ready': False, 'required': name == 'docx',
                                'detail': type(exc).__name__+'. '+instructions[name]})
    return {'checks': results, 'pdf_ready': any(r['ready'] for r in results[:2]),
            'editor_ready': results[2]['ready'],
            'note': w('PDF requiere al menos un motor disponible. Si ambos faltan, puedes editar y exportar Word, Markdown y HTML. Las fuentes de Word dependen del equipo que abra el DOCX.',
                      'PDF requires at least one available engine. If both are missing, you can edit and export Word, Markdown and HTML. Word fonts depend on the computer opening the DOCX.', 'O PDF requer pelo menos um mecanismo disponível. Se ambos estiverem ausentes, você pode editar e exportar Word, Markdown e HTML. As fontes do Word dependem do computador que abre o DOCX.')}


def weasy(root):
    from weasyprint import HTML
    from pypdf import PdfReader
    path = root/'test.pdf'
    HTML(string='<p>report-gen: ÁÉÍÓÚ ñ 0123456789</p>').write_pdf(path)
    assert 'report-gen' in PdfReader(path).pages[0].extract_text()


def chromium(root):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(timeout=8000)
        try:
            page = browser.new_page()
            page.set_content('<p>report-gen</p>')
            assert page.pdf(path=str(root/'chromium.pdf')).startswith(b'%PDF-')
        finally:
            browser.close()


def word(root):
    from docx import Document
    path = root/'test.docx'
    doc = Document(); doc.add_paragraph('report-gen'); doc.save(path)
    assert Document(path).paragraphs[0].text == 'report-gen'


if __name__ == '__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--lang',type=language_argument, choices=['es','en','pt'],default='es')
    print(json.dumps(run(parser.parse_args().lang), ensure_ascii=False))
