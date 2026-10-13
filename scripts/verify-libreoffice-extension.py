#!/usr/bin/env python3
"""Accept the optional OXT in a real, isolated, non-root Writer installation.

The UNO client never imports/registers the extension. Only unopkg installs it.
This is LibreOffice acceptance, not Microsoft Word or human visual acceptance.
"""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile

ROOT = Path(__file__).resolve().parents[1]
TAG = 'zxx-ZZ-x-rgwrap'
SERVICE = 'com.sun.star.i18n.BreakIterator_qlt_ZZ_zxx-ZZ-x-rgwrap'
IMPLEMENTATION = 'org.reportgen.TechnicalLiteralBreakIterator'
CONTROL_SERVICE = 'org.reportgen.acceptance.InstallerControl'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
NATIVE_ENV_UNSET = ('LD_LIBRARY_PATH', 'PYTHONHOME', 'PYTHONPATH', 'URE_BOOTSTRAP', 'UNO_PATH')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def stable_version(value):
    return bool(re.match(r'^LibreOffice\s+\d+(?:\.\d+)+\b', value.strip())) and not re.search(
        r'(?i)(?:dev|alpha|beta|rc\d)', value)


def native_environment():
    """Use distribution libraries for distribution Python/Writer/unopkg.

    In particular setup-python's libpython must not be injected into Writer's
    embedded distribution Python loader. The report-gen Python stays untouched.
    """
    environment = dict(os.environ)
    for name in NATIVE_ENV_UNSET:
        environment.pop(name, None)
    environment['SAL_USE_VCLPLUGIN'] = 'svp'
    return environment


def run(command, output, *, timeout=180, env=None):
    """Bound the whole private process group, retaining failure diagnostics."""
    started = time.monotonic()
    process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                               start_new_session=os.name == 'posix', env=env)
    expired = False
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        expired = True
        if os.name == 'posix':
            os.killpg(process.pid, signal.SIGKILL)
        else:
            process.kill()
        stdout, stderr = process.communicate()
    receipt = {'command': command, 'returncode': process.returncode, 'timeout': expired,
               'elapsed_seconds': round(time.monotonic() - started, 3), 'stdout': stdout, 'stderr': stderr}
    write_json(output, receipt)
    if expired or process.returncode:
        raise ValueError(f'Command failed; see {output}')
    return receipt


def fixture(path):
    """A compact independent native DOCX matrix, with exact bookmarked values."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Cm, Pt
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'DejaVu Sans Mono', Pt(10)
    normal.paragraph_format.space_after = Pt(8)
    url = ('https://settlement.example/batches/ALT%2D015/items/' +
           '0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ' * 4 + '?value=A%26B&literal=%252F#revision(15)')
    token = 'BATCH_ID_' + '0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ' * 5
    rows = []

    def add(parent, name, prefix, value, *, linked=False, table=False, listing=False, language=TAG):
        p = parent.add_paragraph(style='List Number' if listing else None)
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(len(rows) + 1)); start.set(qn('w:name'), name)
        p._p.append(start)
        p.add_run(prefix)
        target = p._p
        if linked:
            target = OxmlElement('w:hyperlink')
            target.set(qn('r:id'), document.part.relate_to(value, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            p._p.append(target)
        r = OxmlElement('w:r'); properties = OxmlElement('w:rPr'); lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), language); properties.append(lang); r.append(properties)
        text = OxmlElement('w:t'); text.text = value; r.append(text); target.append(r)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(len(rows) + 1)); p._p.append(end)
        rows.append({'name': name, 'prefix': prefix, 'value': value, 'literal': prefix + value,
                     'target': value if linked else None, 'table': table, 'language': language})

    document.add_heading('Optional Writer adapter acceptance', 0)
    document.add_paragraph('Synthetic data. Exact literals must remain editable without inserted break characters.')
    add(document, 'rgPlainUrl', 'PLAIN URL: ', url)
    add(document, 'rgLinkUrl', 'LINKED URL: ', url, linked=True)
    add(document, 'rgToken', 'IDENTIFIER: ', token)
    add(document, 'rgUnicode', 'Unicode é e\u0301 \U0001f642: ', url)
    add(document, 'rgList', 'LIST URL: ', url, linked=True, listing=True)
    table = document.add_table(rows=0, cols=1)
    table.autofit = False; table.columns[0].width = Cm(9)
    properties = table._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        item = OxmlElement('w:' + side); item.set(qn('w:w'), '0'); item.set(qn('w:type'), 'dxa'); margins.append(item)
    properties.append(margins)
    for name, prefix, value in [('rgCellUrl', 'CELL URL: ', url), ('rgCellToken', 'CELL ID: ', token)]:
        cell = table.add_row().cells[0]; cell.width = Cm(9)
        add(cell, name, prefix, value, linked=name == 'rgCellUrl', table=True)
    add(document, 'rgControl', 'CONTROL: ', 'An ordinary English sentence with accents café and naïve.', language='en-US')
    document.save(path)
    return rows


def build_installer_control(path, compression=zipfile.ZIP_STORED):
    """A test-only Python UNO control following the official SDK registration.

    https://api.libreoffice.org/examples/python/minimal-extension/main.py
    It has no break iterator, no document edits and is never a release asset.
    """
    module = '''import unohelper
from com.sun.star.task import XJobExecutor
from com.sun.star.lang import XServiceInfo

class Control(unohelper.Base, XJobExecutor, XServiceInfo):
    def __init__(self, context):
        self.context = context
    def trigger(self, argument):
        if argument != "report-gen-native-acceptance":
            raise ValueError("Unexpected control invocation")
    def getImplementationName(self):
        return "org.reportgen.acceptance.InstallerControl"
    def supportsService(self, service):
        return service == self.getImplementationName()
    def getSupportedServiceNames(self):
        return (self.getImplementationName(),)

g_ImplementationHelper = unohelper.ImplementationHelper()
g_ImplementationHelper.addImplementation(Control, "org.reportgen.acceptance.InstallerControl",
                                         ("org.reportgen.acceptance.InstallerControl",))
'''
    payload = {
        'main.py': module,
        'description.xml': '<?xml version="1.0" encoding="UTF-8"?>\n<description xmlns="http://openoffice.org/extensions/description/2006"><identifier value="org.reportgen.acceptance.installer-control"/><version value="1.0"/><display-name><name>Native installer control</name></display-name></description>',
        'META-INF/manifest.xml': '<?xml version="1.0" encoding="UTF-8"?>\n<manifest:manifest xmlns:manifest="http://openoffice.org/2001/manifest"><manifest:file-entry manifest:media-type="application/vnd.sun.star.uno-component;type=Python" manifest:full-path="main.py"/></manifest:manifest>',
    }
    if compression not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
        raise ValueError('Unsupported installer-control compression')
    with zipfile.ZipFile(path, 'w', compression) as archive:
        for name, data in sorted(payload.items()):
            entry = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            entry.compress_type = compression
            entry.external_attr = 0o100644 << 16
            archive.writestr(entry, data.encode('utf-8'))
    return path


def preserve_profile_metadata(profile, output):
    """Retain only this disposable test profile's bookkeeping for diagnosis."""
    source = Path(profile)
    entries = []
    for path in sorted(source.rglob('*')):
        if path.is_symlink() or not path.is_file():
            continue
        relative = path.relative_to(source)
        entries.append({'path': relative.as_posix(), 'bytes': path.stat().st_size, 'sha256': digest(path)})
        if path.suffix in ('.xml', '.xcu', '.xcs', '.ini', '.pmap', '.properties') and path.stat().st_size <= 1024 * 1024:
            target = output / 'profile-metadata' / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(path, target)
    write_json(output / 'profile-files.json', entries)


def initialize_profile(args, profile, output, environment):
    """Complete Writer's documented first initialization, then close it."""
    run([args.executable, '-env:UserInstallation=' + Path(profile).as_uri(), '--headless',
         '--nologo', '--nodefault', '--norestore', '--terminate_after_init'],
        output / 'profile-initialization.json', timeout=30, env=environment)


def inspect_docx(path, rows):
    from lxml import etree
    with zipfile.ZipFile(path) as archive:
        tree = etree.fromstring(archive.read('word/document.xml'))
        relationships = etree.fromstring(archive.read('word/_rels/document.xml.rels'))
    targets = {item.get('Id'): item.get('Target') for item in relationships}
    observed = []
    for row in rows:
        paragraphs = tree.xpath('//w:p[w:bookmarkStart[@w:name=$name]]', namespaces=NS, name=row['name'])
        if len(paragraphs) != 1:
            raise ValueError('Missing unique preserved bookmark: ' + row['name'])
        paragraph = paragraphs[0]
        value = ''.join(paragraph.xpath('.//w:t/text()', namespaces=NS))
        if value != row['literal']:
            raise ValueError('DOCX changed literal: ' + row['name'])
        marked = ''.join(paragraph.xpath('.//w:r[w:rPr/w:lang[@w:val=$tag]]/w:t/text()',
                                         namespaces=NS, tag=TAG))
        if row['language'] == TAG and row['value'] != marked:
            raise ValueError('DOCX lost literal language: ' + row['name'])
        if row['language'] != TAG and marked:
            raise ValueError('DOCX marked ordinary prose as a technical literal: ' + row['name'])
        if row['target']:
            ids = paragraph.xpath('.//w:hyperlink/@r:id', namespaces={**NS, 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})
            destinations = [targets.get(key) for key in ids]
            # Writer can save an OOXML hyperlink as a complex HYPERLINK field,
            # retaining its fragment in the native \\l switch. Inspect that
            # authored destination, never the visible URL text as a substitute.
            codes = ''.join(paragraph.xpath('.//w:instrText/text()', namespaces=NS))
            for match in re.finditer(r'\bHYPERLINK\s+"([^"\r\n]*)"(?:\s+\\l\s+"([^"\r\n]*)")?', codes):
                destinations.append(match[1] + ('#' + match[2] if match[2] is not None else ''))
            if row['target'] not in destinations:
                raise ValueError('DOCX changed exact hyperlink: ' + row['name'])
        observed.append(row['name'])
    return {'exact_literals': observed, 'sha256': digest(path)}


def compact(value):
    return ''.join(character for character in value if not character.isspace())


def physical_characters(page, page_number):
    """Order this single-column LTR fixture by baseline/x, not font paint order.

    Writer can paint an emoji fallback font after the following punctuation.
    Keep original codepoints and combining marks; never normalize source text.
    The small baseline tolerance merges font runs, not adjacent report lines.
    """
    raw = [character for block in page.get_text('rawdict')['blocks']
           for line in block.get('lines', []) for span in line['spans']
           for character in span['chars'] if not character['c'].isspace()]
    lines = []
    for character in sorted(raw, key=lambda item: item['origin'][1]):
        baseline = character['origin'][1]
        if not lines or abs(lines[-1][0] - baseline) > .5:
            lines.append((baseline, []))
        lines[-1][1].append(character)
    result = []
    for line_number, (_, characters) in enumerate(lines):
        for character in sorted(characters, key=lambda item: item['origin'][0]):
            result.append({'c': character['c'], 'bbox': character['bbox'], 'line': (page_number, line_number)})
    return result


def inspect_pdf(path, rows, *, require_fill, reflow=False):
    import fitz
    all_chars, links, pages = [], [], []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document):
            all_chars.extend(physical_characters(page, page_number))
            links.extend(link.get('uri') for link in page.get_links() if link.get('uri'))
            png = path.with_name(path.stem + f'-{page_number+1:02}.png')
            page.get_pixmap(dpi=110, alpha=False).save(png)
            pages.append({'page': page_number + 1, 'png': png.name, 'sha256': digest(png)})
    text = ''.join(item['c'] for item in all_chars)
    checks = []
    for row in rows:
        needle = compact(row['literal'])
        if text.count(needle) != 1:
            raise ValueError('PDF lost or duplicated exact literal: ' + row['name'])
        first = text.index(needle) + len(compact(row['prefix']))
        characters = all_chars[first:first + len(compact(row['value']))]
        lines = defaultdict(list)
        for character in characters:
            lines[character['line']].append(character)
        # The fixture's one-column table has explicit 9 cm width and zero cell
        # padding. Body margins are 2.5 cm, or 3 cm after native editing.
        right = ((3 if reflow else 2.5) + 9) * 72 / 2.54 if row['table'] else (18 if reflow else 18.5) * 72 / 2.54
        prefix_last = all_chars[first - 1] if compact(row['prefix']) else None
        same_prefix_line = prefix_last is None or prefix_last['line'] == characters[0]['line']
        if require_fill and row['language'] == TAG and not same_prefix_line:
            first_width = characters[0]['bbox'][2] - characters[0]['bbox'][0]
            available = right - prefix_last['bbox'][2]
            if available > first_width * 2 + 4:
                raise ValueError('Premature break between prefix and literal: ' + row['name'])
        samples = []
        line_values = list(lines.values())
        if require_fill and row['language'] == TAG and any(line[-1]['bbox'][2] > right + 4 for line in line_values):
            raise ValueError('Premature/overflow literal wrap in ' + row['name'])
        for index, line in enumerate(line_values[:-1]):
            last = line[-1]['bbox']
            next_width = line_values[index + 1][0]['bbox'][2] - line_values[index + 1][0]['bbox'][0]
            gap = right - last[2]
            samples.append({'gap_pt': round(gap, 3), 'next_character_pt': round(next_width, 3)})
            if require_fill and row['language'] == TAG and (gap > next_width + 4 or gap < -4):
                raise ValueError(f'Premature/overflow literal wrap in {row["name"]}: gap={gap:.2f}, next={next_width:.2f}')
        checks.append({'name': row['name'], 'lines': len(lines), 'prefix_continues_on_same_line': same_prefix_line,
                       'nonfinal_line_fill': samples})
        if row['target'] and row['target'] not in links:
            # LibreOffice may normalize only unreserved URI escapes. The exact
            # authored destination remains mandatory in the saved DOCX above.
            def normalize(value):
                return re.sub(r'%([0-9a-fA-F]{2})', lambda m: chr(int(m[1], 16)) if chr(int(m[1], 16)) in 'ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-._~' else m[0].upper(), value)
            if normalize(row['target']) not in {normalize(value) for value in links}:
                raise ValueError('PDF lost hyperlink: ' + row['name'])
    return {'sha256': digest(path), 'pages': pages, 'literal_checks': checks, 'text_exact_ignoring_layout_whitespace': True}


def uno_worker(config_path):
    """Run using the distribution Python that owns pyuno, with no pip imports."""
    import uno
    config = json.loads(Path(config_path).read_text())
    output = Path(config['output'])
    connection = 'pipe,name=report-gen-oxt-' + uuid.uuid4().hex
    environment = native_environment()
    environment['SAL_USE_VCLPLUGIN'] = 'svp'
    def prop(name, value):
        result = uno.createUnoStruct('com.sun.star.beans.PropertyValue')
        result.Name, result.Value = name, value
        return result
    log = (output / 'office.log').open('wb')
    office = subprocess.Popen([config['executable'], '-env:UserInstallation=' + Path(config['profile']).as_uri(),
                               '--headless', '--nodefault', '--norestore', '--nologo',
                               '--accept=' + connection + ';urp;StarOffice.ServiceManager'],
                              stdout=log, stderr=log, env=environment)
    receipt = {'status': 'running', 'factory_registered_by_client': False, 'documents': []}
    document, desktop = None, None
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
        context, last_error = None, None
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline and office.poll() is None:
            try:
                context = resolver.resolve('uno:' + connection + ';urp;StarOffice.ComponentContext')
                break
            except Exception as error:
                last_error = str(error)
                time.sleep(.2)
        if context is None:
            raise ValueError(f'Writer did not become ready: exit={office.poll()}, {last_error}')
        manager = context.ServiceManager
        if config.get('control'):
            control = manager.createInstanceWithContext(CONTROL_SERVICE, context)
            if control is None or control.getImplementationName() != CONTROL_SERVICE:
                raise ValueError('Independently installed Python UNO control is unavailable')
            control.trigger('report-gen-native-acceptance')
            receipt['installer_control'] = {'implementation': control.getImplementationName(), 'invoked': True}
        adapter = manager.createInstanceWithContext(SERVICE, context)
        if config['enabled']:
            if adapter is None or adapter.getImplementationName() != IMPLEMENTATION:
                raise ValueError('The independently installed OXT is not available in Writer')
            receipt['implementation'] = adapter.getImplementationName()
            delegate = manager.createInstanceWithContext('com.sun.star.i18n.BreakIterator_Unicode', context)
            cases = [('plain', 'URL: ' + 'A1/' * 50, 39, TAG),
                     ('unicode-prefix', 'é e\u0301 \U0001f642: ' + 'A1/' * 50, 43, TAG),
                     ('other-private-locale', 'URL: ' + 'A1/' * 50, 39, 'zxx-ZZ-x-other'),
                     ('english', 'URL: ' + 'A1/' * 50, 39, 'en-US')]
            receipt['service_checks'] = []
            for name, text, position, variant in cases:
                locale = uno.createUnoStruct('com.sun.star.lang.Locale')
                locale.Language, locale.Country, locale.Variant = 'qlt', 'ZZ', variant
                if variant == 'en-US':
                    locale.Language, locale.Country, locale.Variant = 'en', 'US', ''
                hyphenation = uno.createUnoStruct('com.sun.star.i18n.LineBreakHyphenationOptions')
                options = uno.createUnoStruct('com.sun.star.i18n.LineBreakUserOptions')
                result = adapter.getLineBreak(text, position, locale, 0, hyphenation, options)
                default = delegate.getLineBreak(text, position, locale, 0, hyphenation, options)
                if variant == TAG and result.breakIndex != position:
                    raise ValueError('Installed adapter did not fill an eligible literal: ' + name)
                if variant != TAG and (result.breakIndex, result.breakType) != (default.breakIndex, default.breakType):
                    raise ValueError('Installed adapter modified an unrelated locale: ' + name)
                receipt['service_checks'].append({'name': name, 'break_index': result.breakIndex, 'default_index': default.breakIndex})
        elif adapter is not None and adapter.getImplementationName() == IMPLEMENTATION:
            raise ValueError('Control profile unexpectedly contains the optional adapter')
        desktop = manager.createInstanceWithContext('com.sun.star.frame.Desktop', context)
        source = Path(config['source'])
        stages = ['original', 'reopened', 'reflow', 'reopened-reflow'] if config['enabled'] else ['fallback']
        if config.get('control'):
            stages = []
        for stage in stages:
            document = desktop.loadComponentFromURL(source.as_uri(), '_blank', 0,
                                                     (prop('Hidden', True), prop('ReadOnly', False), prop('UpdateDocMode', 3)))
            if document is None:
                raise ValueError('Writer did not load the fixture')
            if stage == 'reflow':
                cursor = document.Text.createTextCursor(); cursor.gotoStart(False); cursor.gotoEnd(True)
                cursor.CharHeight = 12
                styles = document.StyleFamilies.getByName('PageStyles')
                style = styles.getByName(document.Text.createTextCursor().PageStyleName)
                style.LeftMargin = style.RightMargin = 3000
            verified = []
            for row in config['rows']:
                if not document.Bookmarks.hasByName(row['name']):
                    raise ValueError('Writer lost bookmark: ' + row['name'])
                actual = document.Bookmarks.getByName(row['name']).Anchor.String
                if actual != row['literal']:
                    raise ValueError('Writer changed literal: ' + row['name'])
                verified.append(row['name'])
            target = output / (stage + '.docx')
            document.storeToURL(target.as_uri(), (prop('FilterName', 'Office Open XML Text'), prop('Overwrite', True)))
            document.storeToURL((output / (stage + '.pdf')).as_uri(), (prop('FilterName', 'writer_pdf_Export'),))
            receipt['documents'].append({'stage': stage, 'exact_bookmark_literals': verified})
            document.close(True); document = None
            source = target
        if not desktop.terminate():
            raise ValueError('Writer refused graceful termination')
        office.wait(timeout=15)
        if office.returncode != 0:
            raise ValueError('Writer crashed during teardown: ' + str(office.returncode))
        receipt['status'] = 'passed'
    except Exception as error:
        receipt.update(status='failed', error=str(error))
        raise
    finally:
        if office.poll() is None:
            office.terminate()
            try:
                office.wait(timeout=8)
            except subprocess.TimeoutExpired:
                office.kill(); office.wait()
        log.close()
        receipt['office_returncode'] = office.returncode
        write_json(output / 'uno.json', receipt)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--executable', default='/usr/bin/libreoffice')
    parser.add_argument('--unopkg', default='/usr/bin/unopkg')
    parser.add_argument('--uno-python', default='/usr/bin/python3')
    args = parser.parse_args(argv)
    output = args.output.resolve(); output.mkdir(parents=True, exist_ok=False)
    receipt = {'status': 'running', 'created_utc': datetime.now(timezone.utc).isoformat(),
               'scope': 'Standalone optional OXT in stable LibreOffice; native editable DOCX and PDF, not Microsoft Word.',
               'visual_review': 'pending', 'source_sha': os.environ.get('GITHUB_SHA'), 'checks': {}}
    try:
        if hasattr(os, 'geteuid') and os.geteuid() == 0:
            raise ValueError('User OXT acceptance requires an ordinary non-root account; no shared installation fallback.')
        native_env = native_environment()
        write_json(output / 'native-environment.json', {
            'unset_for_distribution_commands': {name: os.environ[name] for name in NATIVE_ENV_UNSET if name in os.environ},
            'set_for_distribution_commands': {'SAL_USE_VCLPLUGIN': 'svp'},
            'scope': 'Only explicit distribution executables and their children; report-gen Python environment unchanged.'})
        version = run([args.executable, '--version'], output / 'version.json', timeout=20, env=native_env)['stdout'].strip()
        if not stable_version(version):
            raise ValueError('Acceptance requires a stable LibreOffice build: ' + version)
        receipt['libreoffice_version'] = version
        builder_path = ROOT / 'integrations/libreoffice/build_oxt.py'
        spec = importlib.util.spec_from_file_location('oxt_builder', builder_path)
        builder = importlib.util.module_from_spec(spec); spec.loader.exec_module(builder)
        if (builder.LANGUAGE_TAG, builder.SERVICE, builder.IMPLEMENTATION) != (TAG, SERVICE, IMPLEMENTATION):
            raise ValueError('Acceptance fixture and OXT identity disagree')
        oxt = builder.build_oxt(output / 'report-gen-technical-literals.oxt')
        receipt['oxt_sha256'] = digest(oxt)
        source = output / 'fixture.docx'; rows = fixture(source)
        write_json(output / 'expected.json', rows)
        receipt['checks']['source'] = inspect_docx(source, rows)
        phase_errors = {}
        for mode, compression in [('stored', zipfile.ZIP_STORED), ('deflated', zipfile.ZIP_DEFLATED)]:
            name = 'installer-control-' + mode
            control_oxt = build_installer_control(output / (name + '.oxt'), compression)
            control_folder = output / name; control_folder.mkdir()
            try:
                with tempfile.TemporaryDirectory(prefix='report-gen-oxt-control-') as profile:
                    try:
                        initialize_profile(args, profile, control_folder, native_env)
                        run([args.unopkg, '-env:UserInstallation=' + Path(profile).as_uri(), 'add', '--verbose', '--log-file', str(control_folder / 'unopkg.log'), '--suppress-license', str(control_oxt)], control_folder / 'installation.json', timeout=60, env=native_env)
                        config = {'control': True, 'enabled': False, 'output': str(control_folder), 'profile': profile,
                                  'source': str(source), 'rows': [], 'executable': args.executable}
                        write_json(control_folder / 'config.json', config)
                        run([args.uno_python, str(Path(__file__).resolve()), '--worker', str(control_folder / 'config.json')], control_folder / 'worker.json', env=native_env)
                    finally:
                        preserve_profile_metadata(profile, control_folder)
                receipt['checks'][name] = {'status': 'passed', 'sha256': digest(control_oxt)}
            except Exception as error:
                phase_errors[name] = str(error)
                write_json(control_folder / 'phase-error.json', {'status': 'failed', 'error': str(error)})
        for enabled in (False, True):
            name = 'enabled' if enabled else 'fallback'
            folder = output / name; folder.mkdir()
            try:
                with tempfile.TemporaryDirectory(prefix='report-gen-oxt-' + name + '-') as profile:
                    try:
                        if enabled:
                            initialize_profile(args, profile, folder, native_env)
                            run([args.unopkg, '-env:UserInstallation=' + Path(profile).as_uri(), 'add', '--verbose', '--log-file', str(folder / 'unopkg.log'), '--suppress-license', str(oxt)], folder / 'installation.json', timeout=60, env=native_env)
                            run([args.unopkg, '-env:UserInstallation=' + Path(profile).as_uri(), 'list'], folder / 'installed-list.json', timeout=30, env=native_env)
                        config = {'enabled': enabled, 'output': str(folder), 'profile': profile, 'source': str(source),
                                  'rows': rows, 'executable': args.executable}
                        write_json(folder / 'config.json', config)
                        run([args.uno_python, str(Path(__file__).resolve()), '--worker', str(folder / 'config.json')], folder / 'worker.json', env=native_env)
                    finally:
                        preserve_profile_metadata(profile, folder)
                stages = ['original', 'reopened', 'reflow', 'reopened-reflow'] if enabled else ['fallback']
                results = {}
                for stage in stages:
                    results[stage] = {'docx': inspect_docx(folder / (stage + '.docx'), rows),
                                      'pdf': inspect_pdf(folder / (stage + '.pdf'), rows,
                                                         require_fill=enabled, reflow='reflow' in stage)}
                receipt['checks'][name] = results
            except Exception as error:
                # Continue to the independent profile to collect its evidence.
                # Any phase failure still fails the complete acceptance gate.
                phase_errors[name] = str(error)
                write_json(folder / 'phase-error.json', {'status': 'failed', 'error': str(error)})
        if phase_errors:
            receipt['phase_errors'] = phase_errors
            raise ValueError('Native acceptance phase failed: ' + json.dumps(phase_errors, ensure_ascii=False))
        fallback_checks = receipt['checks']['fallback']['fallback']['pdf']['literal_checks']
        if not any(sample['gap_pt'] > sample['next_character_pt'] + 24 for check in fallback_checks
                   if check['name'] != 'rgControl' for sample in check['nonfinal_line_fill']):
            raise ValueError('Negative control did not expose the known premature-wrap defect')
        original = receipt['checks']['enabled']['original']['pdf']['literal_checks']
        reflow = receipt['checks']['enabled']['reflow']['pdf']['literal_checks']
        if original == reflow:
            raise ValueError('Native edit did not exercise reflow')
        receipt['status'] = 'passed'
        return 0
    except Exception as error:
        receipt.update(status='failed', error=str(error))
        print(str(error), file=sys.stderr)
        return 1
    finally:
        write_json(output / 'validation.json', receipt)


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--worker':
        uno_worker(sys.argv[2])
    else:
        raise SystemExit(main())
