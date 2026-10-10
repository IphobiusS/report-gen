#!/usr/bin/env python3
"""Exercise 0.17 with real browser, disk, worker processes and document downloads."""
import argparse
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import zipfile
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'webapp')]
import app as api  # noqa: E402
from version import __version__  # noqa: E402

parser = argparse.ArgumentParser()
parser.add_argument('--output', required=True, type=Path)
args = parser.parse_args()
out = args.output.resolve(); out.mkdir(parents=True, exist_ok=True)
checks, errors, external = [], [], []


def check(name, value, detail=None):
    checks.append(dict(check=name, passed=bool(value), detail=detail))
    (out / 'checks.json').write_text(json.dumps(dict(checks=checks, errors=errors, external=external), ensure_ascii=False, indent=2), encoding='utf-8')
    assert value, (name, detail)


with tempfile.TemporaryDirectory(prefix='report-gen-0170-audit-') as temporary:
    api.PROJECTS = Path(temporary) / 'projects'; api.PROJECTS.mkdir()
    api.app.config['LIBRARY_PATH'] = str(Path(temporary) / 'library/findings.json')
    api.app.config['DESIGN_LIBRARY_PATH'] = str(Path(temporary) / 'library/designs.json')
    server = api.create_local_server(0)
    threading.Thread(target=server.run, daemon=True).start()
    base = f'http://127.0.0.1:{server.effective_port}'
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(executable_path=os.environ.get('RG_BROWSER_PATH'))
            context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True, reduced_motion='reduce')
            page = context.new_page()
            page.on('pageerror', lambda error: errors.append(str(error)))
            context.on('request', lambda request: external.append(request.url) if not request.url.startswith((base + '/', 'data:', 'blob:', 'about:')) else None)
            page.goto(base)
            page.locator('.workspace-dashboard').wait_for()
            check('Release version from API shown in dashboard', __version__ in page.locator('.brand-caption').inner_text())
            check('No project cannot import scanner findings', page.locator('#intakeBtn').is_disabled())
            page.get_by_role('button', name='Probar demo Nimbus', exact=True).click()
            page.wait_for_function('()=>S.slug && !projectStore.loading && projectBusyCount===0')
            original = page.evaluate('structuredClone(S.data)'); slug = page.evaluate('S.slug')
            check('Both dogs retained', 'Prometeo y Evaristo' in page.locator('.workspace-mascots img').get_attribute('alt'))

            def request(path, body=None, etag=None):
                return page.evaluate('''async ({path,body,etag})=>{
                  const r=await appFetch(path,body===null?{}:{method:'POST',headers:{'Content-Type':'application/json',...(etag?{'If-Match':etag}:{})},body:JSON.stringify(body)});
                  return {status:r.status,data:await r.json(),etag:r.headers.get('ETag')};
                }''', dict(path=path, body=body, etag=etag))

            def save(data):
                page.evaluate('''async data=>{S.data=data;projectStore.data=S.data;scheduleSave();if(!await doSave())throw Error('save failed');}''', data)

            page.route('**/api/projects/' + slug, lambda route: route.fulfill(status=503, content_type='application/json', body='{"error":"Simulated disk failure"}') if route.request.method == 'PUT' else route.continue_())
            edited = deepcopy(original); edited['meta']['report_title'] = 'Nimbus validation 0.17'
            page.evaluate('''data=>{S.data=data;projectStore.data=S.data;scheduleSave();}''', edited)
            page.locator('#saveError').wait_for()
            check('Save failure persists and pending edits remain', page.evaluate("projectStore.dirty && document.querySelector('#saveState').dataset.state==='error'"))
            page.evaluate("setSaveState('export succeeded')")
            check('Unrelated export cannot overwrite failed save status', page.locator('#saveState').get_attribute('data-state') == 'error')
            page.screenshot(path=str(out / 'save-error.png'), full_page=True)
            page.unroute('**/api/projects/' + slug)
            page.get_by_role('button', name='Reintentar guardado', exact=True).click()
            page.locator('#saveError').wait_for(state='detached')
            check('Retry preserves and saves latest text', request('/api/projects/' + slug)['data']['meta']['report_title'] == 'Nimbus validation 0.17')
            check('Retry returns focus to editor', page.evaluate("document.activeElement.id==='editor'"))
            check('Successful save has actual timestamp', page.evaluate('!!projectStore.lastSavedAt && !projectStore.dirty'))

            # Assisted import, remapping, preview details, dedup and review gating.
            page.locator('#projectActions > summary').click()
            page.locator('#intakeBtn').click()
            d = page.locator('dialog.progress-dialog')
            csv = 'Nombre;Nivel;URL;Texto\nImport demo;low;https://example.test;Two lines\nImport demo;low;https://example.test;Duplicate row\n'
            d.locator('input[type=file]').set_input_files(dict(name='scanner.csv', mimeType='text/csv', buffer=csv.encode()))
            d.locator('[data-map=title]').wait_for()
            for field, name in [('title', 'Nombre'), ('severity', 'Nivel'), ('affected', 'URL'), ('description_md', 'Texto')]:
                d.locator('[data-map=' + field + ']').select_option(name)
            d.get_by_role('button', name='Previsualizar mapeo', exact=True).click()
            d.locator('[data-intake-row="0"]').wait_for()
            check('CSV column mapping presents actual title', d.locator('.intake-rows').inner_text().startswith('Import demo'))
            check('Duplicate row disabled before any import', d.locator('[data-intake-row="1"]').is_disabled())
            check('Preview does not modify source', len(request('/api/projects/' + slug)['data']['findings']) == len(original['findings']))
            for width in (390, 768, 1440):
                page.set_viewport_size(dict(width=width, height=1000))
                check('Importer fits viewport ' + str(width), d.evaluate('(d)=>d.scrollWidth<=d.clientWidth+1'))
            page.set_viewport_size(dict(width=1440, height=1000))
            page.locator('.toast').first.wait_for(state='detached') if page.locator('.toast').count() else None
            page.screenshot(path=str(out / 'import-review.png'), full_page=True)
            d.get_by_role('button', name='Importar selección como pendiente', exact=True).click()
            d.wait_for(state='detached')
            check('Importer returns focus to opener', page.evaluate("document.activeElement===document.querySelector('#projectActions > summary')"))
            check('Imported finding unreviewed and not verified', page.evaluate("S.data.findings.at(-1).status==='not_verified'&&!S.data.findings.at(-1).import_source.reviewed"))
            review = request('/api/projects/' + slug + '/review')['data']
            check('Delivery review blocks unreviewed imported finding', any(i['level'] == 'error' and 'importación' in i['message'] for i in review['issues']))
            page.evaluate("select('finding', S.data.findings.length-1)")
            page.locator('.import-origin input').check()
            page.evaluate('doSave()')
            check('Explicit review clears import block without confirming finding', page.evaluate("S.data.findings.at(-1).status==='not_verified'&&S.data.findings.at(-1).import_source.reviewed"))
            check('Import provenance persists after reload', request('/api/projects/' + slug)['data']['findings'][-1]['import_source']['reviewed'])
            # A second import uses Burp with a literal HTML attack payload.
            page.evaluate('ProgressUI.intake()'); d = page.locator('dialog.progress-dialog')
            d.locator('select[aria-label=Formato]').select_option('burp')
            xml = '<issues><issue><name>Burp fixture</name><severity>Medium</severity><host>https://example.test</host><path>/b</path><issueDetail><![CDATA[<p>&lt;img src=x onerror=alert(1)&gt;</p>]]></issueDetail></issue></issues>'
            d.locator('input[type=file]').set_input_files(dict(name='burp.xml', mimeType='application/xml', buffer=xml.encode()))
            d.locator('[data-intake-row="0"]').wait_for()
            d.locator('.intake-rows summary').click()
            check('Burp payload displayed as inert text', d.locator('.intake-rows img').count() == 0 and 'onerror' in d.locator('.intake-rows').inner_text())
            d.get_by_role('button', name='Cerrar', exact=True).click()
            # A slow response for an older selection must never replace a newer file preview.
            page.evaluate('ProgressUI.intake()'); d = page.locator('dialog.progress-dialog')
            held=[]
            def delay_first(route):
                if 'Race A' in (route.request.post_data or ''): held.append(route)
                else: route.continue_()
            page.route('**/intake/preview', delay_first)
            d.locator('input[type=file]').set_input_files(dict(name='a.csv', mimeType='text/csv', buffer=b'title,severity\nRace A,low'))
            for _ in range(100):
                if held: break
                page.wait_for_timeout(20)
            check('Delayed preview fixture actually in flight', bool(held))
            d.locator('input[type=file]').set_input_files(dict(name='b.csv', mimeType='text/csv', buffer=b'title,severity\nRace B,high'))
            d.get_by_text('Race B', exact=True).wait_for()
            held[0].continue_(); page.wait_for_timeout(300)
            check('Stale preview cannot replace newer file selection', 'Race B' in d.locator('.intake-rows').inner_text() and 'Race A' not in d.locator('.intake-rows').inner_text())
            page.unroute('**/intake/preview')
            d.get_by_role('button', name='Cerrar', exact=True).click()

            # Reset disposable fixture, then create two named deliveries and evidence variants.
            save(edited)
            tag = page.evaluate('projectStore.etag')
            a = request('/api/projects/' + slug + '/history', dict(label='First delivery', stage='delivered'), tag)['data']['id']
            changed = deepcopy(edited)
            changed['findings'][0]['description_md'] += '\n\nAdditional reproduction context verified in the lab.'
            changed['findings'][0]['status'] = 'in_progress'
            save(changed)
            # Replace a processed screenshot inside this isolated audit project, keeping its name.
            from PIL import Image, ImageDraw
            source = api.PROJECTS / slug / changed['evidence'][0]['src']
            with Image.open(source) as im:
                im = im.convert('RGB'); ImageDraw.Draw(im).rectangle((5, 5, 160, 40), fill='#183e58'); im.save(source)
            b = request('/api/projects/' + slug + '/history', dict(label='Retest draft', stage='retest'), page.evaluate('projectStore.etag'))['data']['id']
            page.evaluate("select('workflow','history')")
            page.locator('.wf-diff').wait_for(state='attached')
            page.evaluate('''({a,b})=>ProgressUI.compare(document.querySelector('.wf-diff'),S.slug,a,b)''', dict(a=a, b=b))
            frame = page.frame_locator('.comparison-frame')
            frame.locator('ins').first.wait_for()
            check('Visual comparison highlights edits', frame.locator('ins').count() > 0 and frame.locator('del').count() > 0)
            check('Old and new images are actually embedded', frame.locator('img').count() == 2)
            page.screenshot(path=str(out / 'visual-comparison.png'), full_page=True)
            with page.expect_download() as download:
                page.get_by_role('link', name='Descargar comparación HTML', exact=True).click()
            html_path = out / 'comparison.html'; download.value.save_as(html_path)
            check('Readable comparison downloads as standalone HTML', 'data:image/' in html_path.read_text(encoding='utf-8'))
            comparison_page = context.new_page(); comparison_page.goto(html_path.as_uri())
            comparison_page.wait_for_function('()=>document.fonts.status==="loaded" && [...document.images].every(i=>i.complete&&i.naturalWidth)')
            comparison_page.pdf(path=str(out / 'comparison.pdf'), print_background=True, prefer_css_page_size=True)
            comparison_page.close()

            # Full Nimbus package via queue; no client-side document construction.
            save(edited)
            page.evaluate("select('home')")
            page.evaluate('ExperienceUI.prepareDelivery()')
            d = page.locator('dialog.experience-dialog')
            d.locator('[data-delivery=ack]').check()
            d.get_by_role('button', name='Generar paquete', exact=True).click()
            d = page.locator('dialog.progress-dialog'); d.locator('.job-card').wait_for()
            check('Queue shows real snapshot revision', d.locator('.job-card small').inner_text().endswith(page.evaluate('projectStore.etag.replaceAll(\'"\',\'\').slice(0,12)')))
            page.screenshot(path=str(out / 'export-queue.png'), full_page=True)
            d.locator('.job-card[data-stage=done]').wait_for(timeout=180000)
            with page.expect_download() as download:
                d.get_by_role('button', name='Descargar', exact=True).click()
            bundle = out / 'nimbus-delivery.zip'; download.value.save_as(bundle)
            with zipfile.ZipFile(bundle) as z:
                check('Three real formats and verification manifest delivered', {'informe.pdf','informe.docx','markdown.zip','verificacion.json'}.issubset(z.namelist()))
                manifest = json.loads(z.read('verificacion.json'))
                check('Delivery manifest uses current release', manifest['tool_version'] == __version__)
                for filename in ('informe.pdf', 'informe.docx'): (out / filename).write_bytes(z.read(filename))
                check('All Nimbus image placements verified by worker', manifest['files']['informe.pdf']['checks']['image_placements'] >= manifest['files']['informe.pdf']['checks']['expected_image_placements'])
                from docx import Document
                doc = Document(BytesIO(z.read('informe.docx')))
                word_text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                import pdfplumber
                with pdfplumber.open(BytesIO(z.read('informe.pdf'))) as pdf: pdf_text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
                with zipfile.ZipFile(BytesIO(z.read('markdown.zip'))) as md: md_text = md.read('report.md').decode()
                for fmt, text in [('PDF', pdf_text), ('Word', word_text), ('Markdown', md_text)]:
                    check('All 12 Nimbus findings retained in ' + fmt, all(f['title'] in text for f in original['findings']))
            d.get_by_role('button', name='Cerrar', exact=True).click()

            # Active and queued cancellation use actual subprocesses.
            payload = dict(kind='export', formats=['pdf'], label='Cancel active PDF')
            first = request('/api/projects/' + slug + '/jobs', payload, page.evaluate('projectStore.etag'))['data']
            payload['label'] = 'Cancel queued PDF'
            second = request('/api/projects/' + slug + '/jobs', payload, page.evaluate('projectStore.etag'))['data']
            page.evaluate('ProgressUI.jobs()'); d = page.locator('dialog.progress-dialog')
            d.locator('[data-job="' + second['id'] + '"] button').first.click()
            d.locator('[data-job="' + first['id'] + '"] button').first.click()
            for job in (first, second):
                d.locator('[data-job="' + job['id'] + '"][data-stage=cancelled]').wait_for(timeout=30000)
                check('Cancellation leaves no downloadable partial ' + job['label'], request('/api/jobs/' + job['id'] + '/download')['status'] == 422)
            page.keyboard.press('Escape'); d.wait_for(state='detached')
            page.evaluate("setUiLang('en');select('home')")
            check('New flows localized in English', page.locator('#jobsBtn').text_content() == 'Exports', page.evaluate("[S.uiLang,document.querySelector('#jobsBtn').textContent]"))
            for width in (390, 768, 1366):
                page.set_viewport_size(dict(width=width, height=1000))
                check('Dashboard no horizontal overflow ' + str(width), page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'))
            page.set_viewport_size(dict(width=1440, height=1000))
            page.screenshot(path=str(out / 'dashboard-en.png'), full_page=True)
            check('No uncaught browser exceptions', not errors, errors)
            # file:// comparison is our explicit offline QA navigation, not a network request.
            check('No external HTTP requests', not [u for u in external if u.startswith(('https:', 'http:'))], external)
            browser.close()
    finally:
        queue = api.app.extensions.pop('export_queue', None)
        if queue: queue.close()
        server.close()
print(json.dumps(dict(passed=len(checks), output=str(out))))
