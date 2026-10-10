#!/usr/bin/env python3
"""Real browser regression for the 0.18 workbench, against isolated projects."""
import argparse
from copy import deepcopy
from io import BytesIO
import json
import os
import platform
from datetime import datetime, timezone
from pathlib import Path
import sys
import tempfile
import threading
import zipfile
from PIL import Image
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'webapp')]
import app as api  # noqa: E402
from version import __version__  # noqa: E402

parser = argparse.ArgumentParser()
parser.add_argument('--output', required=True, type=Path)
args = parser.parse_args()
out = args.output.resolve(); out.mkdir(parents=True, exist_ok=True)
checks, errors, external = [], [], []


def check(name, value, detail=None):
    checks.append(dict(check=name, passed=bool(value), detail=detail))
    (out / 'checks.json').write_text(json.dumps(dict(checks=checks, errors=errors, external=external), ensure_ascii=False, indent=2), encoding='utf-8')
    assert value, (name, detail)


with tempfile.TemporaryDirectory(prefix='report-gen-0180-audit-', ignore_cleanup_errors=True) as temporary:
    api.PROJECTS = Path(temporary) / 'projects'; api.PROJECTS.mkdir()
    api.app.config['LIBRARY_PATH'] = str(Path(temporary) / 'library/findings.json')
    api.app.config['DESIGN_LIBRARY_PATH'] = str(Path(temporary) / 'library/designs.json')
    server = api.create_local_server(0)
    threading.Thread(target=server.run, daemon=True).start()
    base = f'http://127.0.0.1:{server.effective_port}'
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(executable_path=os.environ.get('RG_BROWSER_PATH'))
            browser_version=browser.version
            context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True, reduced_motion='reduce')
            page = context.new_page()
            page.on('pageerror', lambda e: errors.append(str(e)))
            page.on('dialog', lambda d: d.accept())
            context.on('request', lambda r: external.append(r.url) if not r.url.startswith((base + '/', 'data:', 'blob:', 'about:')) else None)
            page.goto(base)
            page.locator('.workspace-dashboard').wait_for()
            check('Release identity visible', __version__ in page.locator('.brand-caption').inner_text())
            check('Capture unavailable without a project', page.locator('#captureBtn').is_disabled())
            page.get_by_role('button', name='Probar demo Nimbus', exact=True).click()
            page.wait_for_function('()=>S.slug && !projectStore.loading && projectBusyCount===0')
            page.wait_for_function("()=>document.querySelector('.attention-task')")
            original = page.evaluate('structuredClone(S.data)'); slug = page.evaluate('S.slug')
            check('Capture enabled after project load', page.locator('#navCaptures').is_enabled())
            check('Contextual pending retest from actual data', 'corrección por verificar' in page.locator('.workspace-next').inner_text())
            check('Three navigation groups', page.locator('#navWork,#navReview,#navTools').count() == 3)
            check('Simplified dog symbol rendered', page.locator('.brand-logo').evaluate('(i)=>i.complete&&i.naturalWidth>0&&i.src.endsWith("brand-symbol.svg")'))
            check('Detailed dogs retained on welcome', 'Prometeo y Evaristo' in page.locator('.workspace-mascots img').get_attribute('alt'))
            check('Credits moved out of work area', not page.locator('.app-credit').count())
            page.wait_for_function("()=>!document.querySelector('.workspace-next [role=status]').textContent.includes('Comprobando')")
            page.screenshot(path=str(out / 'dashboard-es.png'), full_page=True)
            page.locator('#aboutBtn').click();page.keyboard.press('Escape')
            page.wait_for_function("()=>!document.querySelector('dialog.workbench-dialog')")
            check('About close restores focus',page.locator('#aboutBtn').evaluate('(e)=>e===document.activeElement'))

            def request(path, body=None, etag=None, method=None):
                return page.evaluate('''async ({path,body,etag,method})=>{
                  const r=await appFetch(path,{method:method||(body===null?'GET':'POST'),headers:{'Content-Type':'application/json',...(etag?{'If-Match':etag}:{})},...(body===null?{}:{body:JSON.stringify(body)})});
                  return {status:r.status,data:await r.json(),etag:r.headers.get('ETag')};
                }''', dict(path=path, body=body, etag=etag, method=method))

            def save(data):
                page.evaluate('''async data=>{S.data=data;projectStore.data=S.data;scheduleSave();if(!await doSave())throw Error('save failed');}''', data)

            def idle():
                page.wait_for_function('()=>!projectStore.loading&&projectBusyCount===0')

            page.locator('#navReport').click()
            page.locator('#findingList li').nth(1).click()
            page.locator('.finding-context').wait_for()
            page.locator('#editor').evaluate('(e)=>e.scrollTop=450')
            page.wait_for_timeout(180)
            position = page.locator('#editor').evaluate('(e)=>e.scrollTop')
            uid = page.evaluate('S.data.findings[S.sel.idx].uid')
            page.locator('#navHome').click()
            page.get_by_role('button', name='Continuar informe', exact=True).click()
            page.wait_for_timeout(120)
            check('Resume restores finding by permanent identity', page.evaluate('S.data.findings[S.sel.idx].uid') == uid)
            check('Resume restores editor scroll', abs(page.locator('#editor').evaluate('(e)=>e.scrollTop') - position) < 5, {'before': position, 'after': page.locator('#editor').evaluate('(e)=>e.scrollTop')})
            page.get_by_role('button', name='Remediación', exact=True).click()
            check('Section shortcut focuses remediation field', page.evaluate("document.activeElement.dataset.field==='remediation_md'"))
            check('Focused field not hidden by contextual header', page.locator('[data-field=remediation_md]').bounding_box()['y'] >= page.locator('.finding-context').bounding_box()['y'] + page.locator('.finding-context').bounding_box()['height'] - 2)
            page.locator('.finding-tabs').get_by_role('button', name='Evidencias', exact=True).click()
            check('Evidence section reachable directly', page.evaluate("!!document.activeElement.closest('[data-workpart=evidence]')"))
            page.evaluate("select('finding',0)")
            page.wait_for_function("()=>{try{return document.querySelector('#pdfPreview').contentDocument?.getElementById(S.data.findings[0].id)}catch(e){return false}}")
            page.evaluate('WorkbenchUI.followPreview()')
            check('HTML preview follows selected finding', page.locator('#pdfPreview').evaluate('(f)=>f.contentWindow.scrollY>0'))
            page.locator('#previewWidth').fill('44')
            check('Preview width adjustable and retained', page.evaluate("localStorage.getItem('rg.wb.previewWidth')==='44'"))
            data_before = page.evaluate('JSON.stringify(S.data)')
            page.locator('#focusBtn').click()
            check('Focus mode frees sidebar space', not page.locator('.sidebar').is_visible())
            page.locator('#focusBtn').click(); page.locator('#densityBtn').click()
            check('Compact density applied', page.locator('body').get_attribute('data-density') == 'compact')
            check('Appearance controls do not edit report', page.evaluate('JSON.stringify(S.data)') == data_before)
            page.locator('#densityBtn').click()
            page.screenshot(path=str(out / 'editor-es.png'), full_page=True)

            steps_before = page.evaluate('JSON.stringify(S.data.findings[0].walkthrough)')
            page.locator('.block-head .danger').first.click()
            page.locator('[data-field=description_md]').fill('Description retained after undo')
            page.locator('#undoNotice button').click()
            check('Undo restores only deleted step', page.evaluate('JSON.stringify(S.data.findings[0].walkthrough)') == steps_before)
            check('Undo preserves subsequent field edit', page.locator('[data-field=description_md]').input_value() == 'Description retained after undo')
            page.locator('.block-head').nth(1).get_by_role('button', name='↑', exact=True).click()
            page.locator('#undoNotice button').click()
            check('Undo restores step order', page.evaluate('JSON.stringify(S.data.findings[0].walkthrough)') == steps_before)
            finding_ids = page.evaluate('S.data.findings.map(f=>f.uid)')
            page.locator('.finding-actions .danger').click()
            page.locator('#undoNotice button').click()
            check('Undo restores deleted finding identity and order', page.evaluate('S.data.findings.map(f=>f.uid)') == finding_ids)
            page.locator('[data-field=owner]').fill('UX assessor')
            page.locator('#navFindings').click()
            page.locator('.finding-filters select[aria-label=Responsable]').select_option('UX assessor')
            check('Owner filter narrows findings', page.locator('.dash-table tbody tr').count() == 1)
            page.get_by_role('button', name='Guardar vista', exact=True).click()
            d = page.locator('dialog.workbench-dialog')
            d.locator('input').fill('My review'); d.get_by_role('button', name='Guardar', exact=True).click()
            page.get_by_role('button', name='Limpiar filtros', exact=True).click()
            check('Clear filters restores all findings', page.locator('.dash-table tbody tr').count() == 12)
            page.locator('select[aria-label="Vistas guardadas"]').select_option('0')
            check('Saved view restores owner filter', page.locator('.dash-table tbody tr').count() == 1)
            page.get_by_role('button', name='Limpiar filtros', exact=True).click()
            page.locator('#focusBtn').click()
            check('Finding board fills focus mode', page.locator('#editor').bounding_box()['width'] > 1300)
            page.locator('#focusBtn').click()
            page.keyboard.press('Control+k'); page.locator('#globalSearch').fill('Ir a evidencias')
            page.wait_for_timeout(170);page.keyboard.press('ArrowDown');page.wait_for_timeout(180)
            check('Action search remains open with keyboard focus', page.locator('#searchResults').is_visible())
            page.keyboard.press('Enter')
            check('Action search executes navigation', page.evaluate("S.sel.type==='workflow'&&S.sel.idx==='evidence'"))

            page.evaluate('setUiLang("en");select("finding",0)')
            check('Finding field labels translated', page.get_by_text('Short term', exact=True).count() > 0 and not page.get_by_text('Corto plazo', exact=True).count())
            page.evaluate('select("section",S.data.report.sections[0].key)')
            page.locator('.mded-bar').first.wait_for()
            check('English Markdown tools translated', page.get_by_role('button',name='Bold',exact=True).count()>0 and not page.get_by_role('button',name='Negrita',exact=True).count())
            page.evaluate('select("finding",0)')
            page.locator('select[data-field=verification]').wait_for()
            check('Technical verification distinct from remediation status', page.locator('select[data-field=verification]').count() == 1 and page.locator('select[data-field=status]').count() == 1)
            before_lang = page.evaluate('S.data.meta.lang')
            page.evaluate('setUiLang("es")')
            check('UI language preserves report language', page.evaluate('S.data.meta.lang') == before_lang)
            page.evaluate('ProgressUI.intake()')
            d = page.locator('dialog.progress-dialog')
            check('Importer begins with file step', d.locator('[data-import-panel="1"]').is_visible() and not d.locator('[data-import-panel="2"]').is_visible())
            csv = 'Nombre;Nivel;URL;Texto\nImported UX;low;https://example.test;Context\nImported UX;low;https://example.test;Duplicate\n'
            d.locator('input[type=file]').set_input_files(dict(name='scanner.csv', mimeType='text/csv', buffer=csv.encode()))
            d.locator('[data-import-panel="2"]').wait_for(state='visible')
            for field, name in [('title', 'Nombre'), ('severity', 'Nivel'), ('affected', 'URL'), ('description_md', 'Texto')]:
                d.locator('[data-map='+field+']').select_option(name)
            check('Optional import columns initially collapsed', not d.locator('[data-map=cwe]').is_visible())
            page.screenshot(path=str(out / 'import-mapping.png'), full_page=True)
            d.get_by_role('button', name='Revisar selección', exact=True).click()
            d.locator('[data-import-panel="3"]').wait_for(state='visible')
            check('Duplicate is disabled at review', d.locator('[data-intake-row="1"]').is_disabled())
            d.get_by_role('button', name='Anterior', exact=True).click()
            check('Wizard back preserves selected mapping', d.locator('[data-map=title]').input_value() == 'Nombre')
            d.get_by_role('button', name='Revisar selección', exact=True).click()
            d.get_by_role('button', name='Importar selección como pendiente', exact=True).click();d.wait_for(state='detached');idle()
            imported=page.evaluate('S.data.findings.at(-1)')
            check('New import unverified, separate remediation state', imported['verification'] == 'unverified' and imported['status'] == 'open' and not imported['import_source']['reviewed'])
            page.evaluate('select("home")');page.wait_for_function("()=>document.querySelector('.workspace-next').textContent.includes('importación por revisar')")
            check('Dashboard shows new import attention', 'importación por revisar' in page.locator('.workspace-next').inner_text())
            page.locator('.attention-task').filter(has_text='importación por revisar').click()
            page.locator('.import-origin input').check(); page.evaluate('doSave()')
            check('Reviewing import never confirms vulnerability', page.evaluate('S.data.findings.at(-1).verification') == 'unverified')

            # A slow preview from an older file must not replace the newer file.
            delayed=[]
            intake_url=base+'/api/projects/'+slug+'/intake/preview'
            def hold_first(route):
                if 'Slow scanner' in route.request.post_data:
                    delayed.append(route)
                else: route.continue_()
            page.route(intake_url,hold_first)
            page.evaluate('ProgressUI.intake()');d=page.locator('dialog.progress-dialog')
            with page.expect_request(lambda r:r.url==intake_url and 'Slow scanner' in (r.post_data or '')):
                d.locator('input[type=file]').set_input_files(dict(name='slow.csv',mimeType='text/csv',buffer=b'title,severity\nSlow scanner,low'))
            page.wait_for_timeout(100)
            assert delayed, 'Delayed preview was not intercepted'
            d.locator('input[type=file]').set_input_files(dict(name='new.csv',mimeType='text/csv',buffer=b'title,severity\nCurrent scanner,high'))
            d.locator('[data-import-panel="2"]').wait_for(state='visible')
            with page.expect_response(lambda r:r.url==intake_url and 'Slow scanner' in (r.request.post_data or '')):
                delayed[0].continue_()
            page.wait_for_timeout(100)
            check('Stale import response cannot replace current file','Current scanner' in d.locator('.mapping-example').first.inner_text())
            d.get_by_role('button',name='Cerrar',exact=True).click();page.unroute(intake_url,hold_first)

            # Save failure is persistent and retryable without losing typed text.
            project_url=base+'/api/projects/'+slug
            def reject_save(route):
                if route.request.method=='PUT': route.fulfill(status=507,json={'error':'Simulated disk failure'})
                else: route.continue_()
            page.route(project_url,reject_save)
            page.locator('[data-field=title]').fill('Import reviewed without confirmation')
            saved=page.evaluate('doSave()')
            check('Failed save retains text and visible retry',not saved and page.locator('#saveError').is_visible() and page.locator('[data-field=title]').input_value()=='Import reviewed without confirmation')
            page.unroute(project_url,reject_save)
            page.get_by_role('button',name='Reintentar guardado',exact=True).click()
            page.locator('#saveError').wait_for(state='detached')
            check('Successful retry clears error and preserves edit',page.locator('[data-field=title]').input_value()=='Import reviewed without confirmation')

            findings_before = page.evaluate('structuredClone(S.data.findings)')
            page.evaluate('WorkbenchUI.capture()');d=page.locator('dialog.workbench-dialog')
            d.locator('[data-capture-kind]').select_option('http')
            d.locator('[data-capture-title]').fill('Captured HTTP request')
            d.locator('[data-capture-text]').fill('GET /private HTTP/1.1\nHost: example.test\nX-Lab: private')
            d.locator('[data-capture-note]').fill('Observed in an authorized lab')
            d.get_by_role('button',name='Guardar en bandeja',exact=True).click();d.wait_for(state='detached');idle()
            findings_after = page.evaluate('structuredClone(S.data.findings)')
            changes = [{'uid':a.get('uid'), 'fields':[k for k in set(a)|set(b) if a.get(k)!=b.get(k)]} for a,b in zip(findings_before,findings_after) if a!=b]
            check('Quick capture leaves findings unchanged',findings_after==findings_before,changes)
            page.locator('#navCaptures').click()
            check('Capture visible in private inbox',page.locator('.capture-card').count()==1)
            check('Inbox breadcrumb survives save updates',page.locator('#workspaceContext').inner_text()=='Bandeja de captura')
            page.locator('.capture-card select').select_option(original['findings'][0]['uid'])
            page.get_by_role('button',name='Incorporar al hallazgo',exact=True).click();idle()
            check('Explicit link adds exact command to finding',page.evaluate('S.data.findings[0].walkthrough.at(-1).command').startswith('GET /private'))
            check('Linked capture retained and filtered from pending',page.locator('.capture-card').count()==0 and page.evaluate('S.data.capture_inbox.length')==1)
            page.evaluate('WorkbenchUI.capture()');d=page.locator('dialog.workbench-dialog')
            image_bytes=BytesIO();Image.new('RGB',(120,70),'#164255').save(image_bytes,format='PNG')
            # Exercise a paste event with real image bytes. This is not an OS clipboard test.
            d.evaluate('''(d,bytes)=>{const f=new File([new Uint8Array(bytes)],'clipboard.png',{type:'image/png'});const transfer=new DataTransfer();transfer.items.add(f);d.dispatchEvent(new ClipboardEvent('paste',{clipboardData:transfer,bubbles:true,cancelable:true}));}''',list(image_bytes.getvalue()))
            check('Image paste switches capture kind',d.locator('[data-capture-kind]').input_value()=='image')
            d.locator('[data-capture-title]').fill('Pasted proof')
            d.get_by_role('button',name='Guardar en bandeja',exact=True).click();d.wait_for(state='detached');idle()
            page.locator('.capture-card img').wait_for()
            check('Pasted screenshot renders after save',page.locator('.capture-card img').evaluate('(i)=>i.complete&&i.naturalWidth===120'))
            page.screenshot(path=str(out/'capture-inbox.png'),full_page=True)
            page.get_by_role('button',name='Retirar de la bandeja',exact=True).click()
            page.locator('#undoNotice button').click()
            check('Undo restores inbox entry',page.locator('.capture-card img').count()==1)
            page.evaluate('doSave()');page.reload();page.wait_for_function('()=>S.slug&&!projectStore.loading&&projectBusyCount===0')
            page.locator('#navCaptures').click()
            check('Capture survives page reload',page.locator('.capture-card img').count()==1)

            # Saved-version isolation and a genuine complete Nimbus delivery.
            save(deepcopy(original));page.evaluate('select("home")')
            page.get_by_role('button',name='Preparar entrega',exact=True).click()
            d=page.locator('dialog.experience-dialog');d.locator('[data-delivery=ack]').check()
            revision=page.evaluate('projectStore.etag.replaceAll(\'"\',\'\')')
            d.get_by_role('button',name='Generar paquete',exact=True).click()
            page.locator('#exportDrawer .job-card').wait_for()
            check('Exports use a nonmodal region',page.locator('#exportDrawer').get_attribute('role')=='region' and not page.locator('#exportDrawer').evaluate('(d)=>d.matches(":modal")'))
            check('Export snapshot revision visible',page.locator('.job-card small').inner_text().endswith(revision[:12]))
            page.screenshot(path=str(out/'exports.png'),full_page=True)
            page.locator('#navReport').click()
            page.locator('[data-field=report_title]').fill('Edited while export runs')
            check('Editing remains available during export',page.locator('[data-field=report_title]').input_value()=='Edited while export runs')
            check('Panel closes when focus returns to editor',not page.locator('#exportDrawer').count())
            page.locator('#exportIndicator').click()
            page.locator('.job-card[data-stage=done]').wait_for(timeout=180000)
            with page.expect_download() as download:page.locator('.job-card[data-stage=done]').get_by_role('button',name='Descargar',exact=True).click()
            archive=out/'nimbus-delivery.zip';download.value.save_as(archive)
            with zipfile.ZipFile(archive) as z:
                manifest=json.loads(z.read('verificacion.json'))
                check('Delivery uses captured original revision',manifest['source_revision']==revision)
                check('Current release recorded in delivery',manifest['tool_version']==__version__)
                from docx import Document
                doc=Document(BytesIO(z.read('informe.docx')))
                word='\n'.join(p.text for p in doc.paragraphs)+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                import pdfplumber
                with pdfplumber.open(BytesIO(z.read('informe.pdf'))) as pdf:pdftext='\n'.join(p.extract_text() or '' for p in pdf.pages)
                with zipfile.ZipFile(BytesIO(z.read('markdown.zip'))) as md:mdtext=md.read('report.md').decode()
                for fmt,text in [('PDF',pdftext),('Word',word),('Markdown',mdtext)]:
                    check('All 12 Nimbus titles preserved in '+fmt,all(f['title'] in text for f in original['findings']))
                    check('Later edit absent from snapshot '+fmt,'Edited while export runs' not in text)
                check('Expected PDF image placements retained',manifest['files']['informe.pdf']['checks']['image_placements']>=manifest['files']['informe.pdf']['checks']['expected_image_placements'])
                for filename in ('informe.pdf','informe.docx'):(out/filename).write_bytes(z.read(filename))
                (out/'nimbus-verification.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
            page.keyboard.press('Escape')
            check('Export close returns focus to indicator',page.locator('#exportIndicator').evaluate('(e)=>e===document.activeElement'))
            page.evaluate('doSave()')
            jobs=[]
            for n in range(2):
                job=request('/api/projects/'+slug+'/jobs',dict(kind='export',formats=['pdf'],label='Cancellation '+str(n)),page.evaluate('projectStore.etag'))
                check('Cancellation fixture queued '+str(n),job['status']==202,job['status'])
                jobs.append(job['data']['id'])
            page.locator('#exportIndicator').click()
            # Cancel queued work first, then the active process.
            for job_id in reversed(jobs):
                row=page.locator('[data-job="'+job_id+'"]')
                row.get_by_role('button',name='Cancelar tarea',exact=True).click()
                page.locator('[data-job="'+job_id+'"][data-stage=cancelled]').wait_for()
                check('Drawer confirms cancellation '+job_id,request('/api/jobs/'+job_id+'/download')['status']==422)
            page.keyboard.press('Escape')
            page.set_viewport_size({'width':390,'height':1000})
            for name, action in [('finding', 'select("finding",0)'),('inbox','select("capture")'),('findings','select("findings")')]:
                page.evaluate(action)
                check('No mobile horizontal overflow in '+name,page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'),page.evaluate('({viewport:innerWidth,document:document.documentElement.scrollWidth})'))
            page.evaluate('ProgressUI.intake()');d=page.locator('dialog.progress-dialog')
            check('Import wizard fits mobile viewport',d.bounding_box()['width']<=390)
            d.get_by_role('button',name='Cerrar',exact=True).click()
            for lang in ('en','es'):
                page.evaluate('(l)=>{setUiLang(l);select("home")}',lang)
                for theme in ('dark','light'):
                    page.evaluate('(t)=>applyUiTheme(t)',theme)
                    for width in (360,390,768,1366):
                        page.set_viewport_size({'width':width,'height':1000})
                        check(f'No horizontal overflow: {lang} {theme} {width}',page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'))
                if lang=='en':
                    page.set_viewport_size({'width':1440,'height':1000});page.screenshot(path=str(out/'dashboard-en-light.png'),full_page=True)
            page.evaluate('applyUiTheme("dark");select("home")');page.set_viewport_size({'width':390,'height':1000});page.screenshot(path=str(out/'dashboard-mobile.png'),full_page=True)
            check('No uncaught browser exceptions',not errors,errors)
            check('No external HTTP requests',not [x for x in external if x.startswith(('http:','https:'))],external)
            browser.close()
    finally:
        queue=api.app.extensions.pop('export_queue',None)
        if queue:queue.close()
        server.close()
result=dict(release=__version__,date=datetime.now(timezone.utc).isoformat(),python=platform.python_version(),platform=platform.platform(),browser='Chromium '+browser_version,checks=checks,errors=errors,external=external)
final=out/'checks.json.partial'
final.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
final.replace(out/'checks.json')
assert len(json.loads((out/'checks.json').read_text())['checks'])==len(checks)
print(json.dumps({'passed':len(checks),'browser':result['browser'],'output':str(out)}))
